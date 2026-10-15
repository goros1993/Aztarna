#!/usr/bin/env python3
"""
Genera un único DOCX a partir de los YAML de pruebas.

Uso:
    python yaml_a_docx.py salida/yaml                 -> pruebas.docx (usa indice.yaml)
    python yaml_a_docx.py salida/yaml -o mi_doc.docx
    python yaml_a_docx.py capitulo.yaml               (un solo archivo)

Requisitos:
    pip install python-docx ruamel.yaml
"""
import argparse
import re
import struct
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls
from docx.oxml.ns import qn
from docx.oxml.shape import CT_Inline
from docx.opc.constants import CONTENT_TYPE, RELATIONSHIP_TYPE
from docx.parts.image import ImagePart
from docx.shared import Cm, Pt, RGBColor
from ruamel.yaml import YAML

# ----------------------------------------------------------------------
# CONFIGURACIÓN (todo el aspecto visual se cambia aquí, en un solo sitio)
# ----------------------------------------------------------------------
FUENTE = "Calibri"
TAMANO = 10
COLOR_CABECERA = "00A651"      # verde de la cabecera
COLOR_TEXTO_CABECERA = RGBColor(0xFF, 0xFF, 0xFF)
ANCHO_TABLA_CM = 17.0          # ancho útil de la página A4 con márgenes de 2 cm
ANCHO_IMAGEN_CM = 10.0         # ancho de las imágenes incrustadas…
ALTO_MAX_IMAGEN_CM = 9.0       # …salvo que así queden más altas que esto (capturas verticales)
COLUMNAS = 6                   # rejilla base de la tabla

# Nombres de color -> hex (los mismos que usa convertir.py)
COLORES = {"rojo": "FF0000", "azul": "0070C0", "verde": "00B050"}
RESALTADOS = {
    "amarillo": WD_COLOR_INDEX.YELLOW,
    "verde": WD_COLOR_INDEX.BRIGHT_GREEN,
    "cian": WD_COLOR_INDEX.TURQUOISE,
    "yellow": WD_COLOR_INDEX.YELLOW,
    "bright_green": WD_COLOR_INDEX.BRIGHT_GREEN,
    "turquoise": WD_COLOR_INDEX.TURQUOISE,
}

# Cómo se muestra cada valor guardado en el YAML
AUTOMATIZABLE_VISIBLE = {"si": "Sí", "no": "No", "corto": "Corto",
                         "medio": "Medio", "largo": "Largo"}


# ----------------------------------------------------------------------
# Utilidades de formato de tabla
# ----------------------------------------------------------------------
def sombrear(celda, hex_color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def fijar_anchos(tabla, anchos_cm):
    """Ancho fijo de tabla y columnas (más rápido de maquetar para Word)."""
    tabla.autofit = False
    tblPr = tabla._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for col, ancho in zip(tabla.columns, anchos_cm):
        col.width = Cm(ancho)
    for fila in tabla.rows:
        for i, celda in enumerate(fila.cells):
            celda.width = Cm(anchos_cm[i])


def compactar(parrafo):
    pf = parrafo.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


# ----------------------------------------------------------------------
# Marcas inline: **negrita**, *cursiva*, [[color|texto]], [[res:color|texto]]
# ----------------------------------------------------------------------
TOKEN = re.compile(r"(\[\[[^\]|]+\||\]\]|\*\*|\*)")


def escribir_inline(parrafo, texto):
    negrita = cursiva = False
    pila = []  # lista de ("color"|"res", valor)
    for trozo in TOKEN.split(texto or ""):
        if not trozo:
            continue
        if trozo == "**":
            negrita = not negrita
        elif trozo == "*":
            cursiva = not cursiva
        elif trozo.startswith("[[") and trozo.endswith("|"):
            clave = trozo[2:-1]
            if clave.startswith("res:"):
                pila.append(("res", clave[4:]))
            else:
                pila.append(("color", clave))
        elif trozo == "]]":
            if pila:
                pila.pop()
        else:
            run = parrafo.add_run(trozo)
            run.bold = negrita or None
            run.italic = cursiva or None
            for tipo, valor in pila:
                if tipo == "color":
                    hexa = COLORES.get(valor, valor.lstrip("#"))
                    try:
                        run.font.color.rgb = RGBColor.from_string(hexa.upper())
                    except ValueError:
                        pass
                else:
                    ind = RESALTADOS.get(valor)
                    if ind is not None:
                        run.font.highlight_color = ind


# ----------------------------------------------------------------------
# Listas con numeración real de Word (la numeración reinicia en cada lista)
# ----------------------------------------------------------------------
def nueva_numeracion(doc, estilo):
    """Crea una instancia de numeración que empieza en 1 para el estilo dado."""
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    num_id = doc.styles[estilo].element.pPr.numPr.numId.val
    abstract_id = numbering.num_having_numId(num_id).abstractNumId.val
    nuevo = numbering.add_num(abstract_id)
    nuevo.add_lvlOverride(ilvl=0).add_startOverride(1)
    return nuevo.numId


def poner_numeracion(parrafo, num_id, nivel):
    pPr = parrafo._p.get_or_add_pPr()
    numPr = pPr.get_or_add_numPr()
    numPr.get_or_add_ilvl().val = nivel
    numPr.get_or_add_numId().val = num_id


def insertar_imagen(celda, base_dir, imagen):
    ruta = base_dir / imagen
    if ruta.exists():
        run = celda.add_paragraph().add_run()
        agregar_imagen(run, ruta, Cm(ANCHO_IMAGEN_CM))
    else:
        print(f"  AVISO: imagen no encontrada: {ruta}")


def _tamaño_imagen(ruta):
    """(ancho, alto) en píxeles de una imagen normal, o None si no se puede leer."""
    try:
        from PIL import Image

        with Image.open(ruta) as imagen:
            return imagen.size
    except Exception:  # noqa: BLE001 - si no se puede leer, se usa el ancho fijo
        return None


def _ancho_limitado(ancho, ancho_px, alto_px):
    """Reduce el ancho si con él la imagen superaría ALTO_MAX_IMAGEN_CM de alto."""
    if ancho_px > 0 and alto_px > 0:
        alto_max = Cm(ALTO_MAX_IMAGEN_CM)
        if ancho * alto_px / ancho_px > alto_max:
            return int(alto_max * ancho_px / alto_px)
    return ancho


def agregar_imagen(run, ruta, ancho):
    if ruta.suffix.lower() != ".emf":
        tamaño = _tamaño_imagen(ruta)
        if tamaño:
            ancho = _ancho_limitado(ancho, *tamaño)
        run.add_picture(str(ruta), width=ancho)
        return

    blob = ruta.read_bytes()
    if len(blob) < 108 or blob[40:44] != b" EMF":
        raise ValueError(f"Archivo EMF inválido o truncado: {ruta}")

    izquierda, arriba, derecha, abajo = struct.unpack_from("<4i", blob, 24)
    ancho_marco = derecha - izquierda
    alto_marco = abajo - arriba
    if ancho_marco <= 0 or alto_marco <= 0:
        raise ValueError(f"El EMF tiene dimensiones inválidas: {ruta}")

    ancho = _ancho_limitado(ancho, ancho_marco, alto_marco)
    alto = int(ancho * alto_marco / ancho_marco)
    part = run.part
    image_part = ImagePart(
        part.package.next_partname("/word/media/image%d.emf"),
        CONTENT_TYPE.X_EMF,
        blob,
    )
    relationship_id = part.relate_to(image_part, RELATIONSHIP_TYPE.IMAGE)
    inline = CT_Inline.new_pic_inline(
        part.next_id,
        relationship_id,
        ruta.name,
        ancho,
        alto,
    )
    run._r.add_drawing(inline)


def escribir_lista(doc, celda, items, numerada, base_dir):
    estilo = "List Number" if numerada else "List Bullet"
    num_id = nueva_numeracion(doc, estilo) if numerada else None
    for it in items or []:
        if isinstance(it, dict):
            texto, nivel = it.get("texto", ""), int(it.get("nivel", 0))
            imagenes = it.get("imagenes", [])
            if it.get("imagen"):
                imagenes = [*imagenes, it["imagen"]]
        else:
            texto, nivel, imagenes = str(it), 0, []
        p = celda.add_paragraph(style=estilo)
        compactar(p)
        if numerada:
            poner_numeracion(p, num_id, nivel)
        elif nivel:
            poner_numeracion(p, doc.styles[estilo].element.pPr.numPr.numId.val, nivel)
        escribir_inline(p, texto)
        for imagen in imagenes:
            insertar_imagen(celda, base_dir, imagen)


# ----------------------------------------------------------------------
# Construcción de la tabla de una prueba
# ----------------------------------------------------------------------
def etiqueta_valor(celda, etiqueta, valor):
    p = celda.paragraphs[0]
    compactar(p)
    r = p.add_run(f"{etiqueta}: ")
    r.bold = True
    escribir_inline(p, valor or "")


def tabla_prueba(doc, prueba, base_dir):
    t = doc.add_table(rows=6, cols=COLUMNAS)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fijar_anchos(t, [ANCHO_TABLA_CM / COLUMNAS] * COLUMNAS)

    def fusionar(fila, ini, fin):
        return t.cell(fila, ini).merge(t.cell(fila, fin))

    # Fila 0: ID | título (verde)
    c_id, c_tit = fusionar(0, 0, 1), fusionar(0, 2, 5)
    for c, txt, cursiva in ((c_id, prueba.get("id", ""), False),
                            (c_tit, prueba.get("titulo", ""), True)):
        sombrear(c, COLOR_CABECERA)
        p = c.paragraphs[0]
        compactar(p)
        r = p.add_run(txt)
        r.bold, r.italic = True, cursiva
        r.font.color.rgb = COLOR_TEXTO_CABECERA

    # Fila 1: descripción
    etiqueta_valor(fusionar(1, 0, 5), "Descripción", prueba.get("descripcion"))

    # Fila 2: componente | plataforma
    etiqueta_valor(fusionar(2, 0, 2), "Componente", prueba.get("componente"))
    etiqueta_valor(fusionar(2, 3, 5), "Plataforma", prueba.get("plataforma"))

    # Fila 3: situación inicial / pasos / situación final
    c = fusionar(3, 0, 5)
    primera = True
    for clave, titulo, numerada in (("situacion_inicial", "Situación inicial", False),
                                    ("pasos", "Pasos", True),
                                    ("situacion_final", "Situación final", False)):
        p = c.paragraphs[0] if primera else c.add_paragraph()
        primera = False
        compactar(p)
        p.paragraph_format.space_before = Pt(3)
        p.add_run(f"{titulo}:").bold = True
        escribir_lista(doc, c, prueba.get(clave), numerada, base_dir)

    # Fila 4: automatizable | criticidad | entorno
    auto = prueba.get("automatizable") or ""
    etiqueta_valor(fusionar(4, 0, 1), "Automatizable",
                   AUTOMATIZABLE_VISIBLE.get(auto, auto))
    etiqueta_valor(fusionar(4, 2, 3), "Criticidad", prueba.get("criticidad"))
    etiqueta_valor(fusionar(4, 4, 5), "Entorno de pruebas", prueba.get("entorno"))

    # Fila 5: referencia | fecha | versión
    etiqueta_valor(fusionar(5, 0, 2), "Referencia", prueba.get("referencia"))
    etiqueta_valor(fusionar(5, 3, 4), "Fecha", prueba.get("fecha"))
    etiqueta_valor(fusionar(5, 5, 5), "Versión", prueba.get("version"))

    for etiqueta, valor in (prueba.get("campos_adicionales") or {}).items():
        fila = t.add_row()
        celda = fila.cells[0].merge(fila.cells[-1])
        p = celda.paragraphs[0]
        compactar(p)
        p.add_run(f"{etiqueta}: ").bold = True
        lineas = valor if isinstance(valor, list) else str(valor or "").splitlines()
        if lineas:
            escribir_inline(p, str(lineas[0]))
            for linea in lineas[1:]:
                parrafo = celda.add_paragraph()
                compactar(parrafo)
                escribir_inline(parrafo, str(linea))

    for imagen in prueba.get("imagenes", []):
        fila = t.add_row()
        celda = fila.cells[0].merge(fila.cells[-1])
        insertar_imagen(celda, base_dir, imagen)

    # Espacio entre tablas
    sep = doc.add_paragraph()
    compactar(sep)
    return t


def tabla_informativa(doc, tabla, base_dir):
    """Tabla que no es una prueba (historial, parámetros…), tal como venía en el Word.

    Cada fila es una lista de celdas; una celda es un texto (con marcas de
    formato y saltos de línea) o {texto, imagenes, columnas}.
    """
    filas = [f for f in (tabla.get("filas") or []) if isinstance(f, list)]
    ancho = lambda celda: max(int(celda.get("columnas", 1)), 1) if isinstance(celda, dict) else 1
    columnas = max([sum(ancho(c) for c in f) for f in filas] or [1]) or 1
    t = doc.add_table(rows=len(filas) or 1, cols=columnas)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fijar_anchos(t, [ANCHO_TABLA_CM / columnas] * columnas)
    for i, fila in enumerate(filas):
        col = 0
        for j, celda in enumerate(fila):
            # La última celda ocupa lo que quede de la fila.
            fin = columnas - 1 if j == len(fila) - 1 else min(col + ancho(celda), columnas) - 1
            if col > columnas - 1:
                break
            destino = t.cell(i, col).merge(t.cell(i, fin)) if fin > col else t.cell(i, col)
            texto = celda.get("texto", "") if isinstance(celda, dict) else str(celda or "")
            for k, linea in enumerate(texto.split("\n")):
                p = destino.paragraphs[0] if k == 0 else destino.add_paragraph()
                compactar(p)
                escribir_inline(p, linea)
            for imagen in (celda.get("imagenes") or []) if isinstance(celda, dict) else []:
                insertar_imagen(destino, base_dir, imagen)
            col = fin + 1
    sep = doc.add_paragraph()
    compactar(sep)
    return t


# ----------------------------------------------------------------------
# Títulos con numeración automática (1 / 1.1 / 1.1.1 ...), como en Word
# ----------------------------------------------------------------------
NIVELES_TITULO = 6


def numerar_titulos(doc, primer_capitulo=1):
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(ids + [0]) + 1
    niveles = ""
    for i in range(NIVELES_TITULO):
        texto = ".".join(f"%{k + 1}" for k in range(i + 1))
        inicio = primer_capitulo if i == 0 else 1
        niveles += (
            f'<w:lvl w:ilvl="{i}"><w:start w:val="{inicio}"/><w:numFmt w:val="decimal"/>'
            f'<w:pStyle w:val="Heading{i + 1}"/><w:lvlText w:val="{texto}"/>'
            f'<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="864" w:hanging="864"/></w:pPr></w:lvl>')
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
        f'<w:multiLevelType w:val="multilevel"/>{niveles}</w:abstractNum>')
    primer_num = numbering.find(qn("w:num"))
    if primer_num is not None:
        primer_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    num_id = numbering.add_num(abstract_id).numId
    for i in range(NIVELES_TITULO):
        estilo = doc.styles[f"Heading {i + 1}"]
        estilo.font.name = FUENTE
        estilo.font.color.rgb = RGBColor(0, 0, 0)
        numPr = estilo.element.get_or_add_pPr().get_or_add_numPr()
        numPr.get_or_add_ilvl().val = i
        numPr.get_or_add_numId().val = num_id


# ----------------------------------------------------------------------
# Lectura de YAML y documento
# ----------------------------------------------------------------------
def cargar(entrada):
    yaml = YAML(typ="safe")
    ruta = Path(entrada)
    if ruta.is_file():
        archivos = [ruta]
    else:
        indice = ruta / "indice.yaml"
        if indice.exists():
            with open(indice, encoding="utf-8") as f:
                archivos = [ruta / e["archivo"] for e in yaml.load(f)]
        else:
            archivos = sorted(a for a in ruta.glob("*.yaml") if a.name != "indice.yaml")
    capitulos = []
    for a in archivos:
        with open(a, encoding="utf-8") as f:
            d = yaml.load(f)
        if d and "contenido" in d:
            capitulos.append((d.get("capitulo"), d["contenido"], a.parent))
        else:
            print(f"  AVISO: {a.name} no tiene la clave 'contenido', se ignora")
    return capitulos


def preparar_documento(primer_capitulo):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(TAMANO)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    numerar_titulos(doc, primer_capitulo)
    return doc


def titulo(doc, texto, nivel):
    doc.add_heading(texto, level=min(max(int(nivel or 1), 1), NIVELES_TITULO))


def escribir_contexto(doc, textos, base_dir=Path(".")):
    for texto in textos or []:
        if isinstance(texto, dict) and isinstance(texto.get("tabla"), dict):
            tabla_informativa(doc, texto["tabla"], base_dir)
            continue
        p = doc.add_paragraph()
        compactar(p)
        if isinstance(texto, dict):
            escribir_inline(p, texto.get("texto", ""))
            imagenes = texto.get("imagenes", [])
            if texto.get("imagen"):
                imagenes = [*imagenes, texto["imagen"]]
            for imagen in imagenes:
                ruta = base_dir / imagen
                if ruta.exists():
                    # Cada imagen en su propio párrafo, debajo del texto (no pegada a su última línea).
                    if p.text or p.runs:
                        p = doc.add_paragraph()
                        compactar(p)
                        p.paragraph_format.space_after = Pt(6)
                    agregar_imagen(p.add_run(), ruta, Cm(ANCHO_IMAGEN_CM))
                else:
                    print(f"  AVISO: imagen no encontrada: {ruta}")
        else:
            escribir_inline(p, str(texto))


def escribir_capitulos(doc, capitulos, al_escribir_prueba=None):
    """Escribe títulos, contexto y tablas de los capítulos; devuelve el número de pruebas.

    `al_escribir_prueba(doc, tabla, prueba, base_dir)`, si se indica, se llama
    tras cada tabla (por ejemplo, para añadir el resultado de una ejecución).
    """
    ids, total = set(), 0
    for cap, contenido, base_dir in capitulos:
        if cap:
            titulo(doc, cap, 1)
        for el in contenido:
            if "seccion" in el:
                titulo(doc, el["seccion"], el.get("nivel", 2))
                escribir_contexto(doc, el.get("contexto"), base_dir)
            elif "prueba" in el:
                p = el["prueba"]
                if p.get("nombre"):
                    titulo(doc, p["nombre"], p.get("nivel") or 3)
                escribir_contexto(doc, p.get("contexto"), base_dir)
                if p.get("id") in ids:
                    print(f"  AVISO: ID duplicado: {p.get('id')}")
                ids.add(p.get("id"))
                tabla = tabla_prueba(doc, p, base_dir)
                if al_escribir_prueba is not None:
                    al_escribir_prueba(doc, tabla, p, base_dir)
                total += 1
            elif "tabla" in el:
                tabla_informativa(doc, el["tabla"] or {}, base_dir)
            elif "texto" in el:
                escribir_contexto(doc, [el["texto"]], base_dir)
    return total


def generar_docx(entrada, salida, primer_capitulo=1):
    """Genera el DOCX a partir de los YAML y devuelve (pruebas, capítulos)."""
    capitulos = cargar(entrada)
    if not capitulos:
        raise ValueError("No se han encontrado YAML con la clave 'contenido'.")

    doc = preparar_documento(primer_capitulo)
    total = escribir_capitulos(doc, capitulos)
    doc.save(salida)
    return total, len(capitulos)


def main():
    ap = argparse.ArgumentParser(description="YAML de pruebas -> DOCX")
    ap.add_argument("entrada", help="carpeta con YAML (con indice.yaml) o un archivo YAML")
    ap.add_argument("-o", "--salida", default="pruebas.docx")
    ap.add_argument("--primer-capitulo", type=int, default=1,
                    help="número del primer Título 1 (p. ej. 8 si el original empieza en 8)")
    args = ap.parse_args()

    try:
        total, n_capitulos = generar_docx(args.entrada, args.salida, args.primer_capitulo)
    except ValueError as error:
        raise SystemExit(str(error))
    print(f"{total} pruebas en {n_capitulos} capítulos -> {args.salida}")


if __name__ == "__main__":
    main()
