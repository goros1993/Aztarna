#!/usr/bin/env python3
"""Comprueba que la importación del Word a YAML no pierde información.

Uso:
    python herramientas.py auditar documento.docx importacion/yaml -o importacion/auditoria.html

Compara las palabras del Word original con las de los YAML, tabla a tabla y
párrafo a párrafo (sin mirar el formato ni el orden de los campos, que en la
tabla común son distintos). Cada palabra que falta se clasifica:

- **perdida**: no está en el YAML y ninguna regla la quita. Hay que revisarla.
- **eliminada por regla**: era de un campo que el esquema elimina (componente,
  criticidad, referencia, resultado esperado, campos de resultado…).
- **descartada**: valor de entorno, maniobra o automatizable que no es una
  opción del esquema.

También comprueba los recuentos (pruebas, tablas que no son pruebas,
imágenes, IDs) y marca 2 o 3 tablas de cada familia como muestra para revisar
a ojo que cada texto ha ido al campo que toca.

El informe HTML lleva el texto del documento: no lo compartas. Lo que se
escribe por pantalla (y en auditoria.txt) solo tiene recuentos y estructura.
"""
from __future__ import annotations

import argparse
import html
import json
import re
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from migracion import convertir
from app import esquema
from migracion import inventario
from app.editor_core import read_chapters, tests_in_chapter

_MARCAS = re.compile(r"\[\[[^\]|]+\||\]\]|\*\*|\*")
_ETIQUETAS = sorted({*convertir.ETIQUETAS_CAMPOS, *convertir.ALIAS_ETIQUETAS_ADICIONALES, "EJECUCIÓN DE PRUEBA"},
                    key=len, reverse=True)
_PATRON_ETIQUETA = re.compile(r"(?<!\w)(" + "|".join(re.escape(e) for e in _ETIQUETAS) + r")\s*:", re.IGNORECASE)
_NORM_ETIQUETAS = {convertir.norm(e): e for e in _ETIQUETAS}
NO_TEXTO = {"imagenes", "imagen", "nivel", "columnas"}
# Etiquetas cuyo texto el esquema elimina (para marcar ahí las palabras «eliminadas»)…
ETIQUETAS_ELIMINADAS = {convertir.norm(e) for e, clave in convertir.ETIQUETAS_CAMPOS.items()
                        if clave in esquema.OBSOLETOS and clave != "situacion_final"} | {
                        convertir.norm(e) for e in convertir.ALIAS_ETIQUETAS_ADICIONALES}
# …y las de los campos con opciones (para marcar ahí los valores «descartados»).
ETIQUETAS_OPCIONES = {convertir.norm(e) for e, clave in convertir.ETIQUETAS_CAMPOS.items()
                      if clave in ("entorno", "maniobra", "automatizable")}
ESTADOS = {"perdida": "✖ pérdida", "eliminada": "⚠ solo eliminaciones", "ok": "✔ completa"}


# ----------------------------------------------------------------------
# Palabras
# ----------------------------------------------------------------------
def palabras(texto: Any) -> Counter:
    """Palabras normalizadas (sin tildes, en minúsculas, sin marcas de formato).

    Las marcas se quitan sin dejar hueco: una palabra con parte en negrita o en
    color («Ascen**sor**») sigue siendo una sola palabra.
    """
    texto = _MARCAS.sub("", str(texto or ""))
    return Counter(convertir.sin_tildes(w).lower() for w in re.findall(r"\w+", texto))


def palabras_de(valor: Any) -> Counter:
    """Palabras de todos los textos de un valor del YAML (también los nombres de los campos adicionales)."""
    total: Counter = Counter()
    if isinstance(valor, dict):
        for clave, hijo in valor.items():
            if clave in NO_TEXTO:
                continue
            if clave == "campos_adicionales" and isinstance(hijo, dict):
                for etiqueta, v in hijo.items():
                    total += palabras(etiqueta) + palabras_de(v)
            else:
                total += palabras_de(hijo)
    elif isinstance(valor, list):
        for hijo in valor:
            total += palabras_de(hijo)
    elif valor is not None:
        total += palabras(valor)
    return total


def sin_etiquetas(texto: str) -> str:
    """Texto de una celda sin las etiquetas («Descripción:», «Pasos:»…), que no son contenido."""
    lineas = []
    for linea in texto.splitlines():
        n = convertir.norm(linea).rstrip(":")
        if n in convertir.SECCIONES or n in convertir.ETIQUETAS_IGNORADAS or n in _NORM_ETIQUETAS:
            continue                                   # «Situación inicial:» sola en su línea
        lineas.append(_PATRON_ETIQUETA.sub(" ", linea))
    return "\n".join(lineas)


def segmentos(texto: str) -> list[tuple[str, str | None, bool]]:
    """Trozos del texto de una celda: (texto, etiqueta a la que pertenece, es la propia etiqueta)."""
    trozos: list[tuple[str, str | None, bool]] = []
    actual = None
    for i, linea in enumerate(texto.split("\n")):
        if i:
            trozos.append(("\n", actual, True))
        n = convertir.norm(linea).rstrip(":")
        if n and (n in convertir.SECCIONES or n in convertir.ETIQUETAS_IGNORADAS or n in _NORM_ETIQUETAS):
            actual = n
            trozos.append((linea, actual, True))
            continue
        pos = 0
        for m in _PATRON_ETIQUETA.finditer(linea):
            trozos.append((linea[pos:m.start()], actual, False))
            actual = convertir.norm(m.group(1))
            trozos.append((m.group(0), actual, True))
            pos = m.end()
        trozos.append((linea[pos:], actual, False))
    return trozos


def marcar(textos: list[str], faltan: dict[str, list[str]], quitar_etiquetas: bool = True) -> list[list[list[Any]]]:
    """Reparte las palabras que faltan entre sus apariciones en el texto original.

    Devuelve, por cada texto, sus partes [texto, clase o None]. Se marcan tantas
    apariciones como faltan, empezando por donde es más probable que estén:
    las eliminadas en campos que el esquema elimina, las descartadas en
    entorno, maniobra o automatizable, y las perdidas en texto sin etiqueta.
    """
    apariciones = []        # (n.º de texto, n.º de trozo, posición, palabra normalizada, etiqueta)
    trozos_por_texto = []
    for t, texto in enumerate(textos):
        trozos = segmentos(texto) if quitar_etiquetas else [(texto, None, False)]
        trozos_por_texto.append(trozos)
        for s, (trozo, etiqueta, es_etiqueta) in enumerate(trozos):
            if es_etiqueta:
                continue
            for m in re.finditer(r"\w+", _MARCAS.sub(lambda x: " " * len(x.group(0)), trozo)):
                apariciones.append((t, s, m.start(), convertir.sin_tildes(m.group(0)).lower(), etiqueta))
    prioridad = {
        "perdida": lambda a: a[4] is not None,
        "descartada": lambda a: a[4] not in ETIQUETAS_OPCIONES,
        "eliminada": lambda a: a[4] not in ETIQUETAS_ELIMINADAS,
    }
    clase_de: dict[tuple[int, int, int], str] = {}
    for clase in ("perdida", "descartada", "eliminada"):
        quedan = Counter(faltan.get(clase) or [])
        for a in sorted(apariciones, key=prioridad[clase]):
            if quedan[a[3]] > 0 and (a[0], a[1], a[2]) not in clase_de:
                clase_de[(a[0], a[1], a[2])] = clase
                quedan[a[3]] -= 1
    resultado = []
    for t, trozos in enumerate(trozos_por_texto):
        partes: list[list[Any]] = []
        for s, (trozo, _etiqueta, _es) in enumerate(trozos):
            pos = 0
            limpio = _MARCAS.sub(lambda x: " " * len(x.group(0)), trozo)
            for m in re.finditer(r"\w+", limpio):
                clase = clase_de.get((t, s, m.start()))
                if clase:
                    partes.append([trozo[pos:m.start()], None])
                    partes.append([trozo[m.start():m.end()], clase])
                    pos = m.end()
            partes.append([trozo[pos:], None])
        resultado.append([p for p in partes if p[0]])
    return resultado


def etiqueta_de_celda(texto: str) -> str:
    """Primera etiqueta conocida de la celda (estructura, sin contenido) para el resumen."""
    for linea in texto.splitlines():
        n = convertir.norm(linea).rstrip(":")
        if n in _NORM_ETIQUETAS:
            return _NORM_ETIQUETAS[n]
        m = _PATRON_ETIQUETA.search(linea)
        if m:
            return m.group(1)
    return "sin etiqueta"


# ----------------------------------------------------------------------
# Tablas del Word
# ----------------------------------------------------------------------
def celdas(tabla: Table) -> list[list[dict[str, Any]]]:
    """Filas con sus celdas: {texto, span, unida}. Una celda unida en vertical
    sale en cada fila que ocupa, pero solo la primera lleva el texto."""
    filas, vistas = [], set()
    for fila in tabla.rows:
        actual = []
        for c in convertir.celdas_unicas(fila):
            unida = c._tc in vistas
            vistas.add(c._tc)
            actual.append({"texto": "" if unida else convertir.texto_celda(c), "span": c._tc.grid_span, "unida": unida})
        filas.append(actual)
    return filas


def imagenes_tabla(tabla: Table) -> set[str]:
    rutas: set[str] = set()
    for fila in tabla.rows:
        for c in convertir.celdas_unicas(fila):
            for p in c.paragraphs:
                rutas.update(convertir.imagenes_parrafo(p, lambda parte: str(parte.partname)))
    return rutas


def tablas_anidadas(tabla: Table) -> int:
    return sum(len(c.tables) for fila in tabla.rows for c in convertir.celdas_unicas(fila))


# ----------------------------------------------------------------------
# YAML
# ----------------------------------------------------------------------
def leer_yaml(carpeta: Path) -> dict[str, Any]:
    capitulos = read_chapters(carpeta)
    pruebas: list[dict[str, Any]] = []
    tablas: list[dict[str, Any]] = []
    texto: Counter = Counter()
    imagenes: set[str] = set()

    def contexto(items: Any, capitulo: str) -> None:
        nonlocal texto
        for item in items if isinstance(items, list) else []:
            if isinstance(item, dict) and isinstance(item.get("tabla"), dict):
                tablas.append({"tabla": item["tabla"], "capitulo": capitulo})
            elif isinstance(item, dict):
                texto += palabras(item.get("texto"))
            else:
                texto += palabras(item)

    for c in capitulos:
        nombre = str(c["data"].get("capitulo") or "")
        texto += palabras(nombre)
        imagenes.update(_rutas(c["data"]))
        for e in c["data"]["contenido"]:
            if not isinstance(e, dict):
                continue
            if isinstance(e.get("prueba"), dict):
                p = e["prueba"]
                pruebas.append({"prueba": p, "capitulo": nombre})
                texto += palabras(p.get("nombre"))
                contexto(p.get("contexto"), nombre)
            elif isinstance(e.get("seccion"), str):
                texto += palabras(e["seccion"])
                contexto(e.get("contexto"), nombre)
            elif isinstance(e.get("tabla"), dict):
                tablas.append({"tabla": e["tabla"], "capitulo": nombre})
            elif "texto" in e:
                contexto([e["texto"]], nombre)
    return {"pruebas": pruebas, "tablas": tablas, "texto": texto, "imagenes": imagenes,
            "capitulos": len(capitulos)}


def _rutas(valor: Any) -> set[str]:
    rutas: set[str] = set()
    if isinstance(valor, dict):
        for clave, hijo in valor.items():
            if clave in ("imagenes", "imagen"):
                rutas.update(str(r) for r in (hijo if isinstance(hijo, list) else [hijo]) if r)
            else:
                rutas |= _rutas(hijo)
    elif isinstance(valor, list):
        for hijo in valor:
            rutas |= _rutas(hijo)
    return rutas


def _lista(c: Counter) -> list[str]:
    return sorted(c.elements())


# ----------------------------------------------------------------------
# Auditoría
# ----------------------------------------------------------------------
def auditar_prueba(tabla: Table, final: dict[str, Any] | None) -> dict[str, Any]:
    filas = celdas(tabla)
    original: Counter = Counter()
    for fila in filas[1:]:
        for celda in fila:
            original += palabras(sin_etiquetas(celda["texto"]))
    for celda in filas[0]:                              # cabecera: ID y título
        original += palabras(celda["texto"])
    leida = convertir.parsear(tabla, [], migrar=False)
    perdidas = original - palabras_de(leida)
    resultado: dict[str, Any] = {"perdidas": _lista(perdidas), "eliminadas": [], "descartadas": [], "avisos": []}
    if final is None:
        resultado["avisos"].append("No está en los YAML (¿se borró o cambió el ID después de importar?).")
        return resultado
    # Lo que el esquema convierte (A -> Automático, S -> Simplex, Sí -> si) no cuenta como eliminado.
    reconocidas: Counter = Counter()
    descartadas: Counter = Counter()
    for campo, normalizar in (("entorno", esquema.normalizar_entorno), ("maniobra", esquema.normalizar_maniobra)):
        if leida.get(campo):
            _validos, fuera = normalizar(leida[campo])
            descartadas += palabras(" ".join(fuera))
            reconocidas += palabras_de(leida[campo]) - palabras(" ".join(fuera))
    if leida.get("automatizable"):
        if esquema.normalizar_automatizable(leida["automatizable"]) is None:
            descartadas += palabras(leida["automatizable"])
        else:
            reconocidas += palabras(leida["automatizable"])
    quitadas = (original - perdidas) - palabras_de(final) - reconocidas
    resultado["descartadas"] = _lista(quitadas & descartadas)
    resultado["eliminadas"] = _lista(quitadas - descartadas)
    faltan_imagenes = len(imagenes_tabla(tabla)) - len(_rutas(final))
    if faltan_imagenes > 0:
        resultado["avisos"].append(f"Faltan {faltan_imagenes} imágenes de la tabla.")
    return resultado


def auditar_tabla(tabla: Table, final: dict[str, Any] | None) -> dict[str, Any]:
    original: Counter = Counter()
    for fila in celdas(tabla):
        for celda in fila:
            original += palabras(celda["texto"])
    resultado: dict[str, Any] = {"perdidas": [], "eliminadas": [], "descartadas": [], "avisos": []}
    if final is None:
        if original:
            resultado["avisos"].append("No se encuentra en los YAML.")
            resultado["perdidas"] = _lista(original)
        return resultado
    resultado["perdidas"] = _lista(original - palabras_de(final.get("filas")))
    faltan_imagenes = len(imagenes_tabla(tabla)) - len(_rutas(final))
    if faltan_imagenes > 0:
        resultado["avisos"].append(f"Faltan {faltan_imagenes} imágenes de la tabla.")
    return resultado


def estado(r: dict[str, Any]) -> str:
    if r["perdidas"] or r["avisos"]:
        return "perdida"
    if r["eliminadas"] or r["descartadas"]:
        return "eliminada"
    return "ok"


def auditar(ruta_docx: str | Path, carpeta_yaml: str | Path) -> dict[str, Any]:
    doc = Document(str(ruta_docx))
    yaml = leer_yaml(Path(carpeta_yaml))

    # Pruebas del YAML por ID, en orden (con IDs repetidos, la n-ésima con la n-ésima).
    por_id: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for p in yaml["pruebas"]:
        por_id[str(p["prueba"].get("id") or "").strip()].append(p)
    tablas_yaml = list(yaml["tablas"])
    usadas: set[int] = set()

    familias: dict[tuple, dict[str, Any]] = {}
    tablas, parrafos = [], []
    capitulo, numero = "", 0
    pool = Counter(yaml["texto"])
    imagenes_doc: set[str] = set()
    for el in doc.element.body.iterchildren():
        if el.tag.endswith("}p"):
            p = Paragraph(el, doc)
            imagenes_doc.update(convertir.imagenes_parrafo(p, lambda parte: str(parte.partname)))
            if convertir.nivel_titulo(p) == 1:
                capitulo = convertir.texto(p).strip()
            toks = palabras(convertir.texto(p))
            if not toks:
                continue
            faltan = toks - pool
            pool -= toks
            if faltan:
                parrafos.append({"capitulo": capitulo, "texto": convertir.texto(p), "perdidas": _lista(faltan),
                                 "partes": marcar([convertir.texto(p)], {"perdida": _lista(faltan)}, quitar_etiquetas=False)[0],
                                 "titulo": bool(convertir.nivel_titulo(p))})
            continue
        if not el.tag.endswith("}tbl"):
            continue
        numero += 1
        tabla = Table(el, doc)
        imagenes_doc |= imagenes_tabla(tabla)
        fam = inventario.familia(tabla)
        clave = fam["clave"]
        if clave not in familias:
            familias[clave] = {"id": f"F{len(familias) + 1}", "nombre": inventario.nombre_familia(clave),
                               "clase": fam["clase"], "faltan": list(fam["faltan"]), "sobran": list(fam["sobran"]),
                               "motivo": fam["motivo"], "tablas": []}
        familia = familias[clave]
        if fam["clase"] == "prueba":
            ident = convertir.texto_celda(convertir.celdas_unicas(tabla.rows[0])[0]).strip() if tabla.rows else ""
            cola = por_id.get(ident) or []
            encontrada = cola.pop(0) if cola else None
            final = encontrada["prueba"] if encontrada else None
            r = auditar_prueba(tabla, final)
            nueva = final
        else:
            original = Counter()
            for fila in celdas(tabla):
                for celda in fila:
                    original += palabras(celda["texto"])
            mejor, puntos = None, -1.0
            for i, t in enumerate(tablas_yaml):
                if i in usadas:
                    continue
                comun = sum((original & palabras_de(t["tabla"].get("filas"))).values())
                total = max(sum(original.values()), 1)
                nota = comun / total if original else (1.0 if not palabras_de(t["tabla"].get("filas")) else 0.0)
                if nota > puntos:
                    mejor, puntos = i, nota
            final = None
            if mejor is not None and (puntos >= 0.5 or not original):
                usadas.add(mejor)
                final = tablas_yaml[mejor]["tabla"]
            r = auditar_tabla(tabla, final)
            ident = fam["ident"]
            nueva = final
        anidadas = tablas_anidadas(tabla)
        if anidadas:
            r["avisos"].append(f"Tiene {anidadas} tablas dentro de celdas: su texto no se importa.")
        # Dónde estaban las palabras perdidas (fila, columna y etiqueta de la celda: estructura, sin contenido).
        lugares = []
        if r["perdidas"]:
            perdidas = set(r["perdidas"])
            for f, fila in enumerate(celdas(tabla)):
                for c, celda in enumerate(fila):
                    if perdidas & set(palabras(celda["texto"])):
                        lugares.append([f + 1, c + 1, etiqueta_de_celda(celda["texto"]) if fam["clase"] == "prueba" else "tabla"])
        original = celdas(tabla)
        textos = [c["texto"] for fila in original for c in fila]
        partes = iter(marcar(textos, {"perdida": r["perdidas"], "eliminada": r["eliminadas"], "descartada": r["descartadas"]},
                             quitar_etiquetas=fam["clase"] == "prueba"))
        for fila in original:
            for c in fila:
                c["partes"] = next(partes)
        entrada = {
            "n": numero, "tipo": "prueba" if fam["clase"] == "prueba" else "tabla", "clase": fam["clase"],
            "familia": familia["id"], "id": ident, "capitulo": capitulo, "estado": estado(r), **r,
            "lugares": lugares, "original": original, "nueva": nueva,
        }
        familia["tablas"].append(numero)
        tablas.append(entrada)

    # Muestra: la primera, la del medio y la última de cada familia.
    for f in familias.values():
        n = f["tablas"]
        muestra = {n[0], n[len(n) // 2], n[-1]}
        for t in tablas:
            if t["n"] in muestra:
                t["muestra"] = True

    ids = Counter(str(p["prueba"].get("id") or "").strip() for p in yaml["pruebas"])
    sin_tabla = [p["prueba"].get("id") for cola in por_id.values() for p in cola]
    pruebas_doc = sum(1 for t in tablas if t["tipo"] == "prueba")
    return {
        "documento": Path(ruta_docx).name,
        "recuentos": {
            "pruebas_word": pruebas_doc,
            "pruebas_yaml": len(yaml["pruebas"]),
            "tablas_word": sum(1 for t in tablas if t["tipo"] == "tabla"),
            "tablas_yaml": len(yaml["tablas"]),
            "imagenes_word": len(imagenes_doc),
            "imagenes_yaml": len(yaml["imagenes"]),
            "ids_repetidos": sorted(i for i, n in ids.items() if n > 1 and i),
            "pruebas_sin_id": ids.get("", 0),
            "yaml_sin_tabla": [i for i in sin_tabla if i],
            "capitulos_yaml": yaml["capitulos"],
        },
        "familias": [{k: v for k, v in f.items() if k != "tablas"} | {
            "n": len(f["tablas"]),
            **{e: sum(1 for t in tablas if t["familia"] == f["id"] and t["estado"] == e) for e in ESTADOS},
        } for f in familias.values()],
        "tablas": tablas,
        "parrafos": parrafos,
    }


# ----------------------------------------------------------------------
# Resumen sin contenido (para pantalla y auditoria.txt)
# ----------------------------------------------------------------------
def resumen(datos: dict[str, Any]) -> str:
    r = datos["recuentos"]
    t = datos["tablas"]
    lineas = [
        f"Auditoría de {datos['documento']}",
        "",
        f"Pruebas:      {r['pruebas_word']} tablas en el Word, {r['pruebas_yaml']} en los YAML",
        f"No pruebas:   {r['tablas_word']} tablas en el Word, {r['tablas_yaml']} en los YAML",
        f"Imágenes:     {r['imagenes_word']} en el Word, {r['imagenes_yaml']} referenciadas en los YAML",
        f"IDs repetidos: {len(r['ids_repetidos'])} · pruebas sin ID: {r['pruebas_sin_id']}"
        f" · pruebas del YAML sin tabla en el Word: {len(r['yaml_sin_tabla'])}",
        "",
        "Tablas:  " + " · ".join(f"{ESTADOS[e]}: {sum(1 for x in t if x['estado'] == e)}" for e in ESTADOS),
        f"Párrafos fuera de tablas con palabras perdidas: {len(datos['parrafos'])}",
        "",
        "Por familia (✖ / ⚠ / ✔):",
    ]
    for f in datos["familias"]:
        lineas.append(f"  {f['id']:>4}  {f['nombre']:<22} {f['n']:5d} tablas   ✖ {f['perdida']:4d}  ⚠ {f['eliminada']:4d}  ✔ {f['ok']:4d}")
        if f["sobran"]:
            lineas.append(f"        sobran: {', '.join(s[:40] for s in f['sobran'])}")
    lugares: Counter = Counter()
    for x in t:
        for fila, col, etiqueta in x["lugares"]:
            lugares[(x["familia"], fila, col, etiqueta)] += 1
    if lugares:
        lineas += ["", "Dónde están las palabras perdidas (familia, fila, columna, etiqueta de la celda: nº de tablas):"]
        for (fam, fila, col, etiqueta), n in lugares.most_common(40):
            lineas.append(f"  {fam:>4}  fila {fila:2d} col {col}  {etiqueta[:40]:<40} {n:5d}")
    avisos = Counter(a for x in t for a in x["avisos"])
    if avisos:
        lineas += ["", "Otros avisos:"] + [f"  {n:5d} × {a}" for a, n in avisos.most_common()]
    return "\n".join(lineas)


# ----------------------------------------------------------------------
# Informe HTML
# ----------------------------------------------------------------------
def informe_html(datos: dict[str, Any]) -> str:
    plantilla = (Path(__file__).resolve().parent.parent / "plantillas" / "auditoria.html").read_text(encoding="utf-8")
    datos_json = json.dumps({**datos, "etiquetas": esquema.ETIQUETAS, "automatizable": esquema.AUTOMATIZABLE_VISIBLE},
                            ensure_ascii=False, default=str).replace("</", "<\\/")
    return (plantilla.replace("{{titulo}}", html.escape(f"Auditoría · {datos['documento']}"))
            .replace("{{datos}}", datos_json))


def main(argv=None, prog=None) -> None:
    ap = argparse.ArgumentParser(prog=prog, description="Comprueba que la importación del Word a YAML no pierde información")
    ap.add_argument("docx", help="documento Word original")
    ap.add_argument("yaml", help="carpeta con los YAML importados (con indice.yaml)")
    ap.add_argument("-o", "--salida", default="auditoria.html", help="informe HTML (por defecto: auditoria.html)")
    args = ap.parse_args(argv)
    try:
        datos = auditar(args.docx, args.yaml)
    except (ValueError, OSError) as error:
        raise SystemExit(f"Error: {error}")
    salida = Path(args.salida)
    salida.write_text(informe_html(datos), encoding="utf-8")
    texto = resumen(datos)
    salida.with_suffix(".txt").write_text(texto + "\n", encoding="utf-8")
    print(texto)
    print(f"\nInforme: {salida}  (lleva el texto del documento: no lo compartas)")
    print(f"Resumen sin contenido: {salida.with_suffix('.txt')}")


if __name__ == "__main__":
    main()
