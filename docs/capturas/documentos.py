"""Documentos ficticios para las capturas de importar (Word de requisitos, Word de resultados y Excel).

Se crean en la carpeta de los datos de las capturas, a partir de los ejemplos.

Uso: python documentos.py <repo> <destino>
"""
import sys
from pathlib import Path

repo, destino = Path(sys.argv[1]), Path(sys.argv[2])
sys.path.insert(0, str(repo))
from docx import Document  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from app import yaml_a_docx  # noqa: E402


def requisito(doc, ident, titulo, celdas, tabla=None, despues=None):
    t = doc.add_table(rows=1 + len(celdas), cols=2)
    t.style = "Table Grid"
    t.cell(0, 0).text, t.cell(0, 1).text = ident, titulo
    for i, texto in enumerate(celdas, start=1):
        celda = t.cell(i, 0).merge(t.cell(i, 1))
        lineas = texto.split("\n")
        celda.paragraphs[0].text = lineas[0]
        for linea in lineas[1:]:
            celda.add_paragraph(linea)
        if tabla and texto.startswith("Descripción"):
            anidada = celda.add_table(rows=len(tabla), cols=len(tabla[0]))
            for f, fila in enumerate(tabla):
                for c, valor in enumerate(fila):
                    anidada.cell(f, c).text = valor
            if despues:
                celda.add_paragraph(despues)


# 1. Requisitos de la botonera (proyecto nuevo).
doc = Document()
doc.add_heading("Requisitos de la botonera", 1)
doc.add_heading("Pulsadores", 2)
requisito(doc, "BOT-001", "Los pulsadores se iluminan al registrar la llamada",
          ["Descripción: Cada pulsador tiene un anillo luminoso.", "Origen: Pliego técnico", "Tipo: Obligatorio",
           "Método de verificación: Test", "Prioridad: Alta"],
          tabla=[["Estado", "Anillo"], ["Llamada registrada", "Encendido"], ["Llamada atendida", "Apagado"]],
          despues="El brillo se ajusta con el parámetro P-12.")
requisito(doc, "BOT-002", "El pulsador de alarma suena 5 s como mínimo",
          ["Descripción: Al pulsar la alarma suena el zumbador de cabina.\nNota: también se avisa al centro de rescate.",
           "Origen: Normativa", "Tipo: Obligatorio", "Método de verificación: Test", "Prioridad: Alta"])
doc.add_heading("Indicadores", 2)
requisito(doc, "BOT-003", "El indicador muestra la planta en la que está la cabina",
          ["Descripción: El indicador de cabina y los de piso muestran la planta.", "Origen: Pliego técnico",
           "Tipo: Obligatorio", "Método de verificación: Inspección", "Prioridad: Media"])
requisito(doc, "BOT-003", "El indicador muestra la flecha de sentido",
          ["Descripción: Una flecha indica si la cabina sube o baja.", "Origen: Pliego técnico",
           "Tipo: Obligatorio", "Método de verificación: Inspección"])
doc.add_paragraph("Versión")
t = doc.add_table(rows=2, cols=2)
t.style = "Table Grid"
t.cell(0, 0).text, t.cell(0, 1).text = "Versión", "Fecha"
t.cell(1, 0).text, t.cell(1, 1).text = "1.0", "02/03/2026"
doc.save(destino / "requisitos_botonera.docx")

# 2. Requisitos del proyecto en Word, con un cambio (importar al proyecto abierto).
yaml_a_docx.generar_docx(destino / "yaml" / "requisitos", destino / "requisitos_revisados.docx")
doc = Document(destino / "requisitos_revisados.docx")
for tabla in doc.tables:
    if tabla.cell(0, 0).text.strip() == "REQ-LLAM-002":
        for fila in tabla.rows:
            for celda in fila.cells:
                for p in celda.paragraphs:
                    if p.text.startswith("Descripción"):
                        for run in p.runs[1:]:
                            run.text = ""
                        p.add_run(" La luz del pulsador de piso se enciende al registrar la llamada y se apaga al llegar la cabina.")
                        break
doc.save(destino / "requisitos_revisados.docx")

# 3. Resultados de una ejecución hechos en Word.
yaml_a_docx.generar_docx(destino / "yaml" / "pruebas", destino / "resultados_torre.docx")
doc = Document(destino / "resultados_torre.docx")
resultados = {
    "DEMO-LLAM-001": ("OK", ""), "DEMO-LLAM-002": ("Correcto", ""), "DEMO-LLAM-003": ("NOK", "Atiende primero P1."),
    "DEMO-NIV-001": ("Pasa", ""), "DEMO-CAN-001": ("OK", ""), "DEMO-CAN-002": ("Correcto", ""),
    "DEMO-DIAG-001": ("Bloqueado", "Sin acceso al registro."), "DEMO-PUER-001": ("OK", ""),
    "DEMO-PUER-002": ("NOK", "No reabre con la barrera al 50 %."), "DEMO-SEG-001": ("OK", ""),
    "DEMO-MANT-001": ("N/A", ""), "DEMO-MANT-002": ("OK", ""), "DEMO-VAR-001": ("Parcial", "Solo la mitad de los ciclos."),
    "DEMO-VAR-002": ("", ""),
}
for tabla in doc.tables:
    r = resultados.get(tabla.cell(0, 0).text.strip())
    if not r:
        continue
    for texto in (f"Resultado obtenido: {r[0]}", f"Observaciones: {r[1]}", "Ejecutor prueba: Luis Gómez", "Fecha ejecución: 11/09/2026"):
        fila = tabla.add_row()
        fila.cells[0].merge(fila.cells[-1]).text = texto
doc.save(destino / "resultados_torre.docx")

# 4. Excel de pruebas: hoja «Llamadas» (una fila por paso) y hoja «Notas».
libro = Workbook()
hoja = libro.active
hoja.title = "Llamadas"
hoja.append(["ID", "Título", "Capítulo", "Sección", "Descripción", "Pasos", "Requisitos", "Entorno", "Notas"])
hoja.append(["DEMO-LLAM-001", "Una llamada de cabina lleva el ascensor a la planta pedida", "Ejemplo básico", "Llamadas",
             "Comprobar que una llamada de cabina se atiende.", "Pulsar el botón 2 de la botonera de cabina.", "REQ-LLAM-001", "Automático\nMaqueta", ""])
hoja.append(["", "", "", "", "", "La cabina se detiene en P2 y abre puertas.", "", "", ""])
hoja.append(["DEMO-LLAM-010", "Una llamada de cabina se cancela al pulsar dos veces", "Ejemplo básico", "Llamadas",
             "Comprobar que la segunda pulsación cancela la llamada.", "Pulsar dos veces el botón 3.", "REQ-LLAM-004", "Maqueta", "Revisar con el cliente"])
hoja.append(["DEMO-ALAR-001", "El pulsador de alarma avisa al centro de rescate", "Ejemplo básico", "Alarmas",
             "Comprobar el aviso de alarma.", "Mantener pulsada la alarma 5 s.", "", "Torre", ""])
libro.create_sheet("Notas").append(["Pruebas nuevas de la versión 2.4"])
libro.save(destino / "pruebas_llamadas.xlsx")

# El mismo Excel que importa otra persona después (para el mapeo del proyecto).
hoja.cell(row=4, column=1).value = "DEMO-LLAM-011"
hoja.cell(row=5, column=1).value = "DEMO-ALAR-002"
libro.save(destino / "pruebas_llamadas_luis.xlsx")
