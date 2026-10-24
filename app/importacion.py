#!/usr/bin/env python3
"""Importar pruebas o requisitos de un Word, con un mapeo de sus etiquetas a nuestro esquema.

Un Excel (una fila por elemento) se lee en app/importacion_excel.py y sigue
el mismo camino a partir de la conversión: revisión e importación.

Lo usa el asistente de importación del editor (⋯ → Importar…) y la orden
`python herramientas.py importar`. Tres pasos:

1. `analizar`: agrupa las tablas del Word en familias (tablas con las mismas
   etiquetas: «Descripción:», «Pasos:»…), propone qué es cada familia (prueba,
   requisito, tabla informativa o nada) y a qué campo va cada etiqueta.
2. `convertir`: lee el Word con ese mapeo. Devuelve los capítulos (con sus
   secciones, textos y tablas), los avisos y los IDs que habría que cambiar.
3. `importar`: escribe los YAML en un proyecto nuevo o los añade al abierto.

Se espera la forma de las plantillas (`plantillas/tabla_*.docx`): la primera
fila es «ID | Título» y cada celda de las demás empieza por «Etiqueta: valor».
Un párrafo «Etiqueta:» solo en su línea abre un campo nuevo dentro de la misma
celda (por ejemplo «Situación inicial:» y «Ensayo:»). Las tablas dentro de
una celda se conservan: en un requisito, en `detalle` (tablas e imágenes de la
descripción); en una prueba, en su contexto.

El mapeo es un dict:

    {"tipos": {firma_de_familia: "prueba" | "requisito" | "tabla" | "ignorar"},
     "campos": {"prueba": {etiqueta_normalizada: campo}, "requisito": {...}}}

donde campo es un campo del esquema, "adicional" (campo adicional con su
etiqueta), "anterior" (el párrafo sigue al campo anterior) o "ignorar".
"""
from __future__ import annotations

import argparse
import difflib
import json
import mimetypes
import re
from collections import Counter
from pathlib import Path
from typing import Any, Callable

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from ruamel.yaml import YAML

from app import adjuntos, esquema, formato
from migracion import convertir as cv
from migracion.ids_duplicados import _ID_VALIDO, _siguiente

TIPOS = ("prueba", "requisito", "tabla", "ignorar")
TIPOS_VISIBLES = {"prueba": "Pruebas", "requisito": "Requisitos", "tabla": "Tabla informativa", "ignorar": "No importar"}
ESPECIALES = {
    "adicional": "Campo adicional (con su etiqueta)",
    "anterior": "Parte del campo anterior",
    "ignorar": "No importar",
}

# Clase de cada campo de destino: cómo se lee su valor.
CLASE_PRUEBA = {
    "descripcion": "parrafos", "parametrizacion": "lista", "situacion_inicial": "lista", "pasos": "lista",
    "maniobra": "opciones", "automatizable": "linea", "entorno": "opciones", "requisitos": "textos",
    "etiquetas": "textos", "fecha": "linea", "version": "linea", "contexto": "lista",
}
CLASE_REQUISITO = {
    "descripcion": "parrafos", "fuente": "linea", "referencia": "linea", "tipo": "linea",
    "criterio_validacion": "opciones", "fase": "linea", "fecha": "linea", "version": "linea", "contexto": "lista",
}
CLASES = {"prueba": CLASE_PRUEBA, "requisito": CLASE_REQUISITO}
NOMBRES = {"prueba": esquema.ETIQUETAS, "requisito": esquema.REQ_ETIQUETAS}

# Etiquetas que se proponen para cada campo (además del nombre del campo en el esquema).
SINONIMOS = {
    "prueba": {
        "descripcion": ["Descripción", "Objetivo", "Objeto", "Propósito"],
        "parametrizacion": ["Parametrización", "Parámetros", "Configuración"],
        "situacion_inicial": ["Situación inicial", "Condiciones iniciales", "Precondiciones", "Precondición", "Estado inicial"],
        "pasos": ["Ensayo", "Pasos", "ENSAYO(s) / Proceso(s)", "Procedimiento", "Proceso", "Situación final"],
        "maniobra": ["Maniobra", "Maniobra(s)"],
        "automatizable": ["Automatizable"],
        "entorno": ["Entorno de pruebas", "Entorno de prueba(s)", "Entorno de pruebas(s)", "Entorno"],
        "requisitos": ["Requisitos", "Requisito(s)", "Requisitos cubiertos", "Trazabilidad"],
        "etiquetas": ["Tags", "Etiquetas"],
        "fecha": ["Fecha"],
        "version": ["Versión", "Revisión", "Edición"],
    },
    "requisito": {
        "descripcion": ["Descripción", "Texto", "Enunciado"],
        "fuente": ["Fuente", "Origen"],
        "referencia": ["Referencia", "Referencias", "Documento de referencia"],
        "tipo": ["Tipo", "Tipo de requisito", "Clase"],
        "criterio_validacion": ["Criterio validación", "Criterio de validación", "Método de validación",
                                "Método de verificación", "Verificación"],
        "fase": ["Fase"],
        "fecha": ["Fecha"],
        "version": ["Versión", "Revisión", "Edición"],
    },
}
# Etiquetas del formato antiguo de las pruebas que no pasan al esquema (como en convertir.py), y la
# fila «Adjuntos» del Word que genera el editor (sus objetos se recogen igualmente: ver _Lector.objetos).
IGNORADAS_PRUEBA = ["Componente", "Plataforma", "Criticidad", "Referencia", "Ejecución de prueba",
                    *esquema.ADICIONALES_ELIMINADOS, esquema.ETIQUETAS["adjuntos"]]
# Campos que comparten pruebas y requisitos: no sirven para distinguirlos.
COMUNES = {"descripcion", "fecha", "version"}
# Clave (solo durante la importación, no se escribe) con los campos que trae un elemento de un Excel:
# al actualizar uno que ya existe, solo cambian esos y los demás se conservan.
COLUMNAS = "_columnas"


def _sin_columnas(elemento: dict[str, Any]) -> dict[str, Any]:
    return {k: v for k, v in elemento.items() if k != COLUMNAS}


def clave(etiqueta: str) -> str:
    """Etiqueta normalizada: sin tildes, mayúsculas, espacios ni signos («Requisito(s):» → «requisito»)."""
    return re.sub(r"[^a-z0-9]", "", cv.norm(str(etiqueta)).replace("(s)", ""))


_SINONIMOS = {tipo: {clave(s): campo for campo, lista in campos.items() for s in [*lista, NOMBRES[tipo].get(campo, campo)]}
              for tipo, campos in SINONIMOS.items()}
_IGNORADAS = {clave(e) for e in IGNORADAS_PRUEBA} - set(_SINONIMOS["prueba"])


def proponer_campo(tipo: str, etiqueta: str, solo_interior: bool = False) -> str:
    """Campo propuesto para una etiqueta: el de su sinónimo, uno muy parecido o, si no, campo adicional."""
    if tipo not in CLASES:
        return "ignorar"
    k = clave(etiqueta)
    if k in _SINONIMOS[tipo]:
        return _SINONIMOS[tipo][k]
    # Una «Nota:» en mitad de una celda suele ser texto del campo, no un campo nuevo.
    if solo_interior:
        return "anterior"
    if tipo == "prueba" and k in _IGNORADAS:
        return "ignorar"
    parecida = difflib.get_close_matches(k, list(_SINONIMOS[tipo]), n=1, cutoff=0.85)
    return _SINONIMOS[tipo][parecida[0]] if parecida else "adicional"


# ----------------------------------------------------------------------
# Etiquetas de una tabla
# ----------------------------------------------------------------------
_ETIQUETA = re.compile(r"^\s*([^:\n]{1,40}?)\s*:\s*", re.S)


def etiqueta_de(texto: str) -> tuple[str, int] | None:
    """(«Etiqueta», posición donde empieza el valor) si el párrafo empieza por «Etiqueta:»."""
    m = _ETIQUETA.match(texto)
    if not m:
        return None
    etiqueta = " ".join(m.group(1).split())
    if not etiqueta or len(etiqueta.split()) > 6 or re.fullmatch(r"[\d\s.,/\-]+", etiqueta) \
            or etiqueta.lower().startswith(("http", "www")) or not clave(etiqueta):
        return None
    return etiqueta, m.end()


def bloques(celda) -> list[Paragraph | Table]:
    """Párrafos y tablas de una celda, en orden."""
    return list(celda.iter_inner_content())


def estructura(tabla: Table) -> dict[str, Any]:
    """Cabecera, etiquetas (con si abren la celda o van en mitad) y tablas anidadas de una tabla."""
    filas = cv.filas_sin_repetir(tabla)
    cabecera = [" ".join(cv.texto_celda(c).split()) for c in filas[0]] if filas else []
    etiquetas: dict[str, dict[str, Any]] = {}
    anidadas = 0
    for fila in filas[1:]:
        for celda in fila:
            primero = True
            for b in bloques(celda):
                if isinstance(b, Table):
                    anidadas += 1
                    primero = False
                    continue
                texto = cv.texto(b)
                if not texto.strip():
                    continue
                e = etiqueta_de(texto)
                if e:
                    info = etiquetas.setdefault(clave(e[0]), {"etiqueta": e[0], "inicio": False})
                    info["inicio"] = info["inicio"] or primero
                primero = False
    return {"cabecera": cabecera, "etiquetas": etiquetas, "anidadas": anidadas, "filas": len(filas),
            "columnas": max((len(f) for f in filas), default=0)}


def _proponer_tipo(etiquetas: dict[str, dict[str, Any]], preferido: str | None) -> str:
    if not etiquetas:
        return "tabla"
    campos = {t: {proponer_campo(t, i["etiqueta"]) for i in etiquetas.values()} - set(ESPECIALES) for t in CLASES}
    propios = {t: campos[t] - COMUNES for t in CLASES}
    if preferido in CLASES and campos[preferido]:
        return preferido
    if {"pasos", "situacion_inicial"} & propios["prueba"]:
        return "prueba"
    if len(propios["requisito"]) > len(propios["prueba"]):
        return "requisito"
    if propios["prueba"] or len(campos["prueba"]) >= 2:
        return "prueba"
    return "tabla"


# ----------------------------------------------------------------------
# 1. Analizar
# ----------------------------------------------------------------------
def _familias(estructuras: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[int]]:
    """Agrupa las tablas: primero por etiquetas idénticas y después une las variantes
    (las que comparten al menos el 60 % de sus etiquetas con una familia más grande).
    Devuelve las familias y, para cada tabla, el número de su familia."""
    grupos: dict[str, list[int]] = {}
    for i, e in enumerate(estructuras):
        firma = "|".join(sorted(e["etiquetas"])) or f"sin etiquetas · {e['columnas']} columnas"
        grupos.setdefault(firma, []).append(i)
    familias: list[dict[str, Any]] = []
    de_tabla = [0] * len(estructuras)
    for firma, tablas in sorted(grupos.items(), key=lambda g: (-len(g[1]), g[0])):
        propias = set(estructuras[tablas[0]]["etiquetas"])
        destino = None
        if propias:
            for f in familias:
                if f["claves"] and len(propias & f["claves"]) >= 0.6 * len(propias):
                    destino = f
                    break
        if destino is None:
            destino = {"firma": firma, "claves": set(propias), "tablas": []}
            familias.append(destino)
        destino["claves"] |= propias
        destino["tablas"].extend(tablas)
    for n, f in enumerate(familias):
        for i in f["tablas"]:
            de_tabla[i] = n
    return familias, de_tabla


def analizar_documento(doc, preferido: str | None = None) -> dict[str, Any]:
    estructuras = [estructura(t) for t in doc.tables]
    familias, de_tabla = _familias(estructuras)
    resultado = []
    for n, f in enumerate(familias):
        tablas = sorted(f["tablas"])
        cuenta: Counter[str] = Counter()
        datos: dict[str, dict[str, Any]] = {}
        for i in tablas:
            for k, info in estructuras[i]["etiquetas"].items():
                cuenta[k] += 1
                d = datos.setdefault(k, {"etiqueta": info["etiqueta"], "inicio": False})
                d["inicio"] = d["inicio"] or info["inicio"]
        etiquetas = {k: datos[k] for k in sorted(datos, key=lambda k: -cuenta[k])}
        tipo = _proponer_tipo(etiquetas, preferido)
        resultado.append({
            "numero": n + 1,
            "firma": f["firma"],
            "tablas": len(tablas),
            "posiciones": [i + 1 for i in tablas[:200]],
            "ejemplos": [" · ".join(x for x in estructuras[i]["cabecera"][:2] if x)[:80] or "(sin cabecera)" for i in tablas[:3]],
            "anidadas": sum(estructuras[i]["anidadas"] for i in tablas),
            "tipo": tipo,
            "etiquetas": [{"clave": k, "etiqueta": d["etiqueta"], "tablas": cuenta[k], "inicio": d["inicio"],
                           "propuesta": {t: proponer_campo(t, d["etiqueta"], not d["inicio"]) for t in CLASES}}
                          for k, d in etiquetas.items()],
        })
    return {"tablas": len(estructuras), "familias": resultado, "familia_de_tabla": de_tabla}


def es_excel(ruta: str | Path) -> bool:
    return Path(ruta).suffix.lower() in (".xlsx", ".xlsm")


def analizar(ruta: str | Path, preferido: str | None = None) -> dict[str, Any]:
    """Familias de tablas del Word (o las hojas del Excel), con el tipo y los campos propuestos."""
    if es_excel(ruta):
        from app import importacion_excel
        return importacion_excel.analizar_libro(ruta, preferido)
    resultado = analizar_documento(Document(str(ruta)), preferido)
    resultado.pop("familia_de_tabla")
    resultado["archivo"] = str(ruta)
    return resultado


def mapeo_propuesto(analisis: dict[str, Any], recordado: dict[str, Any] | None = None) -> dict[str, Any]:
    """Mapeo inicial: lo propuesto por el análisis, con lo que se eligió la última vez encima."""
    recordado = recordado or {}
    tipos_rec, campos_rec = recordado.get("tipos") or {}, recordado.get("campos") or {}
    mapeo: dict[str, Any] = {"tipos": {}, "campos": {t: {} for t in CLASES}}
    for f in analisis["familias"]:
        mapeo["tipos"][f["firma"]] = tipos_rec.get(f["firma"], f["tipo"])
        for e in f["etiquetas"]:
            for t in CLASES:
                mapeo["campos"][t].setdefault(e["clave"], (campos_rec.get(t) or {}).get(e["clave"], e["propuesta"][t]))
    return mapeo


# ----------------------------------------------------------------------
# 2. Convertir
# ----------------------------------------------------------------------
def _item_lista(p: Paragraph, desde: int, guardar_imagen) -> Any:
    """Elemento de una lista (paso, situación…): texto con formato, nivel e imágenes."""
    texto = cv.limpiar(cv.texto_md(p, desde))
    ppr = p._p.pPr
    nivel = ppr.numPr.ilvl.val if ppr is not None and ppr.numPr is not None and ppr.numPr.ilvl is not None else 0
    imagenes = cv.imagenes_parrafo(p, guardar_imagen)
    if not nivel and not imagenes:
        return texto
    item: dict[str, Any] = {"texto": texto}
    if nivel:
        item["nivel"] = nivel
    if imagenes:
        item["imagenes"] = imagenes
    return item


def _texto_plano(valor: Any) -> str:
    """Todo el texto de un valor (listas, tablas, campos adicionales), sin marcas de formato."""
    if isinstance(valor, dict):
        return " ".join(_texto_plano(v) for v in valor.values())
    if isinstance(valor, list):
        return " ".join(_texto_plano(v) for v in valor)
    return re.sub(r"\[\[[^\]|]*\||\]\]|\*", "", str(valor or ""))


def _letras(texto: str) -> str:
    """Solo letras y números, en minúsculas: para comparar textos sin fijarse en espacios, viñetas o signos."""
    return re.sub(r"[\W_]+", "", str(texto).lower())


def _tiene_imagen(p: Paragraph) -> bool:
    return bool(p._p.xpath(".//a:blip | .//*[local-name()='imagedata']"))


def _vacio(item: Any) -> bool:
    return (isinstance(item, str) and not item.strip()) or (isinstance(item, dict) and not item.get("texto") and not item.get("imagenes") and "tabla" not in item)


class _Lector:
    """Lee las tablas de un tipo con el mapeo de sus etiquetas."""

    def __init__(self, tipo: str, campos: dict[str, str], guardar_imagen) -> None:
        self.tipo, self.campos, self.guardar = tipo, campos, guardar_imagen
        self.clases = CLASES[tipo]

    def destino(self, etiqueta: str, interior: bool) -> str:
        return self.campos.get(clave(etiqueta)) or proponer_campo(self.tipo, etiqueta, interior)

    def segmentos(self, filas) -> list[dict[str, Any]]:
        """Trozos de la tabla: cada uno con su etiqueta, su campo y sus párrafos y tablas."""
        segmentos: list[dict[str, Any]] = []
        actual = None
        for fila in filas[1:]:
            for celda in fila:
                primero = True
                for b in bloques(celda):
                    if isinstance(b, Table):
                        if primero or actual is None:
                            actual = {"etiqueta": "", "destino": "adicional", "partes": []}
                            segmentos.append(actual)
                        actual["partes"].append(("t", b, 0))
                        primero = False
                        continue
                    texto = cv.texto(b)
                    con_imagen = _tiene_imagen(b)
                    if not texto.strip() and not con_imagen:
                        continue
                    e = etiqueta_de(texto) if texto.strip() else None
                    if e:
                        destino = self.destino(e[0], not primero)
                        if destino != "anterior" or actual is None:
                            actual = {"etiqueta": e[0], "destino": "adicional" if destino == "anterior" else destino, "partes": []}
                            segmentos.append(actual)
                            if texto[e[1]:].strip() or con_imagen:
                                actual["partes"].append(("p", b, e[1]))
                            primero = False
                            continue
                        # «anterior»: el párrafo entero, con su etiqueta, sigue al campo anterior.
                    elif primero or actual is None:
                        # Celda que no empieza por una etiqueta: su texto se conserva aparte.
                        actual = {"etiqueta": "", "destino": "adicional", "partes": []}
                        segmentos.append(actual)
                    actual["partes"].append(("p", b, 0))
                    primero = False
        return segmentos

    def leer(self, tabla: Table, avisos: list) -> dict[str, Any]:
        filas = cv.filas_sin_repetir(tabla)
        cab = filas[0] if filas else []
        d: dict[str, Any] = {"id": " ".join(cv.texto_celda(cab[0]).split()) if cab else "",
                             "titulo": " ".join(" ".join(cv.texto_celda(c).split()) for c in cab[1:]).strip()}
        imagenes_cab = [i for c in cab for p in c.paragraphs for i in cv.imagenes_parrafo(p, self.guardar)]
        extras: dict[str, str] = {}
        detalle: list[Any] = []                      # tablas e imágenes (requisito: `detalle`; prueba: contexto)
        imagenes: list[str] = list(imagenes_cab)
        segmentos = self.segmentos(filas)
        for s in segmentos:
            destino = s["destino"]
            if destino == "ignorar":
                continue
            clase = self.clases.get(destino, "parrafos" if destino == "adicional" else None)
            if clase is None:
                avisos.append(cv.Aviso("campo_desconocido", f"«{s['etiqueta']}» va a un campo que no existe ({destino}): se conserva como campo adicional", d["id"]))
                destino, clase = "adicional", "parrafos"
            valor = self.valor(s, clase, destino, detalle, imagenes, avisos, d["id"])
            if destino == "adicional":
                nombre = s["etiqueta"] or "Sin etiqueta"
                if valor or (s["etiqueta"] and nombre not in extras):
                    extras[nombre] = f"{extras[nombre]}\n{valor}" if extras.get(nombre) else valor
                if not s["etiqueta"] and valor:
                    avisos.append(cv.Aviso("celda_sin_reconocer", f"celda sin etiqueta: se conserva como campo adicional «Sin etiqueta» ({' '.join(valor.split())[:50]})", d["id"]))
            elif clase == "lista":
                d.setdefault(destino, []).extend(valor)
            elif clase == "textos":
                d.setdefault(destino, []).extend(v for v in valor if v not in d[destino])
            elif clase == "parrafos":
                d[destino] = f"{d[destino]}\n{valor}" if d.get(destino) and valor else (valor or d.get(destino, ""))
            else:
                d[destino] = f"{d[destino]} {valor}".strip() if d.get(destino) else valor
        if extras:
            d["campos_adicionales"] = extras
        self.objetos(tabla, d, avisos)
        resultado = self.normalizar(d, detalle, imagenes, avisos)
        self.comprobar_texto(segmentos, resultado, avisos)
        return resultado

    def objetos(self, tabla: Table, d: dict[str, Any], avisos: list) -> None:
        """Los objetos incrustados de la tabla (un Excel…): en una prueba, cada uno pasa a ser un
        adjunto (ver adjuntos.py). Dónde va cada uno lo decide luego _colocar_adjuntos, por su dibujo;
        aquí se recogen también los de los campos que no se importan."""
        if not hasattr(self.guardar, "objeto"):
            return
        for o in adjuntos.objetos(tabla._tbl, tabla.part):
            adjunto = self.guardar.objeto(o, d["id"], avisos, self.tipo == "prueba")
            if adjunto is not None:
                d.setdefault("adjuntos", []).append(adjunto)

    def comprobar_texto(self, segmentos, item, avisos) -> None:
        """Avisa de los párrafos del Word que no han llegado a ningún campo (sin contar los que se
        mapean a «No importar» ni los de opciones, que se normalizan: «Auto» → «Automático»)."""
        plano = _letras(_texto_plano(item))
        for s in segmentos:
            if s["destino"] == "ignorar" or self.clases.get(s["destino"]) == "opciones" or s["destino"] in ("automatizable", "tipo"):
                continue
            for tipo, b, desde in s["partes"]:
                textos = [cv.texto_celda(c) for f in b.rows for c in cv.celdas_unicas(f)] if tipo == "t" else [cv.texto(b)[desde:]]
                for texto in textos:
                    if len(_letras(texto)) >= 3 and _letras(texto) not in plano:
                        avisos.append(cv.Aviso("texto_no_importado", f"«{' '.join(texto.split())[:80]}» ({s['etiqueta'] or 'sin etiqueta'})", item.get("id", "")))

    def valor(self, s, clase, destino, detalle, imagenes, avisos, ident) -> Any:
        partes = s["partes"]
        tras_tabla = False                           # en la descripción: lo que sigue a una tabla va con ella, en orden
        if clase == "lista":
            items = []
            for tipo, b, desde in partes:
                if tipo == "t":
                    self.tabla_anidada(b, s, detalle, avisos, ident)
                    continue
                item = _item_lista(b, desde, self.guardar)
                solo_imagen = isinstance(item, dict) and not item.get("texto") and item.get("imagenes") and not item.get("nivel")
                if solo_imagen and items and not (isinstance(items[-1], dict) and "tabla" in items[-1]):
                    # Una imagen en su propio párrafo, justo después de un paso, es de ese paso (como en el editor).
                    anterior = items[-1] if isinstance(items[-1], dict) else {"texto": items[-1]}
                    anterior["imagenes"] = [*(anterior.get("imagenes") or []), *item["imagenes"]]
                    items[-1] = anterior
                elif not _vacio(item):
                    items.append(item)
            return items
        lineas = []
        for tipo, b, desde in partes:
            if tipo == "t":
                self.tabla_anidada(b, s, detalle, avisos, ident)
                tras_tabla = self.tipo == "requisito" and destino == "descripcion"
                continue
            if clase == "parrafos":
                texto = cv.texto_md(b, desde).strip()
            else:
                texto = cv.texto(b)[desde:].strip()
            fotos = cv.imagenes_parrafo(b, self.guardar)
            if tras_tabla:
                if texto or fotos:
                    detalle.append({"texto": texto, "imagenes": fotos} if fotos else texto)
                continue
            if fotos:
                if self.tipo == "requisito":
                    detalle.append({"texto": "", "imagenes": fotos})
                else:
                    imagenes.extend(fotos)
            if texto:
                lineas.append(texto)
        if clase == "textos":
            return esquema.lista_de_textos("\n".join(lineas), r"[\n,;]")
        if clase in ("linea", "opciones"):
            return " ".join(" ".join(lineas).split())
        return "\n".join(lineas)

    def tabla_anidada(self, tabla: Table, s, detalle, avisos, ident) -> None:
        nombre = s["etiqueta"] or "una celda sin etiqueta"
        detalle.append({"tabla": cv.tabla_informativa(tabla, f"tabla dentro de «{nombre}»", self.guardar)})
        if self.tipo == "prueba":
            avisos.append(cv.Aviso("tabla_en_celda", f"tabla dentro de «{nombre}»: se conserva en el contexto de la prueba", ident))

    def normalizar(self, d, detalle, imagenes, avisos) -> dict[str, Any]:
        extras = d.get("campos_adicionales") or {}
        ident = d["id"]

        def descartar(campo: str, descartados: list[str]) -> None:
            if descartados:
                nombre = NOMBRES[self.tipo][campo]
                avisos.append(cv.Aviso("valor_descartado", f"{nombre}: «{', '.join(descartados)}» no es una opción; se conserva como campo adicional", ident))
                extras[nombre] = ", ".join(descartados)

        if self.tipo == "prueba":
            for campo, normalizar in (("entorno", esquema.normalizar_entorno), ("maniobra", esquema.normalizar_maniobra)):
                if campo in d:
                    d[campo], descartados = normalizar(d[campo])
                    descartar(campo, descartados)
            if "automatizable" in d:
                valor = esquema.normalizar_automatizable(d["automatizable"])
                if valor is None:
                    descartar("automatizable", [d["automatizable"]])
                d["automatizable"] = valor or ""
            if detalle:
                d["contexto"] = [*detalle, *(d.get("contexto") or [])]
            if imagenes:
                d["imagenes"] = list(dict.fromkeys(imagenes))
            if extras:
                d["campos_adicionales"] = extras
            esquema.completar(d)
            orden = esquema.ORDEN
        else:
            if "tipo" in d:
                valor = esquema.normalizar_tipo_requisito(d["tipo"])
                if valor is None:
                    descartar("tipo", [d["tipo"]])
                d["tipo"] = valor or ""
            if "criterio_validacion" in d:
                d["criterio_validacion"], descartados = esquema.normalizar_criterio(d["criterio_validacion"])
                descartar("criterio_validacion", descartados)
            if detalle:
                d["detalle"] = detalle
            if extras:
                d["campos_adicionales"] = extras
            for k, v in esquema.requisito_vacio().items():
                if k not in d and k not in ("nombre", "nivel", "contexto"):
                    d[k] = v
            orden = esquema.REQ_ORDEN
        if not d["titulo"]:
            avisos.append(cv.Aviso("titulo_vacio", f"{'el requisito' if self.tipo == 'requisito' else 'la prueba'} no tiene título (segunda celda de la cabecera)", ident))
        return {k: d[k] for k in sorted(d, key=lambda k: orden.index(k) if k in orden else len(orden))}


class _GuardadorEnMemoria:
    """Guarda las imágenes en memoria con un nombre libre (`imagenes/imagen_001.png`…), y los
    adjuntos con el ID de su prueba (`adjuntos/DEMO-001_1.xlsx`); se escriben al importar, en la
    carpeta de cada tipo."""

    def __init__(self, ocupados: set[str] | None = None) -> None:
        self.ocupados = set(ocupados or ())
        self.por_parte: dict[str, str] = {}
        self.datos: dict[str, bytes] = {}
        self.de_dibujo: dict[str, str | None] = {}    # parte del dibujo de un objeto -> su adjunto (None: no lo es)

    def __call__(self, parte) -> str:
        nombre_parte = str(parte.partname)
        if nombre_parte not in self.por_parte:
            extension = (Path(nombre_parte).suffix or mimetypes.guess_extension(parte.content_type) or ".bin").lower()
            n = len(self.por_parte) + 1
            while f"imagen_{n:03d}{extension}" in self.ocupados:
                n += 1
            nombre = f"imagen_{n:03d}{extension}"
            self.ocupados.add(nombre)
            self.por_parte[nombre_parte] = f"imagenes/{nombre}"
            self.datos[f"imagenes/{nombre}"] = parte.blob
        return self.por_parte[nombre_parte]

    def objeto(self, o: dict[str, Any], ident: str, avisos: list, admite: bool = True) -> dict[str, Any] | None:
        """El adjunto de un objeto incrustado (ver adjuntos.objetos): su archivo y su dibujo en `adjuntos/`
        con el nombre `ident_n`. None si no se puede (queda su imagen, con un aviso) o si ya se ha guardado."""
        dibujo = str(o["vista"].partname) if o["vista"] is not None else None
        if dibujo in self.de_dibujo:
            return None
        que = adjuntos.nombre_tipo(o["prog_id"])
        adjunto = None
        if o["vinculado"]:
            avisos.append(cv.Aviso("objeto_vinculado", f"{que} vinculado a «{o['vinculado']}»: no está dentro del documento; queda su imagen", ident))
        elif o["datos"] is None:
            avisos.append(cv.Aviso("objeto_no_extraido", f"{que} incrustado ({o['prog_id'] or 'sin tipo'}): no se puede extraer; queda su imagen", ident))
        elif not admite:
            avisos.append(cv.Aviso("objeto_no_extraido", f"{que} incrustado: los requisitos no tienen adjuntos; queda su imagen", ident))
        else:
            base = re.sub(r"[^\w.-]+", "_", ident[:40]).strip("._") or "adjunto"
            usados = {Path(n).stem for n in self.ocupados}
            n = 1
            while f"{base}_{n}" in usados:
                n += 1
            raiz = f"{adjuntos.CARPETA}/{base}_{n}"
            adjunto = {"archivo": raiz + o["extension"]}
            self.datos[adjunto["archivo"]] = o["datos"]
            self.ocupados.add(Path(adjunto["archivo"]).name)
            if o["vista"] is not None:
                adjunto["vista_previa"] = raiz + (Path(dibujo).suffix.lower() or ".emf")
                self.datos[adjunto["vista_previa"]] = o["vista"].blob
                self.ocupados.add(Path(adjunto["vista_previa"]).name)
            if o["icono"]:
                adjunto["icono"] = True
            avisos.append(cv.Aviso("adjunto_incrustado", f"{que} incrustado: se guarda como adjunto ({adjunto['archivo']})", ident))
        if dibujo is not None:
            self.de_dibujo[dibujo] = adjunto["archivo"] if adjunto else None
        return adjunto


def convertir_documento(doc, mapeo: dict[str, Any], guardar_imagen=None) -> dict[str, Any]:
    analisis = analizar_documento(doc)
    familias, de_tabla = analisis["familias"], analisis["familia_de_tabla"]
    tipos = mapeo.get("tipos") or {}
    lectores = {t: _Lector(t, dict((mapeo.get("campos") or {}).get(t) or {}), guardar_imagen) for t in CLASES}
    contador = iter(range(len(de_tabla)))
    cuenta: Counter[str] = Counter()

    def leer(tabla, avisos, guardar):
        familia = familias[de_tabla[next(contador)]]
        tipo = tipos.get(familia["firma"], familia["tipo"])
        tipo = tipo if tipo in TIPOS else "tabla"
        cuenta[tipo] += 1
        if tipo == "ignorar":
            return "ignorar", None
        if tipo == "tabla":
            avisos.append(cv.Aviso("tabla_informativa", f"tabla de la familia {familia['numero']}: se conserva tal cual"))
            return "tabla", f"tabla de la familia {familia['numero']}"
        return tipo, lectores[tipo].leer(tabla, avisos)

    avisos: list = []
    capitulos = cv.extraer_capitulos(doc, avisos, guardar_imagen, leer)
    _colocar_adjuntos(capitulos, doc, guardar_imagen, avisos)
    return {"capitulos": capitulos, "avisos": avisos, "totales": {t: cuenta[t] for t in TIPOS}}


def _colocar_adjuntos(capitulos: list[dict[str, Any]], doc, guardador, avisos: list) -> None:
    """Cada imagen que es el dibujo de un objeto incrustado (un Excel…) deja de ser imagen: en el
    párrafo donde estaba queda la referencia a su adjunto. El adjunto es del elemento donde ha
    quedado (la prueba, la sección, el texto del capítulo); los de las tablas de las pruebas ya
    los ha guardado _Lector.objetos. Uno que no se puede extraer sigue siendo imagen."""
    if not hasattr(guardador, "objeto"):
        return
    parte_de = {ruta: parte for parte, ruta in guardador.por_parte.items()}
    objetos = {str(o["vista"].partname): o for o in adjuntos.objetos(doc.element.body, doc.part) if o["vista"] is not None}
    for c in capitulos:
        for e in c["contenido"]:
            if not isinstance(e, dict):
                continue
            tipo = "prueba" if isinstance(e.get("prueba"), dict) else "requisito" if isinstance(e.get("requisito"), dict) else None
            dueño = e[tipo] if tipo else e
            nombre = (str(dueño.get("id") or "") if tipo else e["seccion"] if isinstance(e.get("seccion"), str)
                      else str(c.get("capitulo") or "capitulo"))

            def adjunto_de(ruta: str, dueño=dueño, tipo=tipo, nombre=nombre) -> str | None:
                parte = parte_de.get(ruta)
                if parte not in objetos:
                    return None
                if parte not in guardador.de_dibujo:
                    nuevo = guardador.objeto(objetos[parte], nombre, avisos, tipo != "requisito")
                    if nuevo is not None:
                        dueño.setdefault("adjuntos", []).append(nuevo)
                return guardador.de_dibujo.get(parte)

            _sustituir_dibujos(dueño, adjunto_de, dueño)


def _sustituir_dibujos(valor: Any, adjunto_de: Callable[[str], str | None], dueño: dict[str, Any], en_celda: bool = False) -> None:
    """Quita de las imágenes los dibujos de los adjuntos y pone en su párrafo la referencia (en una
    celda de tabla o entre las imágenes del propio elemento no hay párrafo: van al final)."""
    if isinstance(valor, list):
        for v in valor:
            _sustituir_dibujos(v, adjunto_de, dueño, en_celda)
        return
    if not isinstance(valor, dict):
        return
    for clave in ("imagenes", "imagen"):
        rutas = valor.get(clave)
        rutas = rutas if isinstance(rutas, list) else [rutas] if isinstance(rutas, str) else []
        de_adjunto = {r: adjunto_de(r) for r in rutas}
        if not any(de_adjunto.values()):
            continue
        quedan = [r for r in rutas if not de_adjunto[r]]
        if quedan:
            valor[clave] = quedan if clave == "imagenes" else quedan[0]
        else:
            del valor[clave]
        if valor is not dueño and not en_celda:
            valor.setdefault("adjuntos", []).extend(a for a in de_adjunto.values() if a)
    for clave, hijo in list(valor.items()):
        if clave not in ("imagenes", "imagen", "adjuntos"):
            _sustituir_dibujos(hijo, adjunto_de, dueño, en_celda or clave == "filas")


def convertir_documento_ruta(ruta: str | Path, mapeo: dict[str, Any], guardador: "_GuardadorEnMemoria | None" = None) -> dict[str, Any]:
    """Conversión del Word o del Excel (las imágenes quedan en el guardador, en memoria)."""
    if es_excel(ruta):
        from app import importacion_excel
        return importacion_excel.convertir_libro(ruta, mapeo)
    return convertir_documento(Document(str(ruta)), mapeo, guardador or _GuardadorEnMemoria())


# ----------------------------------------------------------------------
# Revisión: IDs y resumen
# ----------------------------------------------------------------------
def elementos(capitulos: list[dict[str, Any]], tipo: str) -> list[dict[str, Any]]:
    """Pruebas o requisitos en el orden del documento (el mismo objeto, para cambiarles el ID)."""
    return [e[tipo] for c in capitulos for e in c["contenido"] if isinstance(e, dict) and isinstance(e.get(tipo), dict)]


def problemas_ids(capitulos: list[dict[str, Any]], tipo: str, existentes: set[str] | None = None) -> list[dict[str, Any]]:
    """IDs vacíos, repetidos (en el Word o con el proyecto) o con espacios, con un ID nuevo propuesto.

    Cada problema lleva `orden`, la posición del elemento entre los de su tipo.
    """
    existentes = set(existentes or ())
    lista = elementos(capitulos, tipo)
    usados = existentes | {str(e.get("id") or "").strip() for e in lista}
    vistos: dict[str, int] = {}
    problemas = []
    anterior = ""
    for orden, e in enumerate(lista):
        ident = str(e.get("id") or "").strip()
        motivo, propuesta = "", ""
        limpio = re.sub(r"\s+", "", ident)
        if not ident:
            motivo = "sin ID"
            propuesta = _siguiente(anterior, usados) if anterior else ""
        elif ident in existentes:
            motivo = "ya existe en el proyecto"
            propuesta = _siguiente(limpio if _ID_VALIDO.match(limpio) else ident, usados)
        elif ident in vistos:
            motivo = f"repetido (también en la posición {vistos[ident] + 1})"
            propuesta = _siguiente(limpio if _ID_VALIDO.match(limpio) else ident, usados)
        elif limpio != ident:
            motivo = "con espacios"
            propuesta = limpio if limpio not in usados else _siguiente(limpio, usados)
        vistos.setdefault(ident, orden)
        if motivo:
            if propuesta:
                usados.add(propuesta)
            problemas.append({"orden": orden, "id": ident, "titulo": str(e.get("titulo") or ""), "motivo": motivo,
                              "propuesta": propuesta})
        anterior = propuesta or ident or anterior
    return problemas


def aplicar_ids(capitulos: list[dict[str, Any]], tipo: str, nuevos: dict[Any, str]) -> int:
    """Cambia los IDs indicados ({orden: id nuevo}). Devuelve cuántos ha cambiado."""
    lista = elementos(capitulos, tipo)
    cambiados = 0
    for orden, nuevo in (nuevos or {}).items():
        nuevo = str(nuevo or "").strip()
        i = int(orden)
        if nuevo and 0 <= i < len(lista) and lista[i].get("id") != nuevo:
            lista[i]["id"] = nuevo
            cambiados += 1
    return cambiados


def separar(capitulos: list[dict[str, Any]], tipo: str) -> list[dict[str, Any]]:
    """Capítulos con los elementos de un tipo (con sus secciones, textos y tablas); sin los capítulos que no tienen ninguno."""
    otros = set(CLASES) - {tipo}
    resultado = []
    for c in capitulos:
        contenido = [e for e in c["contenido"] if not (isinstance(e, dict) and otros & set(e))]
        if any(isinstance(e, dict) and tipo in e for e in contenido):
            resultado.append({"capitulo": c["capitulo"], "contenido": contenido})
    return resultado


def _resumen_elemento(e: dict[str, Any], tipo: str) -> list[list[str]]:
    """Campos con valor de un elemento, como texto corto, para la vista previa."""
    filas = []
    for k, v in e.items():
        if k in ("nombre", "nivel", COLUMNAS) or v in (None, "", [], {}):
            continue
        nombre = NOMBRES[tipo].get(k, k)
        if k == "campos_adicionales":
            filas.extend([f"{n} (adicional)", str(x)[:300]] for n, x in v.items())
            continue
        if k == "adjuntos":
            filas.append([nombre, "\n".join(f"• {a['archivo'].split('/')[-1]}" for a in adjuntos.lista(v))])
            continue
        if isinstance(v, list):
            partes = []
            for x in v:
                if isinstance(x, dict) and "tabla" in x:
                    partes.append(f"[tabla de {len(x['tabla'].get('filas') or [])} filas]")
                elif isinstance(x, dict):
                    partes.append(str(x.get("texto") or "") + (f" [{len(x.get('imagenes') or [])} imagen(es)]" if x.get("imagenes") else ""))
                else:
                    partes.append(str(x))
            texto = "\n".join(f"• {p}" for p in partes) if k not in ("requisitos", "etiquetas", "entorno", "maniobra", "criterio_validacion") else ", ".join(partes)
        else:
            texto = str(v)
        filas.append([nombre, texto[:600]])
    return filas


# ----------------------------------------------------------------------
# Los que ya existen en el proyecto: actualizar, importar como nuevos u omitir
# ----------------------------------------------------------------------
ACCIONES = ("actualizar", "nuevo", "omitir")
# Rutas de archivos que van con el elemento: se escriben al importar y se comparan por su contenido.
ARCHIVOS = ("imagenes/", f"{adjuntos.CARPETA}/")


def _comparable(valor: Any, imagen: Callable[[str], str]) -> Any:
    """Valor para comparar: sin vacíos y con cada imagen sustituida por una huella de su contenido."""
    if isinstance(valor, dict) and "imagen" in valor:                 # forma antigua: una sola imagen
        valor = {**{k: v for k, v in valor.items() if k != "imagen"}, "imagenes": [*(valor.get("imagenes") or []), valor["imagen"]]}
    if isinstance(valor, dict):
        limpio = {k: _comparable(v, imagen) for k, v in valor.items()}
        return {k: v for k, v in limpio.items() if v is not None} or None
    if isinstance(valor, list):
        # Una imagen en un párrafo propio cuenta como parte del elemento anterior (el Word no distingue).
        unidos: list[Any] = []
        for v in valor:
            solo_imagen = isinstance(v, dict) and not v.get("texto") and (v.get("imagenes") or v.get("imagen")) and not v.get("nivel")
            if solo_imagen and unidos and not (isinstance(unidos[-1], dict) and "tabla" in unidos[-1]):
                anterior = dict(unidos[-1]) if isinstance(unidos[-1], dict) else {"texto": unidos[-1]}
                anterior["imagenes"] = [*(anterior.get("imagenes") or []), *([v["imagen"]] if v.get("imagen") else []), *(v.get("imagenes") or [])]
                unidos[-1] = anterior
            else:
                unidos.append(v)
        return [_comparable(v, imagen) for v in unidos] or None
    if isinstance(valor, str) and valor.startswith(ARCHIVOS):
        return imagen(valor)
    return None if valor in (None, "") else valor


def _huella_bytes(datos: bytes | None) -> str:
    import hashlib
    return hashlib.sha1(datos).hexdigest() if datos is not None else "(no está)"


def del_proyecto(carpeta: Path, tipo: str) -> dict[str, tuple[dict[str, Any], Any]]:
    """{ID: (capítulo, elemento)} de la carpeta (el primero, si un ID está repetido)."""
    from app.editor_core import items_in_chapter, read_chapters
    if not carpeta.is_dir() or not any(carpeta.glob("*.yaml")):
        return {}
    resultado: dict[str, tuple[dict[str, Any], Any]] = {}
    for capitulo in read_chapters(carpeta):
        for _i, e in items_in_chapter(capitulo, tipo):
            resultado.setdefault(str(e.get("id") or "").strip(), (capitulo, e))
    return resultado


def _combinado(nuevo: dict[str, Any], actual: dict[str, Any]) -> dict[str, Any]:
    """Lo que quedaría al actualizar: con un Excel, los campos que no trae siguen como estaban y sus
    campos adicionales se añaden a los que ya había; con un Word, todo es lo del Word."""
    solo = nuevo.get(COLUMNAS)
    nuevo = _sin_columnas(nuevo)
    if solo is None:
        return nuevo
    resultado = {k: v for k, v in actual.items() if k not in solo}
    resultado.update({k: v for k, v in nuevo.items() if k in solo})
    if "campos_adicionales" in solo and isinstance(actual.get("campos_adicionales"), dict):
        resultado["campos_adicionales"] = {**actual["campos_adicionales"], **(nuevo.get("campos_adicionales") or {})}
    return resultado


def diferencias(nuevo: dict[str, Any], actual: Any, tipo: str, imagenes: dict[str, bytes], carpeta: Path) -> list[str]:
    """Campos que cambiarían al actualizar `actual` con `nuevo` (sin contar encabezado ni nivel)."""
    from app.editor_api import to_plain
    actual = to_plain(actual)
    solo = nuevo.get(COLUMNAS)
    nuevo = _combinado(nuevo, actual)
    de_nuevo = lambda ruta: _huella_bytes(imagenes.get(ruta))
    del_actual = lambda ruta: _huella_bytes((carpeta / ruta).read_bytes() if (carpeta / ruta).is_file() else None)
    cambios = []
    for k in [k for k in (*nuevo, *actual) if k not in ("nombre", "nivel")]:
        if k in cambios or NOMBRES[tipo].get(k, k) in cambios or (solo is not None and k not in solo):
            continue
        x, y = _comparable(nuevo.get(k), de_nuevo), _comparable(actual.get(k), del_actual)
        if k == "campos_adicionales" and isinstance(x, dict) and isinstance(y, dict):
            # Un campo adicional en forma de lista sale en el Word como líneas: es lo mismo.
            x, y = ({n: "\n".join(map(str, v)) if isinstance(v, list) else v for n, v in d.items()} for d in (x, y))
        if CLASES[tipo].get(k) == "opciones" and isinstance(x, list) and isinstance(y, list):
            x, y = sorted(map(str, x)), sorted(map(str, y))           # «Test, Análisis» = «Análisis, Test»
        if x != y:
            cambios.append(NOMBRES[tipo].get(k, k))
    return list(dict.fromkeys(cambios))


def coincidencias(capitulos: list[dict[str, Any]], tipo: str, carpeta: Path, imagenes: dict[str, bytes]) -> list[dict[str, Any]]:
    """Elementos del Word cuyo ID ya está en el proyecto, con lo que cambiaría al actualizarlos."""
    proyecto = del_proyecto(carpeta, tipo)
    resultado = []
    for orden, e in enumerate(elementos(capitulos, tipo)):
        ident = str(e.get("id") or "").strip()
        if ident and ident in proyecto:
            capitulo, actual = proyecto[ident]
            resultado.append({"orden": orden, "id": ident, "titulo": str(e.get("titulo") or ""),
                              "capitulo": str(capitulo["data"].get("capitulo") or capitulo["path"].name),
                              "cambios": diferencias(e, actual, tipo, imagenes, carpeta)})
    return resultado


def actualizar_en_proyecto(carpeta: Path, tipo: str, nuevos: list[dict[str, Any]], imagenes: dict[str, bytes],
                           carpeta_pruebas: Path | None = None) -> int:
    """Cambia los elementos del proyecto con el mismo ID por los del Word, en su sitio.

    Se conservan su encabezado, su nivel y su posición; el resto de campos pasa
    a ser el del Word. Un requisito cambiado que ya tiene pruebas queda con sus
    pruebas por revisar (como al editarlo a mano).
    """
    from app import trazabilidad
    from app.editor_api import apply_changes, to_plain
    from app.editor_core import read_chapters, save_chapter, tests_in_chapter
    proyecto = del_proyecto(carpeta, tipo)
    revisiones = trazabilidad.leer_revisiones(carpeta) if tipo == "requisito" else {}
    citados: set[str] = set()
    if tipo == "requisito" and carpeta_pruebas is not None and carpeta_pruebas.is_dir():
        citados = {r for c in read_chapters(carpeta_pruebas) for _i, p in tests_in_chapter(c) for r in trazabilidad.lista(p.get("requisitos"))}
    tocados: dict[Path, dict[str, Any]] = {}
    cambian_revisiones = False
    for nuevo in nuevos:
        ident = str(nuevo.get("id") or "").strip()
        if ident not in proyecto:
            continue
        capitulo, actual = proyecto[ident]
        anterior = to_plain(actual)
        cambios = {k: v for k, v in _combinado(nuevo, anterior).items() if k not in ("nombre", "nivel")}
        for k in [k for k in actual if k not in cambios and k not in ("nombre", "nivel")]:
            del actual[k]
        apply_changes(actual, cambios)
        for ruta in _rutas_imagenes(cambios):
            if ruta in imagenes and not (carpeta / ruta).exists():
                (carpeta / ruta).parent.mkdir(parents=True, exist_ok=True)
                (carpeta / ruta).write_bytes(imagenes[ruta])
        tocados[capitulo["path"]] = capitulo
        if tipo == "requisito" and (ident in citados or ident in revisiones):
            cambian_revisiones = trazabilidad.anotar_cambio(revisiones, anterior, actual) or cambian_revisiones
    for capitulo in tocados.values():
        save_chapter(capitulo)
    if cambian_revisiones:
        trazabilidad.guardar_revisiones(carpeta, revisiones)
    return len(nuevos)


def revision(conversion: dict[str, Any], existentes: dict[str, set[str]] | None = None,
             carpetas: dict[str, Path] | None = None, imagenes: dict[str, bytes] | None = None) -> dict[str, Any]:
    """Lo que se va a importar: cifras, capítulos, avisos, IDs a revisar, los que ya existen y ejemplos."""
    capitulos = conversion["capitulos"]
    existentes = existentes or {}
    tipos = {}
    for tipo in CLASES:
        lista = elementos(capitulos, tipo)
        if not lista:
            continue
        caps = separar(capitulos, tipo)
        ids = problemas_ids(capitulos, tipo, existentes.get(tipo))
        ya = coincidencias(capitulos, tipo, carpetas[tipo], imagenes or {}) if carpetas and existentes.get(tipo) else []
        propuestas = {p["orden"]: p["propuesta"] for p in ids}
        for c in ya:
            c["propuesta"] = propuestas.get(c["orden"], "")
        en_proyecto = {c["orden"] for c in ya}
        del_proyecto_ = capitulos_del_proyecto(carpetas[tipo]) if carpetas and existentes.get(tipo) is not None else {}
        tipos[tipo] = {
            "total": len(lista),
            "capitulos": [{"nombre": c["capitulo"] or "(antes del primer capítulo)",
                           "elementos": sum(1 for e in c["contenido"] if isinstance(e, dict) and tipo in e),
                           "existente": _nombre_clave(c["capitulo"]) in del_proyecto_} for c in caps],
            "ids": [p for p in ids if p["orden"] not in en_proyecto],
            "existentes": ya,
            "ejemplos": [{"id": e.get("id", ""), "titulo": e.get("titulo", ""), "campos": _resumen_elemento(e, tipo)}
                         for e in lista[:3]],
        }
    grupos: dict[str, dict[str, Any]] = {}
    for a in conversion["avisos"]:
        tipo = getattr(a, "tipo", "otro")
        titulo, gravedad = TIPOS_AVISO.get(tipo, cv.TIPOS_AVISO.get(tipo, ("Otros", "aviso")))
        g = grupos.setdefault(tipo, {"tipo": tipo, "titulo": titulo, "gravedad": gravedad, "total": 0, "ejemplos": []})
        g["total"] += 1
        if len(g["ejemplos"]) < 12:
            donde = " · ".join(x for x in (getattr(a, "capitulo", ""), f"tabla {a.tabla}" if getattr(a, "tabla", None) else "", getattr(a, "id", "")) if x)
            g["ejemplos"].append({"donde": donde, "mensaje": getattr(a, "mensaje", str(a))})
    orden = list(TIPOS_AVISO) + list(cv.TIPOS_AVISO)
    return {"totales": conversion["totales"], "tipos": tipos,
            "avisos": sorted(grupos.values(), key=lambda g: orden.index(g["tipo"]) if g["tipo"] in orden else len(orden))}


TIPOS_AVISO = {
    "texto_no_importado": ("Texto del Word que no ha llegado a ningún campo", "aviso"),
    "tabla_informativa": ("Tabla informativa (se conserva tal cual)", "info"),
    "tabla_en_celda": ("Tabla dentro de una celda (se conserva en el contexto)", "aviso"),
    "campo_desconocido": ("Etiqueta mapeada a un campo que no existe", "aviso"),
    "imagen_excel": ("Imágenes del Excel (no se importan)", "aviso"),
    "adjunto_incrustado": ("Documentos incrustados (Excel…) guardados como adjuntos", "info"),
    "objeto_no_extraido": ("Objetos incrustados que no se pueden extraer (queda su imagen)", "aviso"),
    "objeto_vinculado": ("Objetos vinculados a un archivo externo (queda su imagen)", "aviso"),
}


# ----------------------------------------------------------------------
# 3. Importar
# ----------------------------------------------------------------------
def _yaml() -> YAML:
    yaml = YAML()
    yaml.allow_unicode = True
    yaml.width = 4096
    return yaml


def _rutas_imagenes(valor: Any) -> set[str]:
    if isinstance(valor, dict):
        return set().union(*(_rutas_imagenes(v) for v in valor.values())) if valor else set()
    if isinstance(valor, list):
        return set().union(*(_rutas_imagenes(v) for v in valor)) if valor else set()
    return {valor} if isinstance(valor, str) and valor.startswith(ARCHIVOS) else set()


def escribir(capitulos: list[dict[str, Any]], carpeta: Path, imagenes: dict[str, bytes]) -> list[str]:
    """Escribe los capítulos en la carpeta (nueva o existente: se añaden al final de su índice). Devuelve los archivos nuevos."""
    carpeta.mkdir(parents=True, exist_ok=True)
    yaml = _yaml()
    indice_ruta = carpeta / "indice.yaml"
    indice = []
    if indice_ruta.is_file():
        with indice_ruta.open(encoding="utf-8") as f:
            indice = list(yaml.load(f) or [])
    numeros = [int(m.group(1)) for e in indice if isinstance(e, dict)
               for m in [re.match(r"^(\d+)_", str(e.get("archivo") or ""))] if m]
    n = max(numeros + [0])
    nuevos = []
    for c in capitulos:
        n += 1
        archivo = f"{n:02d}_{cv.slug(c['capitulo'] or 'sin_capitulo')}.yaml"
        while (carpeta / archivo).exists():
            n += 1
            archivo = f"{n:02d}_{cv.slug(c['capitulo'] or 'sin_capitulo')}.yaml"
        with open(carpeta / archivo, "w", encoding="utf-8") as f:
            yaml.dump(c, f)
        indice.append({"capitulo": c["capitulo"], "archivo": archivo})
        nuevos.append(archivo)
        for ruta in _rutas_imagenes(c):
            if ruta in imagenes and not (carpeta / ruta).exists():
                (carpeta / ruta).parent.mkdir(parents=True, exist_ok=True)
                (carpeta / ruta).write_bytes(imagenes[ruta])
    with indice_ruta.open("w", encoding="utf-8") as f:
        yaml.dump(indice, f)
    formato.marcar_carpeta(carpeta)
    return nuevos


# ----------------------------------------------------------------------
# Añadir a los capítulos que ya existen (importar poco a poco)
# ----------------------------------------------------------------------
def _nombre_clave(nombre: Any) -> str:
    return clave(str(nombre or ""))


def capitulos_del_proyecto(carpeta: Path) -> dict[str, dict[str, Any]]:
    """{nombre normalizado: capítulo} de la carpeta (el primero, si dos se llaman igual)."""
    from app.editor_core import read_chapters
    if not carpeta.is_dir() or not any(carpeta.glob("*.yaml")):
        return {}
    resultado: dict[str, dict[str, Any]] = {}
    for c in read_chapters(carpeta):
        resultado.setdefault(_nombre_clave(c["data"].get("capitulo")), c)
    return resultado


def _con_contenido(contenido: list[Any], inicio: int, nivel: int) -> bool:
    """Si la sección que empieza en `inicio` tiene alguna prueba, requisito o tabla en su bloque."""
    from app.editor_core import element_level, is_table, item_key
    for e in contenido[inicio + 1:]:
        nivel_e = element_level(e)
        if nivel_e is not None and nivel_e <= nivel:
            return False                      # empieza otra sección (o un elemento) del mismo nivel
        if item_key(e) or is_table(e):
            return True
    return False


def unir_capitulo(capitulo: dict[str, Any], contenido: list[Any]) -> int:
    """Añade al capítulo del proyecto lo importado de un capítulo con el mismo nombre.

    Las secciones se buscan por su ruta de títulos (sección › subsección): si ya
    existe, lo importado va al final de esa sección; si no, se crea al final de
    la sección madre (o del capítulo). Los niveles se ajustan a los del
    proyecto. No se repiten los textos y tablas que ya estén, ni las secciones
    que se quedan vacías (todas sus pruebas ya existían). Devuelve cuántos
    elementos (pruebas o requisitos) se han añadido.
    """
    from app.editor_core import element_level, item_key, normalize_level, section_end_index
    destino = capitulo["data"].setdefault("contenido", [])

    def buscar(ruta: list[str]) -> int | None:
        """Índice de la sección del proyecto con esa ruta de títulos, o None."""
        pila: list[tuple[str, int]] = []
        for i, e in enumerate(destino):
            if isinstance(e, dict) and isinstance(e.get("seccion"), str):
                nivel = normalize_level(e.get("nivel"), 2)
                pila = [p for p in pila if p[1] < nivel] + [(_nombre_clave(e["seccion"]), nivel)]
                if [p[0] for p in pila] == ruta:
                    return i
        return None

    def nivel_de(indice: int | None) -> int:
        return normalize_level(destino[indice].get("nivel"), 2) if indice is not None else 1

    def tras_lo_suyo(indice: int) -> int:
        """Detrás de lo que cuelga directamente de la sección: antes de su primera subsección."""
        fin = section_end_index(destino, indice)
        return next((j for j in range(indice + 1, fin) if isinstance(destino[j], dict) and isinstance(destino[j].get("seccion"), str)), fin)

    def ya_esta(elemento: Any) -> bool:
        from app.editor_api import to_plain
        return to_plain(elemento) in [to_plain(e) for e in destino]

    ruta: list[tuple[str, int]] = []          # secciones abiertas de lo importado: (nombre, nivel en el Word)
    actual: int | None = None                 # sección del proyecto donde va lo que sigue (None: el capítulo)
    añadidos = 0
    for i, e in enumerate(contenido):
        if isinstance(e, dict) and isinstance(e.get("seccion"), str):
            nivel = normalize_level(e.get("nivel"), 2)
            ruta = [p for p in ruta if p[1] < nivel] + [(_nombre_clave(e["seccion"]), nivel)]
            encontrada = buscar([p[0] for p in ruta])
            if encontrada is not None:
                actual = encontrada
                continue
            madre = buscar([p[0] for p in ruta[:-1]]) if len(ruta) > 1 else None
            if not _con_contenido(contenido, i, nivel):
                actual = madre                # sección sin nada nuevo: no se crea
                continue
            nueva = {**e, "nivel": min(nivel_de(madre) + 1, 6)}
            posicion = section_end_index(destino, madre) if madre is not None else len(destino)
            destino.insert(posicion, nueva)
            actual = posicion
            continue
        tipo = item_key(e)
        if not tipo and ya_esta(e):
            continue                          # el mismo texto o tabla de la vez anterior
        if tipo:
            e = {tipo: {**e[tipo], "nivel": min(nivel_de(actual) + 1, 6)}}
            añadidos += 1
        posicion = tras_lo_suyo(actual) if actual is not None else len(destino)
        if actual is None and not tipo:
            # Texto o tabla del capítulo (antes de su primera sección): delante de las secciones.
            posicion = next((j for j, x in enumerate(destino) if element_level(x) is not None), len(destino))
        destino.insert(posicion, e)
    return añadidos


def unir_a_existentes(capitulos: list[dict[str, Any]], carpeta: Path, imagenes: dict[str, bytes]) -> tuple[list[dict[str, Any]], list[str]]:
    """Añade a los capítulos del proyecto los importados con el mismo nombre. Devuelve los que no
    tienen pareja (se escriben como capítulos nuevos) y los nombres de los archivos que han cambiado."""
    from app.editor_core import save_chapter
    del_proyecto_ = capitulos_del_proyecto(carpeta)
    restantes, cambiados = [], []
    for c in capitulos:
        existente = del_proyecto_.get(_nombre_clave(c["capitulo"]))
        if existente is None:
            restantes.append(c)
            continue
        unir_capitulo(existente, c["contenido"])
        save_chapter(existente)
        cambiados.append(existente["path"].name)
        for ruta in _rutas_imagenes(c):
            if ruta in imagenes and not (carpeta / ruta).exists():
                (carpeta / ruta).parent.mkdir(parents=True, exist_ok=True)
                (carpeta / ruta).write_bytes(imagenes[ruta])
    return restantes, cambiados


def carpetas_destino(destino: str | Path, modo: str) -> dict[str, Path]:
    """Carpetas de pruebas y de requisitos: las de un proyecto nuevo o las del abierto."""
    destino = Path(destino)
    if modo == "nuevo":
        return {"prueba": destino / esquema.CARPETA_PRUEBAS, "requisito": destino / esquema.CARPETA_REQUISITOS}
    from app.editor_core import requirements_dir, tests_dir
    pruebas = tests_dir(destino)
    return {"prueba": pruebas, "requisito": requirements_dir(pruebas)}


def ids_existentes(carpetas: dict[str, Path]) -> dict[str, set[str]]:
    from app.editor_core import items_in_chapter, read_chapters
    resultado: dict[str, set[str]] = {}
    for tipo, carpeta in carpetas.items():
        ids = set()
        if carpeta.is_dir() and any(carpeta.glob("*.yaml")):
            for c in read_chapters(carpeta):
                ids |= {str(e.get("id") or "").strip() for _i, e in items_in_chapter(c, tipo)}
        resultado[tipo] = ids - {""}
    return resultado


def imagenes_ocupadas(carpetas: dict[str, Path]) -> set[str]:
    """Nombres ya usados en las carpetas de imágenes y de adjuntos."""
    return {p.name for c in carpetas.values() for sub in ("imagenes", adjuntos.CARPETA)
            if (c / sub).is_dir() for p in (c / sub).iterdir()}


def preparar(ruta: str | Path, mapeo: dict[str, Any], destino: str | Path | None = None, modo: str = "nuevo") -> dict[str, Any]:
    """Convierte el Word para revisarlo o importarlo, teniendo en cuenta lo que ya hay en el destino."""
    carpetas = carpetas_destino(destino, modo) if destino else {}
    existentes = ids_existentes(carpetas) if destino and modo == "anadir" else {}
    guardador = _GuardadorEnMemoria(imagenes_ocupadas(carpetas) if destino and modo == "anadir" else None)
    conversion = convertir_documento_ruta(ruta, mapeo, guardador)
    return {"conversion": conversion, "carpetas": carpetas, "existentes": existentes, "imagenes": guardador.datos}


def importar(ruta: str | Path, mapeo: dict[str, Any], destino: str | Path, modo: str = "nuevo",
             ids: dict[str, dict[Any, str]] | None = None, renombrar_propuestos: bool = False,
             acciones: dict[str, dict[Any, str]] | None = None) -> dict[str, Any]:
    """Importa el Word en `destino` (modo "nuevo": carpeta de un proyecto nuevo; "anadir": el proyecto abierto).

    `ids`: IDs nuevos por tipo ({"prueba": {orden: id}}); con `renombrar_propuestos`,
    los que no se indiquen toman el ID propuesto en la revisión.
    `acciones`: qué hacer con los que ya existen en el proyecto ({"requisito": {orden:
    "actualizar" | "nuevo" | "omitir"}}); sin indicar, se importan como nuevos.
    """
    if modo not in ("nuevo", "anadir"):
        raise ValueError(f"Modo de importación desconocido: {modo}")
    destino = Path(destino)
    if modo == "nuevo" and destino.exists() and any(destino.iterdir()):
        raise ValueError(f"La carpeta {destino} no está vacía: elige una carpeta nueva para el proyecto.")
    p = preparar(ruta, mapeo, destino, modo)
    for tipo, carpeta in p["carpetas"].items():
        formato.comprobar_carpeta(carpeta, f"La carpeta de {TIPOS_VISIBLES[tipo].lower()}")
    conversion, capitulos = p["conversion"], p["conversion"]["capitulos"]
    resultado: dict[str, Any] = {"carpetas": {}, "archivos": {}, "unidos": {}, "totales": {}, "ids_cambiados": {},
                                 "actualizados": {}, "omitidos": {}}
    for tipo in CLASES:
        lista = elementos(capitulos, tipo)
        if not lista:
            continue
        # Los que ya existen: actualizar en su sitio, omitir o (por defecto) importar como nuevos.
        elegidas = {int(k): v for k, v in ((acciones or {}).get(tipo) or {}).items()}
        if any(v not in ACCIONES for v in elegidas.values()):
            raise ValueError(f"Acción desconocida: elige entre {', '.join(ACCIONES)}.")
        ya = p["existentes"].get(tipo) or set()
        apartar = {orden: elegidas.get(orden, "nuevo") for orden, e in enumerate(lista)
                   if str(e.get("id") or "").strip() in ya and elegidas.get(orden, "nuevo") != "nuevo"}
        nuevos = {k: v for k, v in dict((ids or {}).get(tipo) or {}).items() if int(k) not in apartar}
        if renombrar_propuestos:
            for prob in problemas_ids(capitulos, tipo, ya):
                if prob["propuesta"] and prob["orden"] not in apartar:
                    nuevos.setdefault(str(prob["orden"]), prob["propuesta"])
        resultado["ids_cambiados"][tipo] = aplicar_ids(capitulos, tipo, nuevos)
        if apartar:
            quitar = {id(lista[o]) for o in apartar}
            para_actualizar = [lista[o] for o, accion in apartar.items() if accion == "actualizar"]
            for c in capitulos:
                c["contenido"] = [e for e in c["contenido"] if not (isinstance(e, dict) and isinstance(e.get(tipo), dict) and id(e[tipo]) in quitar)]
            if para_actualizar:
                resultado["actualizados"][tipo] = actualizar_en_proyecto(
                    p["carpetas"][tipo], tipo, para_actualizar, p["imagenes"], p["carpetas"].get("prueba"))
                resultado["carpetas"][tipo] = str(p["carpetas"][tipo])
            resultado["omitidos"][tipo] = sum(1 for a in apartar.values() if a == "omitir")
    for tipo in CLASES:                       # la clave de las columnas no se escribe
        for e in elementos(capitulos, tipo):
            e.pop(COLUMNAS, None)
    for tipo in CLASES:
        caps = separar(capitulos, tipo)
        if not caps:
            continue
        carpeta = p["carpetas"][tipo]
        nuevos_caps = caps
        if modo == "anadir":                  # los capítulos que ya existen reciben lo nuevo en su sitio
            nuevos_caps, resultado["unidos"][tipo] = unir_a_existentes(caps, carpeta, p["imagenes"])
        resultado["archivos"][tipo] = escribir(nuevos_caps, carpeta, p["imagenes"]) if nuevos_caps else []
        resultado["carpetas"][tipo] = str(carpeta)
        resultado["totales"][tipo] = len(elementos(caps, tipo))
    if not resultado["carpetas"]:
        if any(resultado["omitidos"].values()):
            raise ValueError("No hay nada que importar: todos los que ya existían se han marcado como «No importar».")
        raise ValueError("Con este mapeo no hay ninguna prueba ni ningún requisito que importar.")
    if modo == "nuevo" and "prueba" not in resultado["carpetas"]:
        # El editor abre siempre una carpeta de pruebas: el proyecto empieza con un capítulo vacío.
        escribir([{"capitulo": "Pruebas", "contenido": []}], p["carpetas"]["prueba"], {})
    if modo == "nuevo":
        cv.escribir_avisos(conversion["avisos"], destino)
        esquema.copiar_configuracion(destino)     # el proyecto nuevo conserva las opciones del abierto
    resultado["avisos"] = len(conversion["avisos"])
    return resultado


# ----------------------------------------------------------------------
# Línea de órdenes
# ----------------------------------------------------------------------
def main(argv=None, prog=None) -> None:
    ap = argparse.ArgumentParser(prog=prog, description="Importa pruebas o requisitos de un Word o un Excel con un mapeo de etiquetas")
    ap.add_argument("docx", help="documento Word (.docx) o Excel (.xlsx)")
    ap.add_argument("-o", "--salida", help="carpeta del proyecto nuevo (sin ella, solo muestra el análisis)")
    ap.add_argument("--tipo", choices=["prueba", "requisito"], help="qué se espera encontrar (ayuda a proponer)")
    ap.add_argument("--mapeo", help="mapeo en JSON (el que guarda el editor o --guardar-mapeo)")
    ap.add_argument("--guardar-mapeo", help="escribe el mapeo propuesto en este JSON para revisarlo")
    ap.add_argument("--renombrar", action="store_true", help="da el ID propuesto a los repetidos o vacíos")
    ap.add_argument("--configuracion", help="aztarna.yaml con las opciones de maniobra y entorno (se copia al proyecto)")
    args = ap.parse_args(argv)
    if args.configuracion:
        esquema.configurar(args.configuracion)
    analisis = analizar(args.docx, args.tipo)
    mapeo = mapeo_propuesto(analisis, json.loads(Path(args.mapeo).read_text(encoding="utf-8")) if args.mapeo else None)
    if analisis.get("origen") == "excel":
        print(f"{analisis['tablas']} filas en {len(analisis['familias'])} hojas:")
    else:
        print(f"{analisis['tablas']} tablas en {len(analisis['familias'])} familias:")
    for f in analisis["familias"]:
        tipo = mapeo["tipos"][f["firma"]]
        print(f"  {f['numero']:3d}. {f['tablas']:5d} {'tabla ' if f['tablas'] == 1 else 'tablas'} → {TIPOS_VISIBLES[tipo]}")
        if tipo in CLASES:
            for e in f["etiquetas"]:
                campo = mapeo["campos"][tipo][e["clave"]]
                print(f"         {e['etiqueta'][:40]:40s} → {ESPECIALES.get(campo) or NOMBRES[tipo].get(campo, campo)}")
    if args.guardar_mapeo:
        Path(args.guardar_mapeo).write_text(json.dumps(mapeo, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"Mapeo -> {args.guardar_mapeo}")
    if args.salida:
        r = importar(args.docx, mapeo, args.salida, "nuevo", renombrar_propuestos=args.renombrar)
        for tipo, total in r["totales"].items():
            print(f"{total} {TIPOS_VISIBLES[tipo].lower()} -> {r['carpetas'][tipo]}")
        print(f"{r['avisos']} avisos (ver {Path(args.salida) / 'avisos.txt'})")


if __name__ == "__main__":
    main()
