#!/usr/bin/env python3
"""
Genera un único DOCX a partir de los YAML de pruebas (o de requisitos).

Uso:
    python herramientas.py word salida/yaml                 -> pruebas.docx (usa indice.yaml)
    python herramientas.py word salida/yaml -o mi_doc.docx
    python herramientas.py word capitulo.yaml               (un solo archivo)
    python herramientas.py word salida/yaml/requisitos -o requisitos.docx

Requisitos:
    pip install python-docx ruamel.yaml
"""
import argparse
import io
import itertools
import re
import struct
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls
from docx.oxml.ns import qn
from docx.oxml.shape import CT_Inline
from docx.opc.constants import CONTENT_TYPE, RELATIONSHIP_TYPE
from docx.opc.part import Part
from docx.parts.image import ImagePart
from docx.shared import Cm, Pt, RGBColor
from ruamel.yaml import YAML

from app import adjuntos as adj
from app import firma
from app.esquema import AUTOMATIZABLE_VISIBLE, ETIQUETAS, REQ_ETIQUETAS
from app.editor_core import NOT_CHAPTERS

# ----------------------------------------------------------------------
# CONFIGURACIÓN (todo el aspecto visual se cambia aquí, en un solo sitio)
# ----------------------------------------------------------------------
FUENTE = "Calibri"
TAMANO = 10
COLOR_CABECERA = "00A651"      # verde de la cabecera
COLOR_TEXTO_CABECERA = RGBColor(0xFF, 0xFF, 0xFF)
ANCHO_TABLA_CM = 17.0          # ancho útil de la página A4 con márgenes de 2 cm
ANCHO_IMAGEN_CM = 10.0         # ancho de las imágenes incrustadas…
ALTO_MAX_IMAGEN_CM = 9.0       # …salvo que así queden más altas que esto (capturas verticales)
ANCHOS_PRUEBA_CM = (4.66, 4.66, 7.35)   # columnas de la plantilla (plantillas/tabla_prueba.docx)
ANCHOS_REQUISITO_CM = (8.91, 3.99, 3.79)   # columnas de la plantilla (plantillas/tabla_requisito.docx)

# Nombres de color -> hex (los mismos que usa convertir.py)
COLORES = {"rojo": "FF0000", "azul": "0070C0", "verde": "00B050"}
RESALTADOS = {
    "amarillo": WD_COLOR_INDEX.YELLOW,
    "verde": WD_COLOR_INDEX.BRIGHT_GREEN,
    "cian": WD_COLOR_INDEX.TURQUOISE,
    "yellow": WD_COLOR_INDEX.YELLOW,
    "bright_green": WD_COLOR_INDEX.BRIGHT_GREEN,
    "turquoise": WD_COLOR_INDEX.TURQUOISE,
}

# Cómo se muestra cada valor guardado en el YAML


# ----------------------------------------------------------------------
# Utilidades de formato de tabla
# ----------------------------------------------------------------------
def sombrear(celda, hex_color):
    tcPr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def fijar_anchos(tabla, anchos_cm):
    """Ancho fijo de tabla y columnas (más rápido de maquetar para Word)."""
    tabla.autofit = False
    tblPr = tabla._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)
    for col, ancho in zip(tabla.columns, anchos_cm):
        col.width = Cm(ancho)
    for fila in tabla.rows:
        for i, celda in enumerate(fila.cells):
            celda.width = Cm(anchos_cm[i])


def compactar(parrafo):
    pf = parrafo.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)


# ----------------------------------------------------------------------
# Marcas inline: **negrita**, *cursiva*, [[color|texto]], [[res:color|texto]]
# ----------------------------------------------------------------------
TOKEN = re.compile(r"(\[\[[^\]|]+\||\]\]|\*\*|\*)")


def escribir_inline(parrafo, texto):
    negrita = cursiva = False
    pila = []  # lista de ("color"|"res", valor)
    for trozo in TOKEN.split(texto or ""):
        if not trozo:
            continue
        if trozo == "**":
            negrita = not negrita
        elif trozo == "*":
            cursiva = not cursiva
        elif trozo.startswith("[[") and trozo.endswith("|"):
            clave = trozo[2:-1]
            if clave.startswith("res:"):
                pila.append(("res", clave[4:]))
            else:
                pila.append(("color", clave))
        elif trozo == "]]":
            if pila:
                pila.pop()
        else:
            run = parrafo.add_run(trozo)
            run.bold = negrita or None
            run.italic = cursiva or None
            for tipo, valor in pila:
                if tipo == "color":
                    hexa = COLORES.get(valor, valor.lstrip("#"))
                    try:
                        run.font.color.rgb = RGBColor.from_string(hexa.upper())
                    except ValueError:
                        pass
                else:
                    ind = RESALTADOS.get(valor)
                    if ind is not None:
                        run.font.highlight_color = ind


# ----------------------------------------------------------------------
# Listas con numeración real de Word (la numeración reinicia en cada lista)
# ----------------------------------------------------------------------
def nueva_numeracion(doc, estilo):
    """Crea una instancia de numeración que empieza en 1 para el estilo dado."""
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    num_id = doc.styles[estilo].element.pPr.numPr.numId.val
    abstract_id = numbering.num_having_numId(num_id).abstractNumId.val
    nuevo = numbering.add_num(abstract_id)
    nuevo.add_lvlOverride(ilvl=0).add_startOverride(1)
    return nuevo.numId


def poner_numeracion(parrafo, num_id, nivel):
    pPr = parrafo._p.get_or_add_pPr()
    numPr = pPr.get_or_add_numPr()
    numPr.get_or_add_ilvl().val = nivel
    numPr.get_or_add_numId().val = num_id


def insertar_imagen(celda, base_dir, imagen):
    ruta = base_dir / imagen
    if ruta.exists():
        run = celda.add_paragraph().add_run()
        agregar_imagen(run, ruta, Cm(ANCHO_IMAGEN_CM))
    else:
        print(f"  AVISO: imagen no encontrada: {ruta}")


def _tamaño_imagen(ruta):
    """(ancho, alto) en píxeles de una imagen normal, o None si no se puede leer."""
    try:
        from PIL import Image

        with Image.open(ruta) as imagen:
            return imagen.size
    except Exception:  # noqa: BLE001 - si no se puede leer, se usa el ancho fijo
        return None


def _ancho_limitado(ancho, ancho_px, alto_px):
    """Reduce el ancho si con él la imagen superaría ALTO_MAX_IMAGEN_CM de alto."""
    if ancho_px > 0 and alto_px > 0:
        alto_max = Cm(ALTO_MAX_IMAGEN_CM)
        if ancho * alto_px / ancho_px > alto_max:
            return int(alto_max * ancho_px / alto_px)
    return ancho


def agregar_imagen(run, ruta, ancho):
    if ruta.suffix.lower() != ".emf":
        tamaño = _tamaño_imagen(ruta)
        if tamaño:
            ancho = _ancho_limitado(ancho, *tamaño)
        run.add_picture(str(ruta), width=ancho)
        return

    blob = ruta.read_bytes()
    if len(blob) < 108 or blob[40:44] != b" EMF":
        raise ValueError(f"Archivo EMF inválido o truncado: {ruta}")

    izquierda, arriba, derecha, abajo = struct.unpack_from("<4i", blob, 24)
    ancho_marco = derecha - izquierda
    alto_marco = abajo - arriba
    if ancho_marco <= 0 or alto_marco <= 0:
        raise ValueError(f"El EMF tiene dimensiones inválidas: {ruta}")

    ancho = _ancho_limitado(ancho, ancho_marco, alto_marco)
    alto = int(ancho * alto_marco / ancho_marco)
    part = run.part
    image_part = ImagePart(
        part.package.next_partname("/word/media/image%d.emf"),
        CONTENT_TYPE.X_EMF,
        blob,
    )
    relationship_id = part.relate_to(image_part, RELATIONSHIP_TYPE.IMAGE)
    inline = CT_Inline.new_pic_inline(
        part.next_id,
        relationship_id,
        ruta.name,
        ancho,
        alto,
    )
    run._r.add_drawing(inline)


# ----------------------------------------------------------------------
# Adjuntos: objetos incrustados (ver adjuntos.py)
# ----------------------------------------------------------------------
# Adjuntos del elemento que se está escribiendo (prueba, sección…) que aún no han salido en el
# documento: [{archivo, vista, icono}] (rutas absolutas). Salen en el párrafo que los nombra
# (insertar_adjuntos); los que no, al final de la tabla de la prueba (fila_adjuntos) o tras el
# contexto (parrafos_adjuntos).
_ADJUNTOS_PENDIENTES: list[dict] = []
_IDS_OBJETO = itertools.count(1)
ANCHO_ICONO_CM = 2.4

_TIPO_FORMA = (
    '<v:shapetype id="_x0000_t75" coordsize="21600,21600" o:spt="75" o:preferrelative="t" '
    'path="m@4@5l@4@11@9@11@9@5xe" filled="f" stroked="f"><v:stroke joinstyle="miter"/><v:formulas>'
    + "".join(f'<v:f eqn="{e}"/>' for e in (
        "if lineDrawn pixelLineWidth 0", "sum @0 1 0", "sum 0 0 @1", "prod @2 1 2", "prod @3 21600 pixelWidth",
        "prod @3 21600 pixelHeight", "sum @0 0 1", "prod @6 1 2", "prod @7 21600 pixelWidth", "sum @8 21600 0",
        "prod @7 21600 pixelHeight", "sum @10 21600 0"))
    + '</v:formulas><v:path o:extrusionok="f" gradientshapeok="t" o:connecttype="rect"/>'
    '<o:lock v:ext="edit" aspectratio="t"/></v:shapetype>'
)


def preparar_adjuntos(elemento, base_dir):
    """Antes de escribir una prueba (su contexto y su tabla) o una sección: los adjuntos que tiene."""
    _ADJUNTOS_PENDIENTES[:] = [
        {"archivo": (base_dir / a["archivo"]).resolve(),
         "vista": (base_dir / a["vista_previa"]).resolve() if a.get("vista_previa") else None,
         "icono": bool(a.get("icono"))}
        for a in adj.lista(elemento.get("adjuntos"))]


def insertar_adjuntos(destino, item, base_dir):
    """Los adjuntos que van en un párrafo (su clave `adjuntos`), cada uno en un párrafo de `destino`."""
    for ruta in adj.referencias(item):
        archivo = (base_dir / ruta).resolve()
        adjunto = next((a for a in _ADJUNTOS_PENDIENTES if a["archivo"] == archivo), None)
        if adjunto is None:                   # no está en la lista del elemento: sin su dibujo
            adjunto = {"archivo": archivo, "vista": None, "icono": True}
        else:
            _ADJUNTOS_PENDIENTES.remove(adjunto)
        parrafo = destino.add_paragraph()
        compactar(parrafo)
        escribir_adjunto(parrafo, adjunto)


def fila_adjuntos(t):
    """Fila «Adjuntos» al final de la tabla de la prueba con los que no han salido en su sitio."""
    if _ADJUNTOS_PENDIENTES:
        fila = t.add_row()
        _escribir_pendientes(fila.cells[0].merge(fila.cells[-1]), primero=True)


def parrafos_adjuntos(doc):
    """Tras el contexto de una sección (o un texto del capítulo), los adjuntos que no han salido en su sitio."""
    if _ADJUNTOS_PENDIENTES:
        _escribir_pendientes(doc, primero=False)


def _escribir_pendientes(destino, primero):
    """«Adjuntos:» y cada uno incrustado (con su vista previa o como icono) o, si Word no
    lo sabe incrustar, con su nombre. `destino` es una celda o el documento."""
    pendientes = list(_ADJUNTOS_PENDIENTES)
    _ADJUNTOS_PENDIENTES.clear()
    p = destino.paragraphs[0] if primero else destino.add_paragraph()
    compactar(p)
    p.add_run(f"{ETIQUETAS['adjuntos']}:").bold = True
    for a in pendientes:
        parrafo = destino.add_paragraph()
        compactar(parrafo)
        escribir_adjunto(parrafo, a)


def escribir_adjunto(parrafo, adjunto):
    """El adjunto incrustado (con su dibujo o con un icono) o, si Word no lo sabe incrustar, su nombre."""
    archivo = adjunto["archivo"]
    if not archivo.is_file():
        print(f"  AVISO: adjunto no encontrado: {archivo}")
        parrafo.add_run(f"{archivo.name} (no se encuentra)")
    elif adj.incrustable(archivo):
        vista = adjunto["vista"] if adjunto["vista"] is not None and adjunto["vista"].is_file() else None
        incrustar_objeto(parrafo.add_run(), archivo, vista, Cm(ANCHO_IMAGEN_CM), adjunto["icono"])
    else:
        parrafo.add_run(archivo.name)


def _imagen_de_vista(part, vista):
    """Relaciona el dibujo con la parte; devuelve (rId, ancho, alto) a su tamaño natural, en EMU."""
    extension = vista.suffix.lower()
    if extension in (".emf", ".wmf"):
        blob = vista.read_bytes()
        if extension == ".emf" and len(blob) >= 108 and blob[40:44] == b" EMF":
            izquierda, arriba, derecha, abajo = struct.unpack_from("<4i", blob, 24)   # centésimas de mm
            escala = 360
        elif extension == ".wmf" and len(blob) >= 22 and blob[:4] == b"\xd7\xcd\xc6\x9a":   # WMF «placeable»
            izquierda, arriba, derecha, abajo = struct.unpack_from("<4h", blob, 6)
            escala = 914400 / (struct.unpack_from("<H", blob, 14)[0] or 1440)            # unidades por pulgada
        else:
            raise ValueError(f"Imagen {extension[1:].upper()} inválida: {vista}")
        tipo = CONTENT_TYPE.X_EMF if extension == ".emf" else CONTENT_TYPE.X_WMF
        imagen = ImagePart(part.package.next_partname(f"/word/media/image%d{extension}"), tipo, blob)
        rid = part.relate_to(imagen, RELATIONSHIP_TYPE.IMAGE)
        return rid, int((derecha - izquierda) * escala), int((abajo - arriba) * escala)
    rid, imagen = part.get_or_add_image(str(vista))
    return rid, imagen.width, imagen.height


def icono_adjunto(nombre):
    """PNG con un icono de documento y el nombre del archivo, para un adjunto sin vista previa."""
    from PIL import Image, ImageDraw, ImageFont

    ancho, alto = 240, 200
    imagen = Image.new("RGB", (ancho, alto), "white")
    dibujo = ImageDraw.Draw(imagen)
    verde = "#" + COLOR_CABECERA
    dibujo.polygon([(80, 10), (140, 10), (165, 35), (165, 120), (80, 120)], fill="white", outline=verde, width=4)
    dibujo.polygon([(140, 10), (140, 35), (165, 35)], fill=verde)
    def fuente(tamaño):
        try:
            return ImageFont.load_default(tamaño)
        except TypeError:                 # Pillow antiguo: una sola fuente, de tamaño fijo
            return ImageFont.load_default()
    dibujo.text((122, 80), Path(nombre).suffix[1:].upper()[:4], fill=verde, font=fuente(26), anchor="mm")
    texto = nombre if len(nombre) <= 30 else nombre[:27] + "…"
    dibujo.text((ancho / 2, 160), texto, fill="black", font=fuente(15), anchor="mm")
    salida = io.BytesIO()
    imagen.save(salida, format="PNG")
    salida.seek(0)
    return salida


def incrustar_objeto(run, archivo, vista, ancho, icono=False):
    """El adjunto como objeto incrustado en el párrafo.

    Se ve con su dibujo (el del documento original, a su tamaño, sin pasar de
    `ancho`) o, sin él, con un icono con el nombre del archivo. Como icono
    (`icono`, o sin dibujo), un doble clic lo abre en su programa (Excel…);
    si no, se edita dentro del documento.
    """
    extension = archivo.suffix.lower()
    part = run.part
    paquete = part.package
    if extension in adj.PAQUETES:
        prog_id, tipo = adj.PAQUETES[extension]
        nombre, relacion = f"/word/embeddings/Adjunto%d{extension}", RELATIONSHIP_TYPE.PACKAGE
    else:
        prog_id, tipo = adj.OLE_ANTIGUOS[extension], CONTENT_TYPE.OFC_OLE_OBJECT
        nombre, relacion = "/word/embeddings/oleObject%d.bin", RELATIONSHIP_TYPE.OLE_OBJECT
    objeto = Part(paquete.next_partname(nombre), tipo, archivo.read_bytes(), paquete)
    rid_objeto = part.relate_to(objeto, relacion)
    if vista is not None:
        rid_vista, ancho_vista, alto_vista = _imagen_de_vista(part, vista)
        if ancho_vista <= 0 or alto_vista <= 0:
            raise ValueError(f"La vista previa tiene dimensiones inválidas: {vista}")
        ancho = _ancho_limitado(min(ancho, ancho_vista), ancho_vista, alto_vista)
        alto, aspecto = int(ancho * alto_vista / ancho_vista), "Icon" if icono else "Content"
    else:
        rid_vista, _imagen = part.get_or_add_image(icono_adjunto(archivo.name))
        ancho, alto, aspecto = Cm(ANCHO_ICONO_CM), Cm(ANCHO_ICONO_CM * 200 / 240), "Icon"
    n = next(_IDS_OBJETO)
    xml = (
        f'<w:object {nsdecls("w", "r")} xmlns:v="urn:schemas-microsoft-com:vml" '
        f'xmlns:o="urn:schemas-microsoft-com:office:office" w:dxaOrig="{int(ancho / 635)}" w:dyaOrig="{int(alto / 635)}">'
        f'{_TIPO_FORMA}'
        f'<v:shape id="_x0000_i{1024 + n}" type="#_x0000_t75" style="width:{ancho / 12700:.1f}pt;height:{alto / 12700:.1f}pt" o:ole="">'
        f'<v:imagedata r:id="{rid_vista}" o:title=""/></v:shape>'
        f'<o:OLEObject Type="Embed" ProgID="{prog_id}" ShapeID="_x0000_i{1024 + n}" DrawAspect="{aspecto}" '
        f'ObjectID="_{1700000000 + n}" r:id="{rid_objeto}"/></w:object>'
    )
    run._r.append(parse_xml(xml))


def escribir_lista(doc, celda, items, numerada, base_dir):
    estilo = "List Number" if numerada else "List Bullet"
    num_id = nueva_numeracion(doc, estilo) if numerada else None
    for it in items or []:
        if isinstance(it, dict):
            texto, nivel = it.get("texto", ""), int(it.get("nivel", 0))
            imagenes = it.get("imagenes", [])
            if it.get("imagen"):
                imagenes = [*imagenes, it["imagen"]]
        else:
            texto, nivel, imagenes = str(it), 0, []
        p = celda.add_paragraph(style=estilo)
        compactar(p)
        if numerada:
            poner_numeracion(p, num_id, nivel)
        elif nivel:
            poner_numeracion(p, doc.styles[estilo].element.pPr.numPr.numId.val, nivel)
        escribir_inline(p, texto)
        for imagen in imagenes:
            insertar_imagen(celda, base_dir, imagen)
        insertar_adjuntos(celda, it, base_dir)


# ----------------------------------------------------------------------
# Construcción de la tabla de una prueba
# ----------------------------------------------------------------------
def etiqueta_valor(celda, etiqueta, valor):
    p = celda.paragraphs[0]
    compactar(p)
    r = p.add_run(f"{etiqueta}: ")
    r.bold = True
    escribir_inline(p, valor or "")


def etiqueta_lista(celda, primera, etiqueta, items, doc, base_dir, numerada=False):
    """«Etiqueta:» en negrita y debajo la lista (viñetas o numerada)."""
    p = celda.paragraphs[0] if primera else celda.add_paragraph()
    compactar(p)
    p.add_run(f"{etiqueta}:").bold = True
    escribir_lista(doc, celda, items, numerada, base_dir)


def filas_adicionales(t, extras):
    """Una fila más al final de la tabla por cada campo adicional: «Etiqueta: valor»."""
    for etiqueta, valor in (extras or {}).items():
        fila = t.add_row()
        celda = fila.cells[0].merge(fila.cells[-1])
        p = celda.paragraphs[0]
        compactar(p)
        p.add_run(f"{etiqueta}: ").bold = True
        lineas = valor if isinstance(valor, list) else str(valor or "").splitlines()
        if lineas:
            escribir_inline(p, str(lineas[0]))
            for linea in lineas[1:]:
                parrafo = celda.add_paragraph()
                compactar(parrafo)
                escribir_inline(parrafo, str(linea))


def escribir_detalle(celda, items, base_dir, ancho):
    """Tablas, imágenes y párrafos que siguen a la descripción, dentro de su celda."""
    for item in items or []:
        if isinstance(item, dict) and isinstance(item.get("tabla"), dict):
            tabla_informativa(celda, item["tabla"], base_dir, ancho)
            continue
        p = celda.add_paragraph()
        compactar(p)
        escribir_inline(p, item.get("texto", "") if isinstance(item, dict) else str(item))
        if isinstance(item, dict):
            for imagen in [*(item.get("imagenes") or []), *([item["imagen"]] if item.get("imagen") else [])]:
                insertar_imagen(celda, base_dir, imagen)
            insertar_adjuntos(celda, item, base_dir)


def unir(valores, separador):
    valores = valores if isinstance(valores, list) else ([valores] if valores else [])
    return separador.join(str(v) for v in valores)


def tabla_prueba(doc, prueba, base_dir):
    """Tabla de una prueba con la plantilla de plantillas/tabla_prueba.docx (campos en esquema.py)."""
    t = doc.add_table(rows=7, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fijar_anchos(t, list(ANCHOS_PRUEBA_CM))

    def fusionar(fila, ini, fin):
        return t.cell(fila, ini).merge(t.cell(fila, fin)) if fin > ini else t.cell(fila, ini)

    # Fila 0: ID | título (verde)
    for c, txt, cursiva in ((t.cell(0, 0), prueba.get("id", ""), False),
                            (fusionar(0, 1, 2), prueba.get("titulo", ""), True)):
        sombrear(c, COLOR_CABECERA)
        p = c.paragraphs[0]
        compactar(p)
        r = p.add_run(str(txt or ""))
        r.bold, r.italic = True, cursiva
        r.font.color.rgb = COLOR_TEXTO_CABECERA

    # Fila 1: descripción
    etiqueta_valor(fusionar(1, 0, 2), ETIQUETAS["descripcion"], prueba.get("descripcion"))

    # Fila 2: parametrización
    etiqueta_lista(fusionar(2, 0, 2), True, ETIQUETAS["parametrizacion"], prueba.get("parametrizacion"), doc, base_dir)

    # Fila 3: situación inicial y ensayo (numerado)
    c = fusionar(3, 0, 2)
    etiqueta_lista(c, True, ETIQUETAS["situacion_inicial"], prueba.get("situacion_inicial"), doc, base_dir)
    etiqueta_lista(c, False, ETIQUETAS["pasos"], prueba.get("pasos"), doc, base_dir, numerada=True)

    # Fila 4: maniobra | automatizable | entorno
    auto = prueba.get("automatizable") or ""
    etiqueta_valor(t.cell(4, 0), ETIQUETAS["maniobra"], unir(prueba.get("maniobra"), "/"))
    etiqueta_valor(t.cell(4, 1), ETIQUETAS["automatizable"], AUTOMATIZABLE_VISIBLE.get(auto, auto))
    etiqueta_valor(t.cell(4, 2), ETIQUETAS["entorno"], unir(prueba.get("entorno"), " / "))

    # Fila 5: requisitos (uno por línea) | tags
    c = t.cell(5, 0)
    p = c.paragraphs[0]
    compactar(p)
    p.add_run(f"{ETIQUETAS['requisitos']}:").bold = True
    for requisito in prueba.get("requisitos") or []:
        linea = c.add_paragraph()
        compactar(linea)
        escribir_inline(linea, str(requisito))
    etiqueta_valor(fusionar(5, 1, 2), ETIQUETAS["etiquetas"], unir(prueba.get("etiquetas"), ", "))

    # Fila 6: fecha | versión
    etiqueta_valor(t.cell(6, 0), ETIQUETAS["fecha"], prueba.get("fecha"))
    etiqueta_valor(fusionar(6, 1, 2), ETIQUETAS["version"], prueba.get("version"))

    filas_adicionales(t, prueba.get("campos_adicionales"))

    for imagen in prueba.get("imagenes", []):
        fila = t.add_row()
        celda = fila.cells[0].merge(fila.cells[-1])
        insertar_imagen(celda, base_dir, imagen)

    fila_adjuntos(t)

    # Espacio entre tablas
    sep = doc.add_paragraph()
    compactar(sep)
    return t


def tabla_requisito(doc, req, base_dir):
    """Tabla de un requisito con la plantilla de plantillas/tabla_requisito.docx (campos en esquema.py)."""
    t = doc.add_table(rows=6, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fijar_anchos(t, list(ANCHOS_REQUISITO_CM))

    def fusionar(fila, ini, fin):
        return t.cell(fila, ini).merge(t.cell(fila, fin)) if fin > ini else t.cell(fila, ini)

    # Fila 0: ID (dos columnas) | título (verde)
    for c, txt, cursiva in ((fusionar(0, 0, 1), req.get("id", ""), False), (t.cell(0, 2), req.get("titulo", ""), True)):
        sombrear(c, COLOR_CABECERA)
        p = c.paragraphs[0]
        compactar(p)
        r = p.add_run(str(txt or ""))
        r.bold, r.italic = True, cursiva
        r.font.color.rgb = COLOR_TEXTO_CABECERA

    # Filas 1-3: descripción, fuente y referencia
    for fila, clave in ((1, "descripcion"), (2, "fuente"), (3, "referencia")):
        etiqueta_valor(fusionar(fila, 0, 2), REQ_ETIQUETAS[clave], req.get(clave))
    escribir_detalle(t.cell(1, 0), req.get("detalle"), base_dir, sum(ANCHOS_REQUISITO_CM) - 0.6)

    # Fila 4: tipo | criterio de validación
    etiqueta_valor(t.cell(4, 0), REQ_ETIQUETAS["tipo"], req.get("tipo"))
    etiqueta_valor(fusionar(4, 1, 2), REQ_ETIQUETAS["criterio_validacion"], unir(req.get("criterio_validacion"), " / "))

    # Fila 5: fase | fecha | versión
    for col, clave in enumerate(("fase", "fecha", "version")):
        etiqueta_valor(t.cell(5, col), REQ_ETIQUETAS[clave], req.get(clave))

    filas_adicionales(t, req.get("campos_adicionales"))

    sep = doc.add_paragraph()
    compactar(sep)
    return t


def tabla_informativa(doc, tabla, base_dir, ancho_cm=ANCHO_TABLA_CM):
    """Tabla que no es una prueba (historial, parámetros…), tal como venía en el Word.

    Cada fila es una lista de celdas; una celda es un texto (con marcas de
    formato y saltos de línea) o {texto, imagenes, columnas}. `doc` puede ser
    el documento o una celda (tabla dentro de la descripción de un requisito).
    """
    filas = [f for f in (tabla.get("filas") or []) if isinstance(f, list)]
    ancho = lambda celda: max(int(celda.get("columnas", 1)), 1) if isinstance(celda, dict) else 1
    columnas = max([sum(ancho(c) for c in f) for f in filas] or [1]) or 1
    t = doc.add_table(rows=len(filas) or 1, cols=columnas)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    fijar_anchos(t, [ancho_cm / columnas] * columnas)
    for i, fila in enumerate(filas):
        col = 0
        for j, celda in enumerate(fila):
            # La última celda ocupa lo que quede de la fila.
            fin = columnas - 1 if j == len(fila) - 1 else min(col + ancho(celda), columnas) - 1
            if col > columnas - 1:
                break
            destino = t.cell(i, col).merge(t.cell(i, fin)) if fin > col else t.cell(i, col)
            texto = celda.get("texto", "") if isinstance(celda, dict) else str(celda or "")
            for k, linea in enumerate(texto.split("\n")):
                p = destino.paragraphs[0] if k == 0 else destino.add_paragraph()
                compactar(p)
                escribir_inline(p, linea)
            for imagen in (celda.get("imagenes") or []) if isinstance(celda, dict) else []:
                insertar_imagen(destino, base_dir, imagen)
            col = fin + 1
    sep = doc.add_paragraph()
    compactar(sep)
    return t


# ----------------------------------------------------------------------
# Títulos con numeración automática (1 / 1.1 / 1.1.1 ...), como en Word
# ----------------------------------------------------------------------
NIVELES_TITULO = 6


def numerar_titulos(doc, primer_capitulo=1):
    numbering = doc.part.numbering_part.numbering_definitions._numbering
    ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    abstract_id = max(ids + [0]) + 1
    niveles = ""
    for i in range(NIVELES_TITULO):
        texto = ".".join(f"%{k + 1}" for k in range(i + 1))
        inicio = primer_capitulo if i == 0 else 1
        niveles += (
            f'<w:lvl w:ilvl="{i}"><w:start w:val="{inicio}"/><w:numFmt w:val="decimal"/>'
            f'<w:pStyle w:val="Heading{i + 1}"/><w:lvlText w:val="{texto}"/>'
            f'<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="864" w:hanging="864"/></w:pPr></w:lvl>')
    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
        f'<w:multiLevelType w:val="multilevel"/>{niveles}</w:abstractNum>')
    primer_num = numbering.find(qn("w:num"))
    if primer_num is not None:
        primer_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    num_id = numbering.add_num(abstract_id).numId
    for i in range(NIVELES_TITULO):
        estilo = doc.styles[f"Heading {i + 1}"]
        estilo.font.name = FUENTE
        estilo.font.color.rgb = RGBColor(0, 0, 0)
        numPr = estilo.element.get_or_add_pPr().get_or_add_numPr()
        numPr.get_or_add_ilvl().val = i
        numPr.get_or_add_numId().val = num_id


# ----------------------------------------------------------------------
# Lectura de YAML y documento
# ----------------------------------------------------------------------
def cargar(entrada):
    yaml = YAML(typ="safe")
    ruta = Path(entrada)
    if ruta.is_dir() and not (ruta / "indice.yaml").exists() and (ruta / "pruebas").is_dir():
        ruta = ruta / "pruebas"             # carpeta del proyecto: el documento de pruebas
    if ruta.is_file():
        archivos = [ruta]
    else:
        indice = ruta / "indice.yaml"
        if indice.exists():
            with open(indice, encoding="utf-8") as f:
                archivos = [ruta / e["archivo"] for e in yaml.load(f)]
        else:
            archivos = sorted(a for a in ruta.glob("*.yaml") if a.name not in NOT_CHAPTERS)
    capitulos = []
    for a in archivos:
        with open(a, encoding="utf-8") as f:
            d = yaml.load(f)
        if d and "contenido" in d:
            capitulos.append((d.get("capitulo"), d["contenido"], a.parent))
        else:
            print(f"  AVISO: {a.name} no tiene la clave 'contenido', se ignora")
    return capitulos


def preparar_documento(primer_capitulo):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    normal = doc.styles["Normal"]
    normal.font.name = FUENTE
    normal.font.size = Pt(TAMANO)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FUENTE)
    numerar_titulos(doc, primer_capitulo)
    return doc


def titulo(doc, texto, nivel):
    doc.add_heading(texto, level=min(max(int(nivel or 1), 1), NIVELES_TITULO))


def escribir_contexto(doc, textos, base_dir=Path(".")):
    for texto in textos or []:
        if isinstance(texto, dict) and isinstance(texto.get("tabla"), dict):
            tabla_informativa(doc, texto["tabla"], base_dir)
            continue
        p = doc.add_paragraph()
        compactar(p)
        if isinstance(texto, dict):
            escribir_inline(p, texto.get("texto", ""))
            imagenes = texto.get("imagenes", [])
            if texto.get("imagen"):
                imagenes = [*imagenes, texto["imagen"]]
            for imagen in imagenes:
                ruta = base_dir / imagen
                if ruta.exists():
                    # Cada imagen en su propio párrafo, debajo del texto (no pegada a su última línea).
                    if p.text or p.runs:
                        p = doc.add_paragraph()
                        compactar(p)
                        p.paragraph_format.space_after = Pt(6)
                    agregar_imagen(p.add_run(), ruta, Cm(ANCHO_IMAGEN_CM))
                else:
                    print(f"  AVISO: imagen no encontrada: {ruta}")
            insertar_adjuntos(doc, texto, base_dir)
        else:
            escribir_inline(p, str(texto))


def escribir_capitulos(doc, capitulos, al_escribir_prueba=None):
    """Escribe títulos, contexto y tablas de los capítulos; devuelve el número de pruebas (o requisitos).

    `al_escribir_prueba(doc, tabla, prueba, base_dir)`, si se indica, se llama
    tras cada tabla (por ejemplo, para añadir el resultado de una ejecución).
    """
    ids, total = set(), 0
    for cap, contenido, base_dir in capitulos:
        if cap:
            titulo(doc, cap, 1)
        for el in contenido:
            if "seccion" in el:
                titulo(doc, el["seccion"], el.get("nivel", 2))
                preparar_adjuntos(el, base_dir)
                escribir_contexto(doc, el.get("contexto"), base_dir)
                parrafos_adjuntos(doc)
            elif "prueba" in el:
                p = el["prueba"]
                if p.get("nombre"):
                    titulo(doc, p["nombre"], p.get("nivel") or 3)
                preparar_adjuntos(p, base_dir)
                escribir_contexto(doc, p.get("contexto"), base_dir)
                if p.get("id") in ids:
                    print(f"  AVISO: ID duplicado: {p.get('id')}")
                ids.add(p.get("id"))
                tabla = tabla_prueba(doc, p, base_dir)
                if al_escribir_prueba is not None:
                    al_escribir_prueba(doc, tabla, p, base_dir)
                total += 1
            elif "requisito" in el:
                r = el["requisito"]
                if r.get("nombre"):
                    titulo(doc, r["nombre"], r.get("nivel") or 3)
                escribir_contexto(doc, r.get("contexto"), base_dir)
                if r.get("id") in ids:
                    print(f"  AVISO: ID duplicado: {r.get('id')}")
                ids.add(r.get("id"))
                tabla_requisito(doc, r, base_dir)
                total += 1
            elif "tabla" in el:
                preparar_adjuntos(el, base_dir)
                tabla_informativa(doc, el["tabla"] or {}, base_dir)
                parrafos_adjuntos(doc)
            elif "texto" in el:
                preparar_adjuntos(el, base_dir)
                escribir_contexto(doc, [el["texto"]], base_dir)
                parrafos_adjuntos(doc)
    return total


def generar_docx(entrada, salida, primer_capitulo=1):
    """Genera el DOCX a partir de los YAML y devuelve (pruebas o requisitos, capítulos)."""
    capitulos = cargar(entrada)
    if not capitulos:
        raise ValueError("No se han encontrado YAML con la clave 'contenido'.")

    doc = preparar_documento(primer_capitulo)
    total = escribir_capitulos(doc, capitulos)
    # En las propiedades del documento (no en el texto): con qué versión se generó.
    doc.core_properties.comments = f"Generado con {firma()}"
    doc.save(salida)
    return total, len(capitulos)


def son_requisitos(capitulos):
    """Si los capítulos son de requisitos (para el texto del resumen)."""
    return any("requisito" in el for _c, contenido, _b in capitulos for el in contenido if isinstance(el, dict))


def main(argv=None, prog=None):
    ap = argparse.ArgumentParser(prog=prog, description="YAML de pruebas -> DOCX")
    ap.add_argument("entrada", help="carpeta con YAML (con indice.yaml) o un archivo YAML")
    ap.add_argument("-o", "--salida", default="pruebas.docx")
    ap.add_argument("--primer-capitulo", type=int, default=1,
                    help="número del primer Título 1 (p. ej. 8 si el original empieza en 8)")
    args = ap.parse_args(argv)

    try:
        total, n_capitulos = generar_docx(args.entrada, args.salida, args.primer_capitulo)
    except ValueError as error:
        raise SystemExit(str(error))
    nombre = "requisitos" if son_requisitos(cargar(args.entrada)) else "pruebas"
    print(f"{total} {nombre} en {n_capitulos} capítulos -> {args.salida}")


if __name__ == "__main__":
    main()
