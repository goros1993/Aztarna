#!/usr/bin/env python3
"""Informe de una ejecución de pruebas: DOCX (y PDF si hay con qué convertirlo).

Mismo aspecto que `yaml_a_docx.py`, con un resumen al principio y el
resultado de la ejecución debajo de la tabla de cada prueba.

Uso:
    python herramientas.py informe ruta/a/la/ejecución [-o informe.docx] [--pdf]
"""
from __future__ import annotations

import argparse
import os
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Pt, RGBColor

from app import ejecucion_core as ej
from app import firma
from app import git_pruebas
from app import yaml_a_docx as yd

COLOR_ESTADO = {"ok": "2B7A4B", "no_ok": "B3261E", "bloqueada": "9A5B00", "no_aplica": "6B7B85", "pendiente": "72838D"}
FONDO_ESTADO = {"ok": "DCEFE3", "no_ok": "F8DEDB", "bloqueada": "F7ECD2", "no_aplica": "E4E9EC", "pendiente": "F2F4F5"}
EXTENSIONES_IMAGEN = {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".emf"}


def _fecha(valor: Any) -> str:
    texto = str(valor or "")
    try:
        return datetime.fromisoformat(texto).strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return texto


def _parrafo(doc, texto: str = "", *, negrita: bool = False, tamaño: float | None = None,
             color: str | None = None, antes: float = 0, despues: float = 4):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(antes)
    p.paragraph_format.space_after = Pt(despues)
    if texto:
        run = p.add_run(texto)
        run.bold = negrita or None
        if tamaño:
            run.font.size = Pt(tamaño)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
    return p


def _apartado(doc, texto: str):
    """Título de apartado del resumen (sin numerar: los títulos numerados son los capítulos)."""
    return _parrafo(doc, texto, negrita=True, tamaño=13, antes=14, despues=6)


def _tabla(doc, cabecera: list[str], filas: list[list[str]], anchos: list[float],
           colores: list[str | None] | None = None):
    tabla = doc.add_table(rows=1, cols=len(cabecera))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    for celda, texto in zip(tabla.rows[0].cells, cabecera):
        yd.sombrear(celda, yd.COLOR_CABECERA)
        p = celda.paragraphs[0]
        yd.compactar(p)
        run = p.add_run(texto)
        run.bold = True
        run.font.color.rgb = yd.COLOR_TEXTO_CABECERA
    for n, fila in enumerate(filas):
        celdas = tabla.add_row().cells
        for celda, texto in zip(celdas, fila):
            p = celda.paragraphs[0]
            yd.compactar(p)
            yd.escribir_inline(p, str(texto))
        if colores and colores[n]:
            yd.sombrear(celdas[0], colores[n])
    yd.fijar_anchos(tabla, anchos)
    doc.add_paragraph()
    return tabla


# ----------------------------------------------------------------------
# Resumen inicial
# ----------------------------------------------------------------------
def escribir_resumen(doc, ejecucion: dict[str, Any], resumen: dict[str, Any]) -> None:
    datos = ejecucion["datos"]
    cerrada = datos.get("estado") == "cerrada"
    _parrafo(doc, "Informe de ejecución", negrita=True, tamaño=22, despues=2)
    _parrafo(doc, str(datos.get("nombre") or ""), tamaño=15, color="485862", despues=10)
    if not cerrada:
        _parrafo(doc, f"BORRADOR: la ejecución sigue abierta ({resumen['totales']['pendiente']} pruebas pendientes).",
                 negrita=True, color=COLOR_ESTADO["no_ok"], despues=8)

    criterio = ej.texto_criterio(datos.get("seleccion"))
    filas = [
        ["Qué se prueba", datos.get("objeto", "")],
        ["Entorno de pruebas", datos.get("entorno_pruebas", "")],
        ["Banco / instalación", datos.get("entorno", "")],
        ["Responsable", datos.get("responsable", "")],
        ["Inicio", _fecha(datos.get("inicio"))],
        ["Fin", _fecha(datos.get("fin")) if cerrada else "(abierta)"],
        ["Criterio de selección", criterio],
        ["Versión de las pruebas", git_pruebas.texto_ejecucion(datos)],
        ["Pruebas", str(resumen["total"])],
        ["Informe generado", datetime.now().strftime("%d/%m/%Y %H:%M")],
    ]
    _tabla(doc, ["Ejecución", ""], [f for f in filas if str(f[1]).strip()], [5, 12])

    _apartado(doc, "Resumen")
    t = resumen["totales"]
    aplican = resumen["total"] - t["no_aplica"]
    def n(cantidad: int, uno: str, varios: str) -> str:
        return f"{cantidad} {uno if cantidad == 1 else varios}"
    frase = (f"{t['ok']} OK de {n(aplican, 'prueba que aplica', 'pruebas que aplican')} ({resumen['porcentaje_ok']} %). "
             f"{t['no_ok']} NO OK, {n(t['bloqueada'], 'bloqueada', 'bloqueadas')}, "
             f"{n(t['no_aplica'], 'no aplica', 'no aplican')}")
    if t["pendiente"]:
        frase += f" y {n(t['pendiente'], 'pendiente', 'pendientes')}"
    _parrafo(doc, frase + ".", despues=8)
    orden = ["ok", "no_ok", "bloqueada", "no_aplica", "pendiente"]
    total = resumen["total"] or 1
    _tabla(doc, ["Resultado", "Pruebas", "%"],
           [[ej.ESTADOS_VISIBLES[e], str(t[e]), f"{round(100 * t[e] / total, 1)} %"] for e in orden],
           [7, 5, 5], [FONDO_ESTADO[e] for e in orden])

    _apartado(doc, "Por capítulo")
    _tabla(doc, ["Capítulo", *[ej.ESTADOS_VISIBLES[e] for e in orden]],
           [[capitulo, *[str(v[e] or "·") for e in orden]] for capitulo, v in resumen["por_capitulo"].items()],
           [5.5, 2.3, 2.3, 2.3, 2.3, 2.3])

    problemas = [p for p in resumen["pruebas"] if p["estado"] in ("no_ok", "bloqueada")]
    if problemas:
        _apartado(doc, "Pruebas NO OK y bloqueadas")
        filas, colores = [], []
        for p in problemas:
            r = ejecucion["resultados"].get(p["id"]) or {}
            filas.append([p["id"], p["titulo"], ej.ESTADOS_VISIBLES[p["estado"]],
                          str(r.get("observaciones") or ""), str(r.get("incidencia") or "")])
            colores.append(None)
        tabla = _tabla(doc, ["ID", "Prueba", "Resultado", "Observaciones", "Incidencia"], filas, [3, 4.2, 2.2, 5.3, 2.3], colores)
        for fila, p in zip(tabla.rows[1:], problemas):
            yd.sombrear(fila.cells[2], FONDO_ESTADO[p["estado"]])

    if resumen["problemas"]:
        _apartado(doc, "Avisos")
        for aviso in resumen["problemas"]:
            p = doc.add_paragraph(style="List Bullet")
            if aviso["id"]:
                p.add_run(f"{aviso['id']}: ").bold = True
            p.add_run(aviso["mensaje"])


# ----------------------------------------------------------------------
# Resultado debajo de cada prueba
# ----------------------------------------------------------------------
def _fila(tabla):
    fila = tabla.add_row()
    celda = fila.cells[0].merge(fila.cells[-1])
    p = celda.paragraphs[0]
    yd.compactar(p)
    return celda, p


def _etiqueta(p, etiqueta: str, valor: str) -> None:
    p.add_run(f"{etiqueta}: ").bold = True
    yd.escribir_inline(p, valor)


def escribir_resultado(tabla, ejecucion: dict[str, Any], prueba: dict[str, Any]) -> None:
    test_id = str(prueba.get("id") or "").strip()
    resultado = ejecucion["resultados"].get(test_id) or {}
    estado = str(resultado.get("estado") or "pendiente")

    celda, p = _fila(tabla)
    yd.sombrear(celda, FONDO_ESTADO.get(estado, "FFFFFF"))
    p.add_run("Resultado de la ejecución: ").bold = True
    run = p.add_run(ej.ESTADOS_VISIBLES.get(estado, estado))
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(COLOR_ESTADO.get(estado, "000000"))
    detalles = [_fecha(resultado.get("fecha")), str(resultado.get("ejecutor") or "")]
    if any(detalles):
        p.add_run("   (" + " · ".join(d for d in detalles if d) + ")")

    if resultado.get("huella") and resultado["huella"] != ej.huella(prueba):
        _, p = _fila(tabla)
        run = p.add_run("Aviso: la prueba se modificó después de ejecutarse; el texto de arriba es el actual.")
        run.italic = True
    if str(resultado.get("observaciones") or "").strip():
        _, p = _fila(tabla)
        _etiqueta(p, "Observaciones", str(resultado["observaciones"]))
    if str(resultado.get("incidencia") or "").strip():
        _, p = _fila(tabla)
        _etiqueta(p, "Incidencia", str(resultado["incidencia"]))

    evidencias = [str(e) for e in resultado.get("evidencias") or []]
    if evidencias:
        celda, p = _fila(tabla)
        p.add_run("Evidencias:").bold = True
        otras = []
        for ruta in evidencias:
            if Path(ruta).suffix.lower() in EXTENSIONES_IMAGEN:
                yd.insertar_imagen(celda, ejecucion["carpeta"], ruta)
            else:
                otras.append(Path(ruta).name)
        if otras:
            q = celda.add_paragraph()
            yd.compactar(q)
            q.add_run("Archivos adjuntos en la carpeta de la ejecución: " + ", ".join(otras))

    intentos = list(resultado.get("intentos") or [])
    if intentos:
        celda, p = _fila(tabla)
        p.add_run(f"Intentos anteriores ({len(intentos)}):").bold = True
        for intento in intentos:
            q = celda.add_paragraph(style="List Bullet")
            yd.compactar(q)
            cabecera = [ej.ESTADOS_VISIBLES.get(str(intento.get("estado")), str(intento.get("estado"))),
                        _fecha(intento.get("fecha")), str(intento.get("ejecutor") or ""), str(intento.get("incidencia") or "")]
            q.add_run(" · ".join(c for c in cabecera if c)).bold = True
            if str(intento.get("observaciones") or "").strip():
                q.add_run(": ")
                yd.escribir_inline(q, str(intento["observaciones"]))


# ----------------------------------------------------------------------
# Documento completo
# ----------------------------------------------------------------------
def generar_informe(ejecucion: dict[str, Any], salida: str | Path,
                    capitulos: list[dict[str, Any]] | None = None) -> dict[str, Any]:
    """Genera el DOCX. Usa las pruebas congeladas si la ejecución está cerrada."""
    capitulos = capitulos if capitulos is not None else ej.capitulos_de(ejecucion)
    resumen = ej.resumen(ejecucion, capitulos)
    ids = set(ej.ids_de(ejecucion))
    salida = Path(salida)
    if salida.suffix.lower() != ".docx":
        salida = salida.with_suffix(".docx")
    salida.parent.mkdir(parents=True, exist_ok=True)

    doc = yd.preparar_documento(1)
    seccion = doc.sections[0]
    seccion.different_first_page_header_footer = False
    escribir_resumen(doc, ejecucion, resumen)
    doc.add_page_break()
    contenido = [
        (c["data"].get("capitulo"), ej.filtrar_contenido(c["data"]["contenido"], ids), c["path"].parent)
        for c in capitulos
    ]
    total = yd.escribir_capitulos(
        doc, [c for c in contenido if c[1]],
        lambda _doc, tabla, prueba, _base: escribir_resultado(tabla, ejecucion, prueba),
    )
    pie = seccion.footer.paragraphs[0]
    pie.text = f"{ejecucion['datos'].get('nombre', '')} · informe de ejecución · generado con {firma()}"
    for run in pie.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("72838D")
    doc.save(salida)
    return {"archivo": str(salida), "pruebas": total, "resumen": resumen}


# ----------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------
def _libreoffice() -> str | None:
    for nombre in ("soffice", "libreoffice"):
        ruta = shutil.which(nombre)
        if ruta:
            return ruta
    candidatos: list[Path] = []
    if sys.platform.startswith("win"):
        candidatos += [Path("C:/Program Files/LibreOffice/program/soffice.exe"),
                       Path("C:/Program Files (x86)/LibreOffice/program/soffice.exe")]
    elif sys.platform == "darwin":
        candidatos.append(Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"))
    else:
        # El paquete oficial de libreoffice.org se instala en /opt/libreofficeX.Y
        # con un lanzador con la versión en el nombre (libreoffice26.8).
        candidatos += sorted(Path("/opt").glob("libreoffice*/program/soffice"), reverse=True)
        for carpeta in os.environ.get("PATH", "").split(os.pathsep):
            if carpeta and Path(carpeta).is_dir():
                candidatos += sorted(Path(carpeta).glob("libreoffice[0-9]*"), reverse=True)
    return next((str(c) for c in candidatos if c.is_file() and os.access(c, os.X_OK)), None)


def pdf_disponible() -> str | None:
    """Con qué se puede generar el PDF: «LibreOffice», «Word» o None."""
    if _libreoffice():
        return "LibreOffice"
    if sys.platform.startswith("win") and shutil.which("powershell"):
        return "Word"          # se comprueba de verdad al convertir
    return None


def convertir_a_pdf(docx: str | Path) -> Path:
    """Convierte el DOCX a PDF con LibreOffice o, en Windows, con Microsoft Word."""
    docx = Path(docx).resolve()
    pdf = docx.with_suffix(".pdf")
    soffice = _libreoffice()
    if soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(docx.parent), str(docx)],
                       check=True, timeout=300, capture_output=True)
    elif sys.platform.startswith("win") and shutil.which("powershell"):
        script = (
            "$ErrorActionPreference='Stop';"
            "$w=New-Object -ComObject Word.Application; $w.Visible=$false;"
            f"$d=$w.Documents.Open('{docx}'); $d.SaveAs([ref]'{pdf}',[ref]17); $d.Close(); $w.Quit()"
        )
        resultado = subprocess.run(["powershell", "-NoProfile", "-Command", script],
                                   timeout=300, capture_output=True, text=True)
        if resultado.returncode != 0:
            raise ValueError("No se pudo convertir a PDF con Word: " + (resultado.stderr.strip()[:300] or "error desconocido"))
    else:
        raise ValueError("Para generar el PDF hace falta LibreOffice (o Microsoft Word en Windows). El DOCX sí se ha generado.")
    if not pdf.is_file():
        raise ValueError("La conversión a PDF no ha generado el archivo.")
    return pdf


def main(argv=None, prog=None) -> None:
    ap = argparse.ArgumentParser(prog=prog, description="Informe de una ejecución de pruebas")
    ap.add_argument("ejecucion", help="carpeta de la ejecución (con ejecucion.yaml)")
    ap.add_argument("-o", "--salida", help="archivo DOCX de salida (por defecto, dentro de la ejecución)")
    ap.add_argument("--pdf", action="store_true", help="generar también el PDF")
    args = ap.parse_args(argv)
    try:
        ejecucion = ej.cargar(args.ejecucion)
        salida = args.salida or ejecucion["carpeta"] / f"Informe_{ejecucion['carpeta'].name}.docx"
        datos = generar_informe(ejecucion, salida)
        print(f"{datos['pruebas']} pruebas -> {datos['archivo']}")
        if args.pdf:
            print(f"PDF -> {convertir_a_pdf(datos['archivo'])}")
    except (ValueError, OSError, subprocess.SubprocessError) as error:
        raise SystemExit(str(error))


if __name__ == "__main__":
    main()
