#!/usr/bin/env python3
"""Importar los resultados de una ejecución desde un Word (⋯ → Importar desde Word… → Resultados).

El Word tiene las tablas de las pruebas con unas filas más con el resultado
(«Resultado: OK», «Observaciones: …», «Ejecutor: …», «Fecha ejecución: …»).
Las pruebas ya están en el proyecto: del Word solo se toman los resultados y
se enlazan por el ID de la cabecera de cada tabla. Con ellos se crea una
ejecución nueva (abierta) con esas pruebas.

Igual que al importar pruebas o requisitos (importacion.py), se eligen qué
familias de tablas tienen resultados y a qué dato va cada etiqueta; además,
qué resultado es cada valor escrito en el Word («Correcto» → OK, «N/A» → No
aplica…). El mapeo es:

    {"tipos": {firma_de_familia: "resultado" | "ignorar"},
     "campos": {"resultado": {etiqueta_normalizada: dato | "ignorar" | "anterior"}},
     "valores": {valor_normalizado: estado}}

Cada tabla tiene como mucho un resultado; si una prueba aparece en varias
tablas, vale la última (y se avisa).
"""
from __future__ import annotations

import re
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document

from app import ejecucion_core as ej
from app import importacion as imp
from migracion import convertir as cv

DATOS = {"estado": "Resultado", "observaciones": "Observaciones", "incidencia": "Incidencia",
         "ejecutor": "Ejecutor", "fecha": "Fecha de ejecución"}
ESPECIALES = {"anterior": imp.ESPECIALES["anterior"], "ignorar": imp.ESPECIALES["ignorar"]}
TIPOS_VISIBLES = {"resultado": "Tiene resultados", "ignorar": "No importar"}
SIN_RESULTADO = "pendiente"

SINONIMOS = {
    "estado": ["Resultado", "Resultado obtenido", "Resultado de la prueba", "Decisión", "Veredicto", "Estado", "OK/NOK"],
    "observaciones": ["Observaciones", "Comentarios", "Notas"],
    "incidencia": ["Incidencia", "Incidencias", "Defecto", "Bug", "Ticket", "Referencia incidencia"],
    "ejecutor": ["Ejecutor", "Ejecutor prueba", "Ejecutado por", "Probado por", "Realizado por"],
    "fecha": ["Fecha ejecución", "Fecha de ejecución", "Fecha de la prueba", "Fecha ensayo"],
}
# Etiquetas que son de la prueba, no del resultado («Fecha» es la de la especificación).
NO_SON_RESULTADO = ["Resultado esperado", "Criterio(s) de aceptación", "Criterio de aceptación", "Fecha", "Versión"]
VALORES = {
    "ok": ["OK", "Correcto", "Correcta", "Pasa", "Pasada", "Superado", "Superada", "Apto", "Conforme", "Válido",
           "Cumple", "Sí", "Si", "Pass", "Passed", "Bien"],
    "no_ok": ["NOK", "NO OK", "KO", "Fallo", "Falla", "Fallida", "Error", "Incorrecto", "Incorrecta", "No apto",
              "No conforme", "No cumple", "No pasa", "Fail", "Failed", "No"],
    "bloqueada": ["Bloqueada", "Bloqueado", "No ejecutable", "Blocked", "Bloqueo"],
    "no_aplica": ["N/A", "NA", "No aplica", "No procede", "No aplicable"],
    SIN_RESULTADO: ["Pendiente", "No ejecutada", "No ejecutado", "Sin ejecutar", "-", "—"],
}

_SINONIMOS = {imp.clave(s): dato for dato, lista in SINONIMOS.items() for s in lista}
_NO_SON = {imp.clave(e) for e in NO_SON_RESULTADO}
_VALORES = {imp._letras(v) or v: estado for estado, lista in VALORES.items() for v in lista}


def clave_valor(texto: str) -> str:
    """Valor normalizado («N/A» → «na», «No OK» → «nook»; «-» se queda como está)."""
    texto = " ".join(str(texto or "").split())
    return imp._letras(texto) or texto


def proponer_campo(etiqueta: str, solo_interior: bool = False) -> str:
    k = imp.clave(etiqueta)
    if k in _SINONIMOS:
        return _SINONIMOS[k]
    if solo_interior:
        return "anterior"
    return "ignorar"


def proponer_valor(texto: str) -> str | None:
    """Estado propuesto para un valor escrito en el Word (None si no se reconoce)."""
    k = clave_valor(texto)
    if not k:
        return SIN_RESULTADO
    if k in _VALORES:
        return _VALORES[k]
    primera = imp._letras(str(texto).split()[0]) if str(texto).split() else ""
    return _VALORES.get(primera)                            # «OK, sin incidencias» → OK


class _Lector(imp._Lector):
    """Reparte las celdas en trozos por etiqueta, con el mapeo de los resultados."""

    def __init__(self, campos: dict[str, str]) -> None:
        super().__init__("prueba", {}, None)
        self.campos_resultado = campos

    def destino(self, etiqueta: str, interior: bool) -> str:
        return self.campos_resultado.get(imp.clave(etiqueta)) or proponer_campo(etiqueta, interior)


def _texto(segmento: dict[str, Any]) -> str:
    lineas = []
    for tipo, b, desde in segmento["partes"]:
        if tipo == "t":
            lineas.extend(" | ".join(cv.texto_celda(c).strip() for c in cv.celdas_unicas(f)) for f in b.rows)
        else:
            lineas.append(cv.texto(b)[desde:].strip())
    return "\n".join(x for x in lineas if x).strip()


# ----------------------------------------------------------------------
# 1. Analizar
# ----------------------------------------------------------------------
def analizar(ruta: str | Path) -> dict[str, Any]:
    """Familias de tablas con la propuesta para resultados y los valores que aparecen en cada etiqueta."""
    doc = Document(str(ruta))
    analisis = imp.analizar_documento(doc)
    de_tabla = analisis.pop("familia_de_tabla")
    valores: dict[int, dict[str, Counter[str]]] = {}
    lector = _Lector({})
    for i, tabla in enumerate(doc.tables):
        por_etiqueta = valores.setdefault(de_tabla[i], {})
        for s in lector.segmentos(cv.filas_sin_repetir(tabla)):
            if s["etiqueta"]:
                texto = " ".join(_texto(s).split())
                if len(texto) <= 60:
                    por_etiqueta.setdefault(imp.clave(s["etiqueta"]), Counter())[texto] += 1
    for n, f in enumerate(analisis["familias"]):
        for e in f["etiquetas"]:
            e["propuesta"] = {"resultado": proponer_campo(e["etiqueta"], not e["inicio"])}
            e["valores"] = [{"valor": v, "tablas": c, "clave": clave_valor(v), "propuesta": proponer_valor(v)}
                            for v, c in (valores.get(n, {}).get(e["clave"]) or Counter()).most_common(40)]
        tiene = any(e["propuesta"]["resultado"] not in ESPECIALES for e in f["etiquetas"])
        f["tipo"] = "resultado" if tiene else "ignorar"
    analisis["archivo"] = str(ruta)
    return analisis


def mapeo_propuesto(analisis: dict[str, Any], recordado: dict[str, Any] | None = None) -> dict[str, Any]:
    recordado = recordado or {}
    campos_rec = dict((recordado.get("campos") or {}).get("resultado") or {})
    valores_rec, tipos_rec = (dict(recordado.get(k) or {}) for k in ("valores", "tipos"))
    mapeo: dict[str, Any] = {"tipos": {}, "campos": {"resultado": {}}, "valores": {}}
    for f in analisis["familias"]:
        mapeo["tipos"][f["firma"]] = tipos_rec.get(f["firma"], f["tipo"])
        for e in f["etiquetas"]:
            mapeo["campos"]["resultado"].setdefault(e["clave"], campos_rec.get(e["clave"], e["propuesta"]["resultado"]))
            for v in e["valores"]:
                propuesta = valores_rec.get(v["clave"], v["propuesta"])
                if propuesta:
                    mapeo["valores"].setdefault(v["clave"], propuesta)
    return mapeo


# ----------------------------------------------------------------------
# 2. Leer los resultados
# ----------------------------------------------------------------------
_FECHA = re.compile(r"^(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{2,4})(?:[ T,]+(\d{1,2})[:.](\d{2}))?$")
_FECHA_ISO = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T](\d{1,2}):(\d{2}))?")


def leer_fecha(texto: str) -> str | None:
    """«15/10/2026», «15-10-26 9:30», «2026-10-15» → «2026-10-15T09:30»; None si no se entiende."""
    texto = " ".join(str(texto or "").split())
    m = _FECHA.match(texto)
    try:
        if m:
            dia, mes, anio, hora, minuto = m.groups()
            anio = int(anio) + (2000 if len(anio) == 2 else 0)
            return datetime(anio, int(mes), int(dia), int(hora or 0), int(minuto or 0)).isoformat(timespec="minutes")
        m = _FECHA_ISO.match(texto)
        if m:
            anio, mes, dia, hora, minuto = m.groups()
            return datetime(int(anio), int(mes), int(dia), int(hora or 0), int(minuto or 0)).isoformat(timespec="minutes")
    except ValueError:
        return None
    return None


def leer(ruta: str | Path, mapeo: dict[str, Any]) -> dict[str, Any]:
    """Un resultado por tabla de las familias con resultados: {filas, avisos}."""
    doc = Document(str(ruta))
    analisis = imp.analizar_documento(doc)
    familias, de_tabla = analisis["familias"], analisis["familia_de_tabla"]
    tipos, valores = mapeo.get("tipos") or {}, mapeo.get("valores") or {}
    lector = _Lector(dict((mapeo.get("campos") or {}).get("resultado") or {}))
    filas, avisos = [], []
    for i, tabla in enumerate(doc.tables):
        familia = familias[de_tabla[i]]
        if tipos.get(familia["firma"], "ignorar") != "resultado":
            continue
        cab = cv.filas_sin_repetir(tabla)[0] if tabla.rows else []
        ident = " ".join(cv.texto_celda(cab[0]).split()) if cab else ""
        fila: dict[str, Any] = {"tabla": i + 1, "id": ident,
                                "titulo": " ".join(" ".join(cv.texto_celda(c).split()) for c in cab[1:]).strip(),
                                "estado": SIN_RESULTADO, "valor": "", "observaciones": [], "incidencia": "", "ejecutor": "", "fecha": ""}
        for s in lector.segmentos(cv.filas_sin_repetir(tabla)):
            dato = s["destino"] if s["destino"] in DATOS else None
            texto = _texto(s)
            if not dato or not texto:
                continue
            if dato == "estado":
                estado = valores.get(clave_valor(texto))
                if estado and estado != SIN_RESULTADO and fila["estado"] == SIN_RESULTADO:
                    fila["estado"], fila["valor"] = estado, texto
                elif not estado:
                    # Un texto que no es un resultado (p. ej., «Resultado obtenido: la cabina para…») va a observaciones.
                    fila["observaciones"].append(f"{s['etiqueta']}: {texto}")
                    fila["valor"] = fila["valor"] or texto
            elif dato == "observaciones":
                fila["observaciones"].append(texto)
            else:
                fila[dato] = f"{fila[dato]} {texto}".strip()
        if fila["fecha"] and not leer_fecha(fila["fecha"]):
            avisos.append(_aviso("fecha_no_entendida", f"fecha «{fila['fecha']}» que no se entiende: se usa la de la importación y se anota en observaciones", fila))
            fila["observaciones"].append(f"Fecha de ejecución: {fila['fecha']}")
            fila["fecha"] = ""
        if fila["estado"] == SIN_RESULTADO and fila["valor"]:
            avisos.append(_aviso("valor_sin_mapear", f"resultado «{fila['valor'][:60]}» que no corresponde a ningún estado: queda pendiente", fila))
        fila["observaciones"] = "\n".join(fila["observaciones"])
        filas.append(fila)
    return {"filas": filas, "avisos": avisos}


def _aviso(tipo: str, mensaje: str, fila: dict[str, Any]) -> cv.Aviso:
    aviso = cv.Aviso(tipo, mensaje, fila.get("id", ""))
    aviso.tabla = fila.get("tabla")
    return aviso


TIPOS_AVISO = {
    "no_en_proyecto": ("Pruebas del Word que no están en el proyecto (se saltan)", "aviso"),
    "repetida": ("Prueba en varias tablas (vale la última)", "aviso"),
    "valor_sin_mapear": ("Resultado que no corresponde a ningún estado", "aviso"),
    "fecha_no_entendida": ("Fecha de ejecución que no se entiende", "aviso"),
    "sin_id": ("Tabla sin ID", "error"),
}


def preparar(ruta: str | Path, mapeo: dict[str, Any], capitulos: list[dict[str, Any]]) -> dict[str, Any]:
    """Resultados que se importarán: uno por prueba del proyecto, con los avisos."""
    leido = leer(ruta, mapeo)
    en_proyecto = ej.indice_por_id(capitulos)
    avisos = list(leido["avisos"])
    por_id: dict[str, dict[str, Any]] = {}
    for fila in leido["filas"]:
        if not fila["id"]:
            avisos.append(_aviso("sin_id", "tabla de resultados sin ID en la cabecera: se salta", fila))
        elif fila["id"] not in en_proyecto:
            avisos.append(_aviso("no_en_proyecto", f"«{fila['titulo'][:50]}» no está en el proyecto: se salta", fila))
        else:
            if fila["id"] in por_id:
                avisos.append(_aviso("repetida", f"también en la tabla {por_id[fila['id']]['tabla']}: vale esta", fila))
            por_id[fila["id"]] = fila
    return {"resultados": list(por_id.values()), "avisos": avisos, "tablas": len(leido["filas"])}


def revision(preparado: dict[str, Any]) -> dict[str, Any]:
    resultados = preparado["resultados"]
    cuenta = Counter(r["estado"] for r in resultados)
    grupos: dict[str, dict[str, Any]] = {}
    for a in preparado["avisos"]:
        titulo, gravedad = TIPOS_AVISO.get(a.tipo, ("Otros", "aviso"))
        g = grupos.setdefault(a.tipo, {"tipo": a.tipo, "titulo": titulo, "gravedad": gravedad, "total": 0, "ejemplos": []})
        g["total"] += 1
        if len(g["ejemplos"]) < 12:
            g["ejemplos"].append({"donde": " · ".join(x for x in (f"tabla {a.tabla}" if a.tabla else "", a.id) if x), "mensaje": a.mensaje})
    nombres = {"estado": "Resultado", **{k: v for k, v in DATOS.items() if k != "estado"}}
    ejemplos = [{"id": r["id"], "titulo": r["titulo"],
                 "campos": [[nombres[k], ej.ESTADOS_VISIBLES.get(r[k], r[k]) if k == "estado" else str(r[k])]
                            for k in DATOS if r.get(k)]}
                for r in sorted(resultados, key=lambda r: r["estado"] == SIN_RESULTADO)[:3]]
    return {"tablas": preparado["tablas"], "pruebas": len(resultados), "por_estado": {e: cuenta[e] for e in ej.ESTADOS},
            "avisos": list(grupos.values()), "ejemplos": ejemplos}


# ----------------------------------------------------------------------
# 3. Crear la ejecución
# ----------------------------------------------------------------------
def importar(ruta: str | Path, mapeo: dict[str, Any], carpeta: str | Path, pruebas: str | Path,
             nombre: str, datos: dict[str, Any] | None = None) -> dict[str, Any]:
    """Crea en `carpeta` una ejecución con las pruebas del Word que están en el proyecto y apunta sus resultados."""
    from app.editor_core import read_chapters, tests_dir
    capitulos = read_chapters(tests_dir(Path(pruebas)))
    preparado = preparar(ruta, mapeo, capitulos)
    resultados = preparado["resultados"]
    if not resultados:
        raise ValueError("Ninguna tabla con resultados corresponde a una prueba del proyecto.")
    if not str((datos or {}).get("entorno_pruebas") or "").strip():
        raise ValueError("Indica en qué entorno de pruebas se ejecutó (Automático, Maqueta o Torre).")
    filtro = {"importado_de": Path(ruta).name, "solo_ids": [r["id"] for r in resultados]}
    ejecucion = ej.crear(carpeta, pruebas, nombre, filtro, datos or {})
    huellas = {p["id"]: p["huella"] for p in ejecucion["datos"]["pruebas"]}
    apuntados = 0
    for r in resultados:
        if r["estado"] == SIN_RESULTADO:
            continue
        ej.registrar(ejecucion, r["id"], r["estado"], observaciones=r["observaciones"], incidencia=r["incidencia"],
                     ejecutor=r["ejecutor"] or str((datos or {}).get("responsable") or ""), huella_actual=huellas.get(r["id"], ""),
                     fecha=leer_fecha(r["fecha"]) if r["fecha"] else None, origen="word", guardar_ahora=False)
        apuntados += 1
    ej.guardar(ejecucion)
    return {"ejecucion": ejecucion, "pruebas": len(resultados), "con_resultado": apuntados, "avisos": len(preparado["avisos"])}
