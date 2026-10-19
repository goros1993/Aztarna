#!/usr/bin/env python3
"""Importar pruebas o requisitos de un Excel (.xlsx): una fila por elemento y una columna por campo.

Lo usa el mismo asistente que el Word (app/importacion.py): `analizar_libro`
devuelve las hojas como «familias» y sus columnas como «etiquetas», y
`convertir_libro` da los capítulos en la misma forma que la conversión del
Word, así que la revisión (IDs, los que ya existen) y la importación son las
mismas.

Cada hoja:

- La cabecera es la primera fila con al menos dos celdas (en las 10 primeras):
  sus textos son los nombres de las columnas.
- Cada fila con ID o título es un elemento. Una fila sin ID ni título continúa
  el de arriba: sus listas (pasos, situación inicial…) se añaden a las suyas.
  Así se puede escribir un paso por fila.
- Las columnas que van a «Capítulo» y «Sección» ordenan los elementos; sin
  ellas, el capítulo es el nombre de la hoja. Una sección «A > B» es la
  subsección B de A.
- En una celda de lista, cada línea es un elemento («1.», «-» o «•» al
  principio se quitan).
"""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from app import esquema
from app import importacion as imp
from migracion import convertir as cv

# Columnas que no son un campo: ordenan los elementos o los identifican.
ESTRUCTURA = {"id": "ID", "titulo": "Título", "capitulo": "Capítulo", "seccion": "Sección"}
SINONIMOS_ESTRUCTURA = {
    "id": ["ID", "Identificador", "Código", "Id prueba", "ID prueba", "ID requisito", "Referencia interna"],
    "titulo": ["Título", "Nombre", "Caso de prueba", "Prueba", "Requisito", "Resumen"],
    "capitulo": ["Capítulo", "Módulo", "Área", "Bloque"],
    "seccion": ["Sección", "Apartado", "Grupo", "Subapartado", "Funcionalidad"],
}
_ESTRUCTURA = {imp.clave(s): campo for campo, lista in SINONIMOS_ESTRUCTURA.items() for s in lista}
# Columnas que no tienen sentido en un Excel (no hay párrafos dentro de una celda que sigan a otro campo).
ESPECIALES = {k: v for k, v in imp.ESPECIALES.items() if k != "anterior"}


def _libro(ruta: str | Path):
    from openpyxl import load_workbook
    try:
        return load_workbook(str(ruta), data_only=True)        # los valores, no las fórmulas
    except Exception as error:                                 # noqa: BLE001 - openpyxl lanza de todo con un archivo dañado
        raise ValueError(f"No se puede leer el Excel {Path(ruta).name}: {error}") from error


def texto_celda(valor: Any, campo: str = "") -> str:
    """Texto de una celda: fechas como dd/mm/aaaa y números sin «.0»."""
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        valor = valor.date() if not (valor.hour or valor.minute) else valor
    if isinstance(valor, date) and not isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y %H:%M")
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    if isinstance(valor, bool):
        return "Sí" if valor else "No"
    return str(valor).replace("\r\n", "\n").replace("\r", "\n").strip()


def proponer_campo(tipo: str, columna: str) -> str:
    """Campo propuesto para una columna: ID, título, capítulo, sección o el de la etiqueta del Word."""
    k = imp.clave(columna)
    if k in _ESTRUCTURA:
        return _ESTRUCTURA[k]
    return imp.proponer_campo(tipo, columna)


def _filas(hoja) -> list[list[Any]]:
    return [list(f) for f in hoja.iter_rows(values_only=True)]


def _cabecera(filas: list[list[Any]]) -> int | None:
    """Índice de la fila de cabecera: la primera (de las 10 primeras) con al menos dos celdas con texto."""
    for i, fila in enumerate(filas[:10]):
        if sum(1 for v in fila if texto_celda(v)) >= 2:
            return i
    return None


def _columnas(fila: list[Any]) -> list[tuple[int, str]]:
    return [(i, " ".join(texto_celda(v).split())) for i, v in enumerate(fila) if texto_celda(v)]


# ----------------------------------------------------------------------
# 1. Analizar
# ----------------------------------------------------------------------
def analizar_libro(ruta: str | Path, preferido: str | None = None) -> dict[str, Any]:
    """Las hojas del Excel como familias del asistente, con el tipo y los campos propuestos."""
    libro = _libro(ruta)
    familias = []
    for n, hoja in enumerate(libro.worksheets):
        filas = _filas(hoja)
        i = _cabecera(filas)
        columnas = _columnas(filas[i]) if i is not None else []
        datos = [f for f in filas[i + 1:]] if i is not None else []
        datos = [f for f in datos if any(texto_celda(v) for v in f)]
        etiquetas = []
        for col, nombre in columnas:
            llenas = sum(1 for f in datos if col < len(f) and texto_celda(f[col]))
            etiquetas.append({"clave": imp.clave(nombre), "etiqueta": nombre, "tablas": llenas, "inicio": True,
                              "propuesta": {t: proponer_campo(t, nombre) for t in imp.CLASES}})
        # Sin las columnas de estructura, el tipo se propone como con las etiquetas del Word.
        campos = {e["clave"]: {"etiqueta": e["etiqueta"], "inicio": True} for e in etiquetas if e["propuesta"]["prueba"] not in ESTRUCTURA}
        tipo = imp._proponer_tipo(campos, preferido) if datos else "ignorar"
        if tipo == "tabla":
            tipo = preferido if preferido in imp.CLASES else "ignorar"
        ident = next((c for c, nombre in columnas if proponer_campo("prueba", nombre) == "id"), None)
        ejemplos = [texto_celda(f[ident]) for f in datos if ident is not None and ident < len(f) and texto_celda(f[ident])][:3]
        familias.append({
            "numero": n + 1, "firma": f"hoja:{hoja.title}", "hoja": hoja.title, "tablas": len(datos),
            "posiciones": [], "ejemplos": ejemplos or ["(sin ID)"], "anidadas": 0, "tipo": tipo,
            "etiquetas": etiquetas, "imagenes": len(getattr(hoja, "_images", []) or []),
        })
    return {"tablas": sum(f["tablas"] for f in familias), "familias": familias, "archivo": str(ruta), "origen": "excel"}


def opciones(tipo: str) -> list[dict[str, str]]:
    """Campos a los que puede ir una columna, para el desplegable: los de estructura primero."""
    return ([{"valor": k, "texto": v} for k, v in ESTRUCTURA.items()]
            + [{"valor": c, "texto": imp.NOMBRES[tipo].get(c, c)} for c in imp.CLASES[tipo]]
            + [{"valor": k, "texto": v} for k, v in ESPECIALES.items()])


# ----------------------------------------------------------------------
# 2. Convertir
# ----------------------------------------------------------------------
_VIÑETA = re.compile(r"^\s*(?:\d{1,3}\s*[.)\-]|[-–•*·▪◦])\s+")


def _lineas(texto: str) -> list[str]:
    return [_VIÑETA.sub("", linea).strip() for linea in texto.split("\n") if linea.strip()]


def _aviso(tipo: str, mensaje: str, ident: str, hoja: str, fila: int) -> cv.Aviso:
    a = cv.Aviso(tipo, mensaje, ident)
    a.capitulo = f"hoja «{hoja}», fila {fila}"
    return a


def _leer_hoja(hoja, tipo: str, campos: dict[str, str], avisos: list) -> list[tuple[str, list[str], dict[str, Any], int]]:
    """(capítulo, ruta de secciones, elemento, fila) de cada elemento de la hoja, en orden."""
    filas = _filas(hoja)
    i = _cabecera(filas)
    if i is None:
        return []
    columnas = [(c, nombre, campos.get(imp.clave(nombre)) or proponer_campo(tipo, nombre)) for c, nombre in _columnas(filas[i])]
    clases = imp.CLASES[tipo]
    lector = imp._Lector(tipo, {}, None)
    leidos: list[tuple[str, list[str], dict[str, Any], int]] = []
    actual: dict[str, Any] | None = None
    capitulo, seccion = hoja.title, ""
    for n, fila in enumerate(filas[i + 1:], start=i + 2):
        valores = {c: texto_celda(fila[c]) if c < len(fila) else "" for c, _nombre, _campo in columnas}
        if not any(valores.values()):
            continue
        por_campo = {campo: valores[c] for c, _nombre, campo in columnas if campo in ESTRUCTURA}
        nuevo = bool(por_campo.get("id") or por_campo.get("titulo"))
        # Capítulo y sección vacíos siguen siendo los de arriba (como en una celda combinada).
        capitulo = por_campo.get("capitulo") or capitulo
        seccion = por_campo.get("seccion") or seccion
        if nuevo:
            actual = {"id": por_campo.get("id", ""), "titulo": " ".join(por_campo.get("titulo", "").split())}
            ruta = [x.strip() for x in re.split(r"\s*>\s*", seccion) if x.strip()]
            leidos.append((capitulo, ruta, actual, n))
        elif actual is None:
            avisos.append(_aviso("celda_sin_reconocer", "fila sin ID ni título antes del primer elemento: no se importa", "", hoja.title, n))
            continue
        for c, nombre, campo in columnas:
            valor = valores[c]
            if not valor or campo in ESTRUCTURA or campo == "ignorar":
                continue
            if campo == "adicional" or campo not in clases:
                extras = actual.setdefault("campos_adicionales", {})
                extras[nombre] = f"{extras[nombre]}\n{valor}" if extras.get(nombre) else valor
                continue
            clase = clases[campo]
            if clase == "lista":
                actual.setdefault(campo, []).extend(_lineas(valor))
            elif clase == "textos":
                actual.setdefault(campo, []).extend(v for v in esquema.lista_de_textos(valor, r"[\n,;]") if v not in actual.get(campo, []))
            elif clase == "parrafos":
                actual[campo] = f"{actual[campo]}\n{valor}" if actual.get(campo) else valor
            else:
                actual[campo] = f"{actual[campo]} {' '.join(valor.split())}".strip() if actual.get(campo) else " ".join(valor.split())
    # Campos que da la hoja: al actualizar uno que ya existe, solo cambian estos (ver importacion.COLUMNAS).
    presentes = {campo for _c, _n, campo in columnas if campo in clases or campo in ("id", "titulo")}
    if any(campo == "adicional" or (campo not in clases and campo not in ESTRUCTURA and campo != "ignorar") for _c, _n, campo in columnas):
        presentes.add("campos_adicionales")
    resultado = []
    for capitulo, ruta, d, n in leidos:
        antes = len(avisos)
        elemento = lector.normalizar(d, [], [], avisos)
        for a in avisos[antes:]:
            a.capitulo = f"hoja «{hoja.title}», fila {n}"
        if elemento.get("campos_adicionales"):
            presentes_fila = presentes | {"campos_adicionales"}   # p. ej., un valor que no es una opción
        else:
            presentes_fila = presentes
        elemento[imp.COLUMNAS] = sorted(presentes_fila)
        resultado.append((capitulo, ruta, elemento, n))
    return resultado


def convertir_libro(ruta: str | Path, mapeo: dict[str, Any]) -> dict[str, Any]:
    """Capítulos (con sus secciones y elementos), avisos y totales, como la conversión del Word."""
    libro = _libro(ruta)
    tipos = mapeo.get("tipos") or {}
    avisos: list = []
    capitulos: dict[str, dict[str, Any]] = {}
    totales = {t: 0 for t in imp.TIPOS}
    vistos: dict[tuple[str, str], tuple[str, int]] = {}
    for hoja in libro.worksheets:
        tipo = tipos.get(f"hoja:{hoja.title}", "ignorar")
        tipo = tipo if tipo in imp.TIPOS else "ignorar"
        if tipo not in imp.CLASES:
            totales["ignorar"] += 1
            continue
        if getattr(hoja, "_images", None):
            avisos.append(_aviso("imagen_excel", f"{len(hoja._images)} imagen(es) en la hoja: no se importan (añádelas en el editor)", "", hoja.title, 1))
        campos = dict((mapeo.get("campos") or {}).get(tipo) or {})
        for capitulo, ruta_secciones, elemento, fila in _leer_hoja(hoja, tipo, campos, avisos):
            totales[tipo] += 1
            ident = elemento.get("id", "")
            if not ident:
                avisos.append(_aviso("id_vacio", f"{'el requisito' if tipo == 'requisito' else 'la prueba'} «{elemento.get('titulo', '')}» no tiene ID", "", hoja.title, fila))
            elif (tipo, ident) in vistos:
                avisos.append(_aviso("id_duplicado", f"ID repetido (también en la hoja «{vistos[(tipo, ident)][0]}», fila {vistos[(tipo, ident)][1]})", ident, hoja.title, fila))
            vistos.setdefault((tipo, ident), (hoja.title, fila))
            c = capitulos.setdefault(imp.clave(capitulo) or capitulo, {"capitulo": capitulo, "sueltos": [], "secciones": {}})
            # Las secciones se crean la primera vez que aparecen; los elementos van al final de la suya.
            camino: tuple[str, ...] = ()
            for nivel, nombre in enumerate(ruta_secciones, start=2):
                camino = (*camino, imp.clave(nombre))
                c["secciones"].setdefault(camino, {"seccion": {"seccion": nombre, "nivel": nivel}, "elementos": []})
            item = {tipo: {"nombre": "", "nivel": len(ruta_secciones) + 2, **elemento}}
            (c["secciones"][camino]["elementos"] if camino else c["sueltos"]).append(item)
    resultado = [{"capitulo": c["capitulo"], "contenido": [*c["sueltos"], *_aplanar(c["secciones"], ())]}
                 for c in capitulos.values()]
    return {"capitulos": resultado, "avisos": avisos, "totales": totales}


def _aplanar(secciones: dict[tuple[str, ...], dict[str, Any]], madre: tuple[str, ...]) -> list[Any]:
    """Cada sección con sus elementos y, detrás, sus subsecciones (en el orden en que aparecen)."""
    salida: list[Any] = []
    for ruta, s in secciones.items():
        if len(ruta) == len(madre) + 1 and ruta[:len(madre)] == madre:
            salida += [s["seccion"], *s["elementos"], *_aplanar(secciones, ruta)]
    return salida


# ----------------------------------------------------------------------
# Plantilla
# ----------------------------------------------------------------------
def plantilla(ruta: str | Path, tipo: str = "prueba") -> Path:
    """Excel vacío con una columna por campo (y Capítulo y Sección), listo para rellenar e importar.

    Las columnas de opciones (maniobra, entorno, automatizable, tipo…) llevan
    una lista desplegable con las del proyecto abierto.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    ruta = Path(ruta)
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Pruebas" if tipo == "prueba" else "Requisitos"
    campos = ["capitulo", "seccion", "id", "titulo", *[c for c in imp.CLASES[tipo] if c != "contexto"]]
    nombres = {**ESTRUCTURA, **imp.NOMBRES[tipo]}
    hoja.append([nombres.get(c, c) for c in campos])
    listas: dict[str, list[str]] = {"automatizable": [esquema.AUTOMATIZABLE_VISIBLE[v] for v in esquema.AUTOMATIZABLE]} if tipo == "prueba" else {"tipo": list(esquema.REQ_TIPOS)}
    opciones_multiples = esquema.OPCIONES_MULTIPLES if tipo == "prueba" else esquema.REQ_OPCIONES_MULTIPLES
    for i, c in enumerate(campos, start=1):
        celda = hoja.cell(row=1, column=i)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F5F6B")
        celda.alignment = Alignment(vertical="top", wrap_text=True)
        letra = celda.column_letter
        hoja.column_dimensions[letra].width = 40 if imp.CLASES[tipo].get(c) in ("lista", "parrafos") or c == "titulo" else 16
        valores = listas.get(c) or (list(opciones_multiples[c]) if len(opciones_multiples.get(c, ())) else None)
        if valores and len(",".join(valores)) < 250:
            # En las de varias opciones se puede escribir más de una («Automático / Torre»): solo se avisa.
            dv = DataValidation(type="list", formula1='"' + ",".join(valores) + '"', allow_blank=True,
                                showErrorMessage=c not in opciones_multiples, errorStyle="warning")
            dv.add(f"{letra}2:{letra}1000")
            hoja.add_data_validation(dv)
    hoja.freeze_panes = "A2"
    ayuda = libro.create_sheet("Cómo se rellena")
    for linea in (
        "Una fila por " + ("prueba" if tipo == "prueba" else "requisito") + ". Las columnas se pueden quitar, reordenar o añadir (las que no conozca el editor se importan como campos adicionales).",
        "Capítulo y Sección ordenan las " + ("pruebas" if tipo == "prueba" else "requisitos") + "; sin Capítulo se usa el nombre de la hoja. «A > B» es la subsección B de A.",
        "En las columnas de lista (pasos, situación inicial…), una línea por elemento (Alt+Intro en la celda). También puedes dejar ID y Título vacíos: la fila continúa la de arriba.",
        "Al importar en un proyecto, lo que ya existe con el mismo ID se puede actualizar o dejar como está, y los capítulos y secciones con el mismo nombre reciben lo nuevo.",
    ):
        ayuda.append([linea])
    ayuda.column_dimensions["A"].width = 140
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(str(ruta))
    return ruta
