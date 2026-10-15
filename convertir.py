#!/usr/bin/env python3
"""
Word -> YAML. Un YAML por capítulo (cada Título 1) + indice.yaml.
Las imágenes se guardan en salida/yaml/imagenes y se referencian desde el YAML.

Cada tabla se clasifica antes de leerla: las pruebas pasan al esquema común
(`prueba:`, definido en esquema.py, que también dice cómo se convierten los
formatos antiguos) y las demás (historial de cambios, tablas de parámetros…) se
conservan tal cual como `tabla:` con sus filas, en el mismo sitio que el
texto que las rodea: en el contexto de la sección o de la prueba siguiente,
o en el capítulo si aún no hay ningún título. Los avisos se
agrupan por tipo en avisos.txt (y avisos.csv para filtrarlos en Excel).

Uso:
    python convertir.py tu_archivo.docx            -> carpeta salida/
    python convertir.py tu_archivo.docx -o otra

Requisitos:
    pip install python-docx ruamel.yaml
"""
import argparse
import csv
import mimetypes
import re
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.dml import MSO_COLOR_TYPE
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from ruamel.yaml import YAML

import esquema


def sin_tildes(s):
    s = unicodedata.normalize("NFD", s)
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def norm(s):
    return re.sub(r"\s+", "", sin_tildes(s)).lower()


def slug(s):
    return re.sub(r"\W+", "_", sin_tildes(s).lower()).strip("_")[:40] or "capitulo"


# Etiquetas del Word -> clave del YAML. Las del formato antiguo (componente,
# criticidad, situación final…) se leen y luego esquema.migrar_prueba las
# convierte o las quita.
CAMPOS = {"descripcion": "descripcion", "componente": "componente",
          "plataforma": "plataforma", "automatizable": "automatizable",
          "criticidad": "criticidad", "entornodepruebas": "entorno",
          "referencia": "referencia", "fecha": "fecha", "version": "version"}
# Títulos de lista dentro de una celda («Situación inicial:», «Ensayo:»…).
SECCIONES = {"parametrizacion": "parametrizacion", "situacioninicial": "situacion_inicial",
             "pasos": "pasos", "ensayo": "pasos", "situacionfinal": "situacion_final"}
# Campos que se guardan como lista: una línea por elemento (etiquetas: separadas por comas).
CAMPOS_EN_LISTA = ("parametrizacion", "situacion_inicial", "pasos", "situacion_final", "requisitos")

ETIQUETAS_CAMPOS = {
    "Descripción": "descripcion",
    "Componente": "componente",
    "Plataforma": "plataforma",
    "Automatizable": "automatizable",
    "Criticidad": "criticidad",
    "Entorno de pruebas": "entorno",
    "Referencia": "referencia",
    "Fecha": "fecha",
    "Versión": "version",
    "Parametrización": "parametrizacion",
    "Situación inicial": "situacion_inicial",
    "Condiciones iniciales": "situacion_inicial",
    "Pasos": "pasos",
    "Ensayo": "pasos",
    "ENSAYO(s) / Proceso(s)": "pasos",
    "Situación final": "situacion_final",
    "Entorno de prueba(s)": "entorno",
    "Entorno de pruebas(s)": "entorno",
    "Maniobra": "maniobra",
    "Maniobra(s)": "maniobra",
    "Requisitos": "requisitos",
    "Requisito(s)": "requisitos",
    "Tags": "etiquetas",
}
CAMPOS.update({norm(etiqueta): clave for etiqueta, clave in ETIQUETAS_CAMPOS.items()})

# Campos conocidos del formato antiguo que no tienen sitio en el esquema: se
# leen como campos adicionales y esquema.migrar_prueba los elimina.
ALIAS_ETIQUETAS_ADICIONALES = {
    "Resultado esperado": "Resultado esperado",
    "Criterio(s) de aceptación": "Criterio(s) de aceptación",
    "Criterio validación": "Criterio validación",
    "Ejecutor prueba": "Ejecutor prueba",
    "Fecha ejecución": "Fecha ejecución",
    "Resultado Obtenido": "Resultado Obtenido",
    "Resultado obtenido": "Resultado Obtenido",
    "Decisión": "Decisión",
    "Resultado": "Resultado",
    "Observaciones": "Observaciones",
}
ETIQUETAS_ADICIONALES = {
    norm(alias): etiqueta
    for alias, etiqueta in ALIAS_ETIQUETAS_ADICIONALES.items()
}

ETIQUETAS_IGNORADAS = {"ejecuciondeprueba"}

# Etiquetas que marcan el inicio de una lista de pasos o situaciones.
ETIQUETAS_LISTA = {n for n, clave in CAMPOS.items() if clave in SECCIONES.values()} | set(SECCIONES)
ID_PRUEBA = re.compile(r"^[A-Za-z][A-Za-z0-9]*(?:[_\-.][A-Za-z0-9]+)+$")


# ----------------------------------------------------------------------
# Avisos
# ----------------------------------------------------------------------
TIPOS_AVISO = {
    # tipo: (título en avisos.txt, gravedad)
    "id_vacio": ("Prueba sin ID", "error"),
    "id_duplicado": ("ID repetido", "error"),
    "id_no_valido": ("ID con espacios o con forma extraña", "aviso"),
    "titulo_vacio": ("Prueba sin título", "error"),
    "tabla_dudosa": ("Tabla dudosa: conservada como tabla, revisar si es una prueba", "aviso"),
    "sin_titulo": ("Prueba sin título previo", "aviso"),
    "valor_descartado": ("Valor que no encaja en el esquema (descartado)", "aviso"),
    "campos_eliminados": ("Datos de campos que ya no existen (eliminados)", "info"),
    "etiqueta_desconocida": ("Etiqueta desconocida (conservada en campos adicionales)", "aviso"),
    "texto_fuera_de_seccion": ("Texto fuera de Situación inicial / Pasos / Situación final", "aviso"),
    "celda_sin_reconocer": ("Celda sin etiqueta reconocida", "aviso"),
    "tabla_informativa": ("Tabla que no es una prueba (conservada tal cual)", "info"),
}
# Tipos que se repiten mucho: en avisos.txt se agrupan por mensaje (en avisos.csv van todos).
AGRUPAR_POR_MENSAJE = {"campos_eliminados", "valor_descartado", "etiqueta_desconocida", "tabla_informativa"}


class Aviso(str):
    """Texto del aviso («ID: mensaje») con su tipo y dónde está, para agruparlos."""

    def __new__(cls, tipo, mensaje, id_prueba=""):
        aviso = super().__new__(cls, f"{id_prueba}: {mensaje}" if id_prueba else mensaje)
        aviso.tipo, aviso.mensaje, aviso.id = tipo, mensaje, id_prueba
        aviso.capitulo, aviso.tabla = "", None
        return aviso


def escribir_avisos(avisos, carpeta):
    """avisos.txt agrupado por tipo (con resumen al principio) y avisos.csv."""
    grupos = {}
    for aviso in avisos:
        grupos.setdefault(getattr(aviso, "tipo", "otro"), []).append(aviso)
    orden = sorted(grupos, key=lambda t: (list(TIPOS_AVISO).index(t) if t in TIPOS_AVISO else len(TIPOS_AVISO)))
    lineas = [f"{len(avisos)} avisos", ""]
    for tipo in orden:
        titulo, gravedad = TIPOS_AVISO.get(tipo, ("Otros", "aviso"))
        lineas.append(f"  {len(grupos[tipo]):5d}  [{gravedad}] {titulo}")
    def donde(aviso, con_id=True):
        return " · ".join(x for x in (
            getattr(aviso, "capitulo", ""),
            f"tabla {aviso.tabla}" if getattr(aviso, "tabla", None) else "",
            getattr(aviso, "id", "") if con_id else "",
        ) if x)

    for tipo in orden:
        titulo, _ = TIPOS_AVISO.get(tipo, ("Otros", "aviso"))
        lineas += ["", f"== {titulo} ({len(grupos[tipo])}) =="]
        if tipo in AGRUPAR_POR_MENSAJE:
            por_mensaje = {}
            for aviso in grupos[tipo]:
                por_mensaje.setdefault(getattr(aviso, "mensaje", aviso), []).append(aviso)
            for mensaje, lista in sorted(por_mensaje.items(), key=lambda x: -len(x[1])):
                ejemplos = ", ".join(a.id or donde(a, False) for a in lista[:8] if a.id or donde(a, False))
                resto = f" y {len(lista) - 8} más" if len(lista) > 8 else ""
                lineas.append(f"  {len(lista):5d} × {mensaje}")
                if ejemplos:
                    lineas.append(f"          {ejemplos}{resto}")
            continue
        for aviso in grupos[tipo]:
            lugar = donde(aviso)
            lineas.append(f"  {lugar}: {getattr(aviso, 'mensaje', aviso)}" if lugar else f"  {aviso}")
    (carpeta / "avisos.txt").write_text("\n".join(lineas) + "\n", encoding="utf-8")
    with open(carpeta / "avisos.csv", "w", encoding="utf-8-sig", newline="") as f:
        escritor = csv.writer(f, delimiter=";")
        escritor.writerow(["gravedad", "tipo", "capitulo", "tabla", "id", "mensaje"])
        for aviso in avisos:
            tipo = getattr(aviso, "tipo", "otro")
            escritor.writerow([TIPOS_AVISO.get(tipo, ("", "aviso"))[1], tipo, getattr(aviso, "capitulo", ""),
                               getattr(aviso, "tabla", "") or "", getattr(aviso, "id", ""),
                               getattr(aviso, "mensaje", aviso)])


# Color de fuente (hex -> nombre) y resaltado (nombre de Word -> nombre)
COLORES = {"FF0000": "rojo", "0070C0": "azul", "00B050": "verde"}
RESALTADOS = {"YELLOW": "amarillo", "BRIGHT_GREEN": "verde", "TURQUOISE": "cian"}


# ----------------------------------------------------------------------
# Texto con formato
# ----------------------------------------------------------------------
def limpiar(l):
    return l.strip().lstrip("•·-–\t ").strip()


# Runs de un párrafo, también los que Word mete dentro de hipervínculos,
# inserciones con control de cambios, etiquetas inteligentes, controles de
# contenido y campos (p.runs y p.text de python-docx se saltan casi todos).
_RUNS = ("./w:r | ./w:hyperlink/w:r | ./w:ins/w:r | ./w:smartTag/w:r | ./w:sdt/w:sdtContent/w:r"
         " | ./w:fldSimple/w:r | ./w:ins/w:hyperlink/w:r | ./w:hyperlink/w:ins/w:r | ./w:customXml/w:r")


def runs(p):
    return [Run(r, p) for r in p._p.xpath(_RUNS)]


def texto(p):
    """Texto completo de un párrafo (sin formato)."""
    return "".join(r.text for r in runs(p))


def texto_celda(c):
    """Texto completo de una celda: sus párrafos, separados por saltos de línea."""
    return "\n".join(texto(p) for p in c.paragraphs)


def texto_md(p):
    out = []
    for r in runs(p):
        t = r.text
        if not t.strip():
            out.append(t)
            continue
        ini = t[:len(t) - len(t.lstrip())]
        fin = t[len(t.rstrip()):]
        core = t.strip()
        if r.bold:
            core = f"**{core}**"
        if r.italic:
            core = f"*{core}*"
        col = r.font.color
        if col is not None and col.type == MSO_COLOR_TYPE.RGB and col.rgb is not None:
            hexa = str(col.rgb)
            if hexa != "000000":
                core = f"[[{COLORES.get(hexa, '#' + hexa)}|{core}]]"
        if r.font.highlight_color is not None:
            n = r.font.highlight_color.name
            core = f"[[res:{RESALTADOS.get(n, n.lower())}|{core}]]"
        out.append(ini + core + fin)
    return "".join(out).replace("****", "").replace("** **", " ")


def imagenes_parrafo(p, guardar_imagen):
    if guardar_imagen is None:
        return []
    imagenes = []
    rel_ids = []
    for blip in p._p.xpath(".//a:blip"):
        rel_id = blip.get(qn("r:embed"))
        if rel_id and rel_id not in rel_ids:
            rel_ids.append(rel_id)
    for imagen in p._p.xpath(".//*[local-name()='imagedata']"):
        rel_id = imagen.get(qn("r:id"))
        if rel_id and rel_id not in rel_ids:
            rel_ids.append(rel_id)
    for rel_id in rel_ids:
        imagenes.append(guardar_imagen(p.part.related_parts[rel_id]))
    return imagenes


def item(p, guardar_imagen=None):
    txt = limpiar(texto_md(p))
    pPr = p._p.pPr
    nivel = 0
    if pPr is not None and pPr.numPr is not None and pPr.numPr.ilvl is not None:
        nivel = pPr.numPr.ilvl.val
    imagenes = imagenes_parrafo(p, guardar_imagen)
    if nivel == 0 and not imagenes:
        return txt
    resultado = {"texto": txt}
    if nivel:
        resultado["nivel"] = nivel
    if imagenes:
        resultado["imagenes"] = imagenes
    return resultado


# ----------------------------------------------------------------------
# Tabla de una prueba
# ----------------------------------------------------------------------
def celdas_unicas(row):
    vistas, out = set(), []
    for c in row.cells:
        if c._tc not in vistas:
            vistas.add(c._tc)
            out.append(c)
    return out


def filas_sin_repetir(tabla):
    """Filas con sus celdas, sin repetir las unidas en vertical (python-docx
    devuelve la celda de arriba en cada fila que ocupa)."""
    vistas, filas = set(), []
    for r in tabla.rows:
        fila = []
        for c in celdas_unicas(r):
            if c._tc not in vistas:
                vistas.add(c._tc)
                fila.append(c)
        filas.append(fila)
    return filas


def guardar_campo(d, etiqueta, valor, avisos):
    n = norm(etiqueta)
    clave = CAMPOS.get(n)
    if clave:
        if clave in CAMPOS_EN_LISTA:
            items = [limpiar(linea) for linea in valor.splitlines() if limpiar(linea)]
            d.setdefault(clave, []).extend(items)
            return
        if clave == "etiquetas":
            d.setdefault(clave, []).extend(e for e in esquema.lista_de_textos(valor, r"[\n,]") if e not in d[clave])
            return
        valor = " ".join(valor.split())
        if valor or clave not in d:            # entorno, maniobra y automatizable los normaliza el esquema
            d[clave] = valor
        return

    if n in ETIQUETAS_IGNORADAS:
        return

    campos = d.setdefault("campos_adicionales", {})
    etiqueta_visible = ETIQUETAS_ADICIONALES.get(n, etiqueta.strip())
    valor = "\n".join(linea.strip() for linea in valor.splitlines()).strip()
    if etiqueta_visible in campos:
        anterior = campos[etiqueta_visible]
        if not anterior:
            campos[etiqueta_visible] = valor
        elif valor:
            campos[etiqueta_visible] = f"{anterior}\n{valor}"
    else:
        campos[etiqueta_visible] = valor
    if n not in ETIQUETAS_ADICIONALES:
        avisos.append(Aviso("etiqueta_desconocida",
                            f"etiqueta desconocida '{' '.join(etiqueta.split())[:80]}' (preservada en campos_adicionales)", d["id"]))


def extraer_etiquetas(texto):
    conocidas = {
        **{etiqueta: clave for etiqueta, clave in ETIQUETAS_CAMPOS.items()},
        **ALIAS_ETIQUETAS_ADICIONALES,
        "EJECUCIÓN DE PRUEBA": None,
    }
    por_normalizada = {}
    for etiqueta in conocidas:
        por_normalizada.setdefault(norm(etiqueta), etiqueta)
    etiquetas = sorted(por_normalizada.values(), key=len, reverse=True)
    patron = re.compile(
        r"(?<!\w)(" + "|".join(re.escape(e) for e in etiquetas) + r")\s*:",
        re.IGNORECASE,
    )
    coincidencias = list(patron.finditer(texto))
    return [
        (m.group(1), texto[m.end():coincidencias[i + 1].start() if i + 1 < len(coincidencias) else None])
        for i, m in enumerate(coincidencias)
    ]


def parsear(tabla, avisos, guardar_imagen=None, migrar=True):
    """Lee la tabla de una prueba. Con `migrar=False` la devuelve tal como se
    leyó, sin pasarla al esquema (auditar.py compara las dos)."""
    filas = filas_sin_repetir(tabla)
    cab = filas[0]
    d = {"id": texto_celda(cab[0]).strip(),
         "titulo": texto_celda(cab[1]).strip() if len(cab) > 1 else ""}
    parrafos_con_imagenes = set()
    for fila in filas[1:]:
        for c in fila:
            if any(norm(texto(p)).rstrip(":") in SECCIONES for p in c.paragraphs):
                seccion = None
                for p in c.paragraphs:
                    l = texto(p)
                    if not l.strip():
                        if seccion:
                            contenido = item(p, guardar_imagen)
                            if isinstance(contenido, dict) and contenido.get("imagenes"):
                                d[seccion].append(contenido)
                                parrafos_con_imagenes.add(p._p)
                        continue
                    n = norm(l).rstrip(":")
                    if n in SECCIONES:
                        seccion = SECCIONES[n]
                        d.setdefault(seccion, [])
                        imagenes = imagenes_parrafo(p, guardar_imagen)
                        if imagenes:
                            d[seccion].append({"texto": "", "imagenes": imagenes})
                            parrafos_con_imagenes.add(p._p)
                    elif seccion:
                        contenido = item(p, guardar_imagen)
                        d[seccion].append(contenido)
                        if isinstance(contenido, dict) and contenido.get("imagenes"):
                            parrafos_con_imagenes.add(p._p)
                    else:
                        avisos.append(Aviso("texto_fuera_de_seccion", f"texto fuera de sección: {l[:60]}", d["id"]))
            elif (campos := extraer_etiquetas(texto_celda(c))):
                for etiqueta, valor in campos:
                    guardar_campo(d, etiqueta, valor, avisos)
            elif ":" in texto_celda(c):
                etiqueta, valor = texto_celda(c).split(":", 1)
                guardar_campo(d, etiqueta, valor, avisos)
            elif texto_celda(c).strip():
                if norm(texto_celda(c).rstrip(":")) not in ETIQUETAS_IGNORADAS:
                    avisos.append(Aviso("celda_sin_reconocer", f"celda sin reconocer: {' '.join(texto_celda(c).split())[:60]}", d["id"]))
    imagenes_tabla = []
    for fila in filas:
        for celda in fila:
            for p in celda.paragraphs:
                if p._p not in parrafos_con_imagenes:
                    imagenes_tabla.extend(imagenes_parrafo(p, guardar_imagen))
    if imagenes_tabla:
        d["imagenes"] = list(dict.fromkeys(imagenes_tabla))
    if not migrar:
        return d
    # Al esquema común: situación final a Ensayo, fuera los campos que ya no
    # existen, entorno y maniobra como opciones… y todos los campos, aunque vacíos.
    for tipo, mensaje in esquema.migrar_prueba(d):
        avisos.append(Aviso(tipo, mensaje, d["id"]))
    esquema.completar(d)
    if not d["titulo"]:
        avisos.append(Aviso("titulo_vacio", "la prueba no tiene título (segunda celda de la cabecera)", d["id"]))
    return {k: d[k] for k in sorted(d, key=lambda k: esquema.ORDEN.index(k) if k in esquema.ORDEN else len(esquema.ORDEN))}


# ----------------------------------------------------------------------
# Clasificación: ¿la tabla es una prueba?
# ----------------------------------------------------------------------
def etiquetas_tabla(tabla):
    """Etiquetas reconocibles del cuerpo de la tabla: {normalizada: texto original}.

    Se buscan como lo hace `parsear`: párrafos que abren una sección
    (Pasos, Situación inicial…) y «Etiqueta:» conocidas dentro de las celdas.
    """
    filas = filas_sin_repetir(tabla)
    encontradas = {}
    for fila in filas[1:]:
        for c in fila:
            for p in c.paragraphs:
                n = norm(texto(p)).rstrip(":")
                if n in ETIQUETAS_LISTA:
                    encontradas[n] = texto(p).strip().rstrip(":")
            contenido = texto_celda(c)
            for etiqueta, _valor in extraer_etiquetas(contenido):
                encontradas.setdefault(norm(etiqueta), etiqueta)
            if ":" in contenido and not extraer_etiquetas(contenido):
                crudo = contenido.split(":", 1)[0].strip()
                if crudo and len(crudo) <= 40:
                    encontradas.setdefault(norm(crudo), crudo)
    return encontradas


def clasificar_tabla(tabla):
    """Devuelve (clase, motivo): clase es "prueba", "dudosa" o "informativa".

    - prueba: tiene Situación inicial / Pasos / Situación final (o sus
      equivalentes), o al menos 3 campos conocidos de una prueba;
    - dudosa: 1 o 2 campos conocidos, o una cabecera con forma de ID sin
      campos; se conserva como tabla y se avisa para revisarla;
    - informativa: ningún campo de prueba (historial de cambios, parámetros…).
    """
    filas = [celdas_unicas(r) for r in tabla.rows]
    if not filas or not any(texto_celda(c).strip() for fila in filas for c in fila):
        return "informativa", "tabla vacía"
    etiquetas = etiquetas_tabla(tabla)
    conocidas = sorted(n for n in etiquetas if n in CAMPOS or n in SECCIONES or n in ETIQUETAS_ADICIONALES)
    if any(n in ETIQUETAS_LISTA for n in conocidas):
        return "prueba", "tiene pasos o situaciones"
    if len(conocidas) >= 3:
        return "prueba", f"{len(conocidas)} campos de prueba"
    cabecera = texto_celda(filas[0][0]).strip() if filas[0] else ""
    if conocidas:
        nombres = ", ".join(etiquetas[n] for n in conocidas)
        return "dudosa", f"solo {len(conocidas)} campo(s) de prueba ({nombres})"
    if ID_PRUEBA.match(cabecera):
        return "dudosa", f"la cabecera «{cabecera}» parece un ID, pero no tiene campos de prueba"
    return "informativa", "sin campos de prueba"


def celda_tabla(celda, guardar_imagen):
    """Contenido de una celda de tabla informativa: texto con formato e imágenes."""
    textos, imagenes = [], []
    for p in celda.paragraphs:
        textos.append(texto_md(p).strip())
        imagenes.extend(imagenes_parrafo(p, guardar_imagen))
    texto = "\n".join(textos).strip()
    columnas = celda._tc.grid_span
    if not imagenes and columnas == 1:
        return texto
    contenido = {"texto": texto}
    if imagenes:
        contenido["imagenes"] = list(dict.fromkeys(imagenes))
    if columnas > 1:
        contenido["columnas"] = columnas
    return contenido


def tabla_informativa(tabla, motivo, guardar_imagen=None):
    """La tabla tal cual: una lista de filas, cada una con sus celdas."""
    filas = [[celda_tabla(c, guardar_imagen) for c in celdas_unicas(r)] for r in tabla.rows]
    return {"filas": filas, "motivo": motivo}


# ----------------------------------------------------------------------
# Estructura del documento
# ----------------------------------------------------------------------
def nivel_titulo(p):
    m = re.match(r"^(?:heading|t[ií]tulo)\s*(\d+)$", p.style.name.strip(), re.I)
    return int(m.group(1)) if m else None


def crear_guardador_imagenes(directorio):
    directorio = Path(directorio)
    guardadas = {}

    def guardar(parte):
        nombre_parte = str(parte.partname)
        if nombre_parte not in guardadas:
            extension = Path(nombre_parte).suffix
            if not extension:
                extension = mimetypes.guess_extension(parte.content_type) or ".bin"
            nombre = f"imagen_{len(guardadas) + 1:03d}{extension.lower()}"
            ruta = directorio / nombre
            directorio.mkdir(parents=True, exist_ok=True)
            ruta.write_bytes(parte.blob)
            guardadas[nombre_parte] = f"imagenes/{nombre}"
        return guardadas[nombre_parte]

    return guardar


def contenido_parrafo(p, guardar_imagen):
    texto = texto_md(p).strip()
    imagenes = imagenes_parrafo(p, guardar_imagen)
    if imagenes:
        return {"texto": texto, "imagenes": imagenes}
    return texto


def extraer_capitulos(doc, avisos, guardar_imagen=None):
    capitulos = [{"capitulo": None, "contenido": []}]
    pendiente = None
    texto_pendiente = []
    grupos = []
    numero_tabla = 0
    ids_vistos = {}

    def volcar():
        nonlocal pendiente, texto_pendiente
        if pendiente:
            entrada = {
                "seccion": pendiente[0],
                "nivel": pendiente[1],
            }
            if texto_pendiente:
                entrada["contexto"] = texto_pendiente
            capitulos[-1]["contenido"].append(entrada)
            grupos.append({"nivel": pendiente[1], "entrada": entrada})
            pendiente = None
            texto_pendiente = []

    def guardar_texto(texto):
        if grupos:
            grupos[-1]["entrada"].setdefault("contexto", []).append(texto)
        else:
            capitulos[-1]["contenido"].append({"texto": texto, "alcance": "capitulo"})

    for el in doc.element.body.iterchildren():
        if el.tag.endswith("}p"):
            p = Paragraph(el, doc)
            parrafo = contenido_parrafo(p, guardar_imagen)
            tiene_texto = parrafo.get("texto", "").strip() if isinstance(parrafo, dict) else parrafo
            if not tiene_texto and not isinstance(parrafo, dict):
                continue
            nivel = nivel_titulo(p)
            if nivel == 1:
                volcar()
                capitulos.append({"capitulo": texto(p).strip(), "contenido": []})
                if isinstance(parrafo, dict) and parrafo.get("imagenes"):
                    capitulos[-1]["contenido"].append({
                        "texto": {"texto": "", "imagenes": parrafo["imagenes"]}
                    })
                grupos = []
            elif nivel:
                volcar()
                grupos = [g for g in grupos if g["nivel"] < nivel]
                pendiente = (texto(p).strip(), nivel)
                if isinstance(parrafo, dict) and parrafo.get("imagenes"):
                    texto_pendiente.append(parrafo)
            elif pendiente:
                texto_pendiente.append(parrafo)
            else:
                guardar_texto(parrafo)
        elif el.tag.endswith("}tbl"):
            numero_tabla += 1
            tabla = Table(el, doc)
            antes = len(avisos)
            clase, motivo = clasificar_tabla(tabla)
            if clase != "prueba":
                if clase == "dudosa":
                    cabecera = " ".join(texto_celda(tabla.rows[0].cells[0]).split())[:60] if tabla.rows else ""
                    avisos.append(Aviso("tabla_dudosa", f"«{cabecera}»: {motivo}"))
                else:
                    avisos.append(Aviso("tabla_informativa", motivo))
                # La tabla va donde iría un párrafo en ese punto: contexto del
                # título pendiente, de la sección actual o, si no hay, del capítulo.
                elemento = {"tabla": tabla_informativa(tabla, motivo, guardar_imagen)}
                if pendiente:
                    texto_pendiente.append(elemento)
                elif grupos:
                    grupos[-1]["entrada"].setdefault("contexto", []).append(elemento)
                else:
                    capitulos[-1]["contenido"].append(elemento)
            else:
                prueba = parsear(tabla, avisos, guardar_imagen)
                if pendiente:
                    cab = {"nombre": pendiente[0], "nivel": pendiente[1]}
                    if texto_pendiente:
                        prueba["contexto"] = texto_pendiente
                    pendiente = None
                    texto_pendiente = []
                else:
                    cab = {"nombre": "", "nivel": None}
                    avisos.append(Aviso("sin_titulo", "tabla sin título previo", prueba["id"]))
                ident = prueba["id"]
                if not ident:
                    avisos.append(Aviso("id_vacio", "la prueba no tiene ID (primera celda de la tabla vacía)"))
                elif not ID_PRUEBA.match(ident):
                    avisos.append(Aviso("id_no_valido", f"ID «{ident}» con espacios o caracteres extraños", ident))
                if ident and ident in ids_vistos:
                    avisos.append(Aviso("id_duplicado", f"ID repetido (también en la tabla {ids_vistos[ident]})", ident))
                ids_vistos.setdefault(ident, numero_tabla)
                capitulos[-1]["contenido"].append({"prueba": {**cab, **prueba}})
            for aviso in avisos[antes:]:
                if isinstance(aviso, Aviso):
                    aviso.capitulo = capitulos[-1]["capitulo"] or ""
                    aviso.tabla = numero_tabla
    volcar()

    if not capitulos[0]["contenido"]:
        capitulos.pop(0)
    return capitulos


def main():
    ap = argparse.ArgumentParser(description="Word -> YAML por capítulos")
    ap.add_argument("docx")
    ap.add_argument("-o", "--salida", default="salida")
    args = ap.parse_args()

    doc = Document(args.docx)
    avisos = []
    out = Path(args.salida) / "yaml"
    out.mkdir(parents=True, exist_ok=True)
    guardar_imagen = crear_guardador_imagenes(out / "imagenes")
    capitulos = extraer_capitulos(doc, avisos, guardar_imagen)

    yaml = YAML()
    yaml.allow_unicode = True
    yaml.width = 4096
    indice = []
    for i, cap in enumerate(capitulos, 1):
        archivo = f"{i:02d}_{slug(cap['capitulo'] or 'sin_capitulo')}.yaml"
        with open(out / archivo, "w", encoding="utf-8") as f:
            yaml.dump(cap, f)
        indice.append({"capitulo": cap["capitulo"], "archivo": archivo})
    with open(out / "indice.yaml", "w", encoding="utf-8") as f:
        yaml.dump(indice, f)
    escribir_avisos(avisos, Path(args.salida))

    n = sum(1 for c in capitulos for e in c["contenido"] if "prueba" in e)
    def tablas_en(elemento):
        contexto = (elemento.get("prueba") or elemento).get("contexto") or []
        return ("tabla" in elemento) + sum(1 for i in contexto if isinstance(i, dict) and "tabla" in i)
    tablas = sum(tablas_en(e) for c in capitulos for e in c["contenido"] if isinstance(e, dict))
    dudosas = sum(1 for a in avisos if getattr(a, "tipo", "") == "tabla_dudosa")
    print(f"{n} pruebas, {tablas} tablas que no son pruebas ({dudosas} dudosas), "
          f"{len(capitulos)} capítulos, {len(avisos)} avisos (ver {Path(args.salida) / 'avisos.txt'})")


if __name__ == "__main__":
    main()
