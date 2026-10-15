"""Pruebas de los requisitos: trazabilidad, API del editor y Word de requisitos."""
import io
import shutil
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from app import editor_api, trazabilidad, yaml_a_docx
from app.editor_core import read_chapters

EJEMPLOS = Path(__file__).resolve().parent.parent / "ejemplos" / "yaml"

PRUEBAS = """\
capitulo: Pruebas
contenido:
- seccion: A
  nivel: 2
- prueba:
    nivel: 3
    id: T-1
    titulo: Uno
    requisitos: [R-1, R-X]
- prueba:
    nivel: 3
    id: T-2
    titulo: Dos
    requisitos: [R-1]
- prueba:
    nivel: 3
    id: T-3
    titulo: Tres
"""

REQUISITOS = """\
capitulo: Requisitos
contenido:
- seccion: Bloque 1
  nivel: 2
- requisito:
    nivel: 3
    id: R-1
    titulo: Cubierto
    descripcion: Texto original
    criterio_validacion: [Test]
- requisito:
    nivel: 3
    id: R-2
    titulo: Sin pruebas
    criterio_validacion: []
- seccion: Bloque 2
  nivel: 2
- requisito:
    nivel: 3
    id: R-3
    titulo: Por inspección
    criterio_validacion: [Inspección]
- requisito:
    nivel: 3
    id: R-4
    titulo: Informativo
    tipo: Informativo
"""


def ok(testcase, resultado):
    testcase.assertTrue(resultado["ok"], resultado.get("error"))
    return resultado["datos"]


class Base(unittest.TestCase):
    def setUp(self):
        self._tmp = TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.dir = Path(self._tmp.name) / "yaml"
        (self.dir / "requisitos").mkdir(parents=True)
        (self.dir / "01_pruebas.yaml").write_text(PRUEBAS, encoding="utf-8")
        (self.dir / "requisitos" / "01_req.yaml").write_text(REQUISITOS, encoding="utf-8")
        self.req_dir = self.dir / "requisitos"


class Calculo(Base):
    def datos(self, revisiones=None):
        return trazabilidad.calcular(read_chapters(self.dir), read_chapters(self.req_dir), revisiones)

    def test_estados_de_cobertura(self):
        d = self.datos()
        estados = {r["id"]: r["estado"] for r in d["requisitos"]}
        self.assertEqual(estados, {"R-1": "cubierto", "R-2": "sin_cobertura", "R-3": "no_aplica", "R-4": "no_aplica"})
        r1 = d["requisitos"][0]
        self.assertEqual([p["id"] for p in r1["pruebas"]], ["T-1", "T-2"])
        self.assertEqual(r1["seccion"], ["Bloque 1"])
        self.assertEqual(d["total"], {"requisitos": 4, "modificados": 0, "cubierto": 1, "sin_cobertura": 1, "no_aplica": 2})
        self.assertEqual(trazabilidad.porcentaje(d["total"]), 50)

    def test_huerfanas_y_pruebas_sin_requisito(self):
        d = self.datos()
        self.assertEqual([(h["id"], h["requisito"]) for h in d["huerfanas"]], [("T-1", "R-X")])
        self.assertEqual([p["id"] for p in d["sin_requisito"]], ["T-3"])

    def test_huella_ignora_metadatos(self):
        base = {"id": "R", "titulo": "T", "descripcion": "D", "version": "1"}
        self.assertEqual(trazabilidad.huella(base), trazabilidad.huella({**base, "version": "2", "fecha": "x"}))
        self.assertNotEqual(trazabilidad.huella(base), trazabilidad.huella({**base, "descripcion": "D2"}))

    def test_anotar_cambio_acumula_y_se_deshace(self):
        v1 = {"id": "R-1", "titulo": "T", "descripcion": "uno"}
        v2 = {**v1, "descripcion": "dos"}
        v3 = {**v1, "descripcion": "tres"}
        rev = {}
        self.assertTrue(trazabilidad.anotar_cambio(rev, v1, v2))
        self.assertEqual(rev["R-1"]["descripcion"], "uno")
        trazabilidad.anotar_cambio(rev, v2, v3)
        self.assertEqual(rev["R-1"]["descripcion"], "uno")        # se guarda el texto revisado, no el intermedio
        self.assertTrue(trazabilidad.anotar_cambio(rev, v3, v1))   # vuelve al texto revisado: ya no está cambiado
        self.assertEqual(rev, {})
        self.assertFalse(trazabilidad.anotar_cambio(rev, v1, {**v1, "version": "2"}))

    def test_anotar_cambio_sigue_al_id(self):
        rev = {}
        trazabilidad.anotar_cambio(rev, {"id": "A", "descripcion": "uno"}, {"id": "A", "descripcion": "dos"})
        trazabilidad.anotar_cambio(rev, {"id": "A", "descripcion": "dos"}, {"id": "B", "descripcion": "dos"})
        self.assertEqual(list(rev), ["B"])

    def test_modificado_solo_si_tiene_pruebas(self):
        rev = {"R-1": {"huella": "otra"}, "R-2": {"huella": "otra"}}
        d = {r["id"]: r for r in self.datos(rev)["requisitos"]}
        self.assertTrue(d["R-1"]["modificado"])
        self.assertFalse(d["R-2"]["modificado"])

    def test_revisiones_en_disco(self):
        trazabilidad.guardar_revisiones(self.req_dir, {"R-1": {"huella": "abc", "titulo": "T"}})
        self.assertEqual(trazabilidad.leer_revisiones(self.req_dir), {"R-1": {"huella": "abc", "titulo": "T"}})
        self.assertEqual(len(read_chapters(self.req_dir)), 1)      # revisiones.yaml no es un capítulo
        trazabilidad.guardar_revisiones(self.req_dir, {})
        self.assertFalse((self.req_dir / "revisiones.yaml").exists())

    def test_matriz_y_orden_de_linea(self):
        salida = Path(self._tmp.name) / "matriz.csv"
        with redirect_stdout(io.StringIO()) as texto:
            trazabilidad.main([str(self.dir), "-o", str(salida)])
        self.assertIn("4 requisitos: 1 cubiertos, 1 sin cobertura, 2 no aplica (50 % cubiertos)", texto.getvalue())
        self.assertIn("T-1 -> R-X", texto.getvalue())
        filas = salida.read_text(encoding="utf-8-sig").splitlines()
        self.assertEqual(filas[0].split(";")[0], "requisito")
        self.assertIn("R-1;Cubierto;Requisitos;Bloque 1;;Test;Cubierto;;T-1, T-2", filas[1])
        self.assertIn("R-X;", filas[-1])

    def test_ejemplos(self):
        d = trazabilidad.calcular(read_chapters(EJEMPLOS), trazabilidad.leer_requisitos(EJEMPLOS),
                                  trazabilidad.leer_revisiones(EJEMPLOS / "requisitos"))
        self.assertEqual(d["total"]["sin_cobertura"], 3)
        self.assertEqual([r["id"] for r in d["requisitos"] if r["modificado"]], ["REQ-PUER-002"])
        self.assertEqual([h["requisito"] for h in d["huerfanas"]], ["REQ-COLA-001"])


class Api(Base):
    def setUp(self):
        super().setUp()
        self.api = editor_api.EditorApi(self.dir)
        resumen = ok(self, self.api.iniciar())
        self.assertTrue(resumen["hay_requisitos"])

    def indice(self, rid):
        return next(n["indice"] for n in ok(self, self.api.navegacion(0, "")) if n.get("id") == rid)

    def test_coleccion_navegacion_y_obtener(self):
        resumen = ok(self, self.api.coleccion("requisitos"))
        self.assertEqual(resumen["capitulos"][0]["pruebas"], 4)
        nodos = ok(self, self.api.navegacion(0, ""))
        r1 = next(n for n in nodos if n.get("id") == "R-1")
        self.assertEqual(r1["tipo"], "requisito")
        self.assertEqual(r1["cobertura"], {"estado": "cubierto", "pruebas": 2, "modificado": False})
        datos = ok(self, self.api.obtener(0, r1["indice"]))
        self.assertEqual(datos["tipo"], "requisito")
        self.assertEqual([p["id"] for p in datos["traza"]["pruebas"]], ["T-1", "T-2"])
        ok(self, self.api.coleccion("pruebas"))
        self.assertEqual(ok(self, self.api.navegacion(0, ""))[1]["tipo"], "prueba")

    def test_guardar_texto_deja_pruebas_por_revisar(self):
        ok(self, self.api.coleccion("requisitos"))
        i = self.indice("R-1")
        ok(self, self.api.guardar_requisito(0, i, {"descripcion": "Texto nuevo", "version": "2"}))
        self.assertTrue(ok(self, self.api.obtener(0, i))["traza"]["modificado"])
        self.assertEqual(ok(self, self.api.obtener(0, i))["traza"]["anterior"]["descripcion"], "Texto original")
        ok(self, self.api.revisar_requisito("R-1"))
        self.assertFalse(ok(self, self.api.obtener(0, i))["traza"]["modificado"])
        # Sin pruebas que lo citen, cambiar el texto no deja nada por revisar.
        ok(self, self.api.guardar_requisito(0, self.indice("R-2"), {"descripcion": "otra"}))
        self.assertFalse((self.req_dir / "revisiones.yaml").exists())

    def test_cambiar_id_tambien_en_pruebas(self):
        ok(self, self.api.coleccion("requisitos"))
        self.assertEqual(ok(self, self.api.requisito_usos("R-1"))["pruebas"], ["T-1", "T-2"])
        r = ok(self, self.api.guardar_requisito(0, self.indice("R-1"), {"id": "R-10"}, True))
        self.assertEqual(r["pruebas_actualizadas"], 2)
        texto = (self.dir / "01_pruebas.yaml").read_text(encoding="utf-8")
        self.assertIn("R-10", texto)
        self.assertNotIn("R-1,", texto)

    def test_id_repetido_y_tipo_no_valido(self):
        ok(self, self.api.coleccion("requisitos"))
        r = self.api.guardar_requisito(0, self.indice("R-2"), {"id": "R-1"})
        self.assertFalse(r["ok"])
        self.assertIn("ya existe", r["error"])
        r = self.api.guardar_requisito(0, self.indice("R-2"), {"tipo": "Raro"})
        self.assertFalse(r["ok"])

    def test_crear_y_borrar(self):
        ok(self, self.api.coleccion("requisitos"))
        nuevo = ok(self, self.api.crear_requisito(0, {"tipo": "seccion", "indice": 0}, {"id": "R-5", "titulo": "Nuevo", "criterio_validacion": ["Test"]}))
        datos = ok(self, self.api.obtener(0, nuevo["indice"]))
        self.assertEqual(datos["elemento"]["id"], "R-5")
        self.assertEqual(datos["traza"]["estado"], "sin_cobertura")
        self.assertIn("fuente", datos["elemento"])           # con todos los campos del esquema
        ok(self, self.api.borrar_requisito(0, nuevo["indice"]))
        self.assertNotIn("R-5", (self.req_dir / "01_req.yaml").read_text(encoding="utf-8"))

    def test_validacion(self):
        avisos = ok(self, self.api.validar())["problemas"]
        self.assertEqual([p["id"] for p in avisos if "catálogo" in p["mensaje"]], ["T-1"])
        ok(self, self.api.coleccion("requisitos"))
        self.assertEqual(ok(self, self.api.validar())["errores"], 0)
        ok(self, self.api.guardar_requisito(0, self.indice("R-2"), {"titulo": ""}))
        problemas = ok(self, self.api.validar())["problemas"]
        self.assertTrue(any(p["mensaje"] == "El requisito no tiene título." for p in problemas))

    def test_trazabilidad_y_lista(self):
        d = ok(self, self.api.trazabilidad())
        self.assertEqual(d["total"]["porcentaje"], 50)
        lista = ok(self, self.api.requisitos_lista())
        self.assertEqual([r["id"] for r in lista], ["R-1", "R-2", "R-3", "R-4"])
        self.assertEqual(ok(self, self.api.ir_a_requisito("R-3")), {"capitulo": 0, "indice": 4})
        self.assertEqual(ok(self, self.api.ir_a_prueba("T-2"))["indice"], 2)

    def test_matriz_desde_la_api(self):
        salida = Path(self._tmp.name) / "m.csv"
        d = ok(self, self.api.exportar_matriz(str(salida)))
        self.assertEqual(d["requisitos"], 4)
        self.assertTrue(salida.is_file())

    def test_sin_requisitos_se_crean(self):
        shutil.rmtree(self.req_dir)
        api = editor_api.EditorApi(self.dir)
        self.assertFalse(ok(self, api.iniciar())["hay_requisitos"])
        self.assertEqual(ok(self, api.requisitos_lista()), [])
        self.assertEqual(ok(self, api.validar())["avisos"], 0)     # sin catálogo no se avisa de referencias
        resumen = ok(self, api.crear_catalogo_requisitos())
        self.assertEqual((resumen["coleccion"], len(resumen["capitulos"])), ("requisitos", 1))
        self.assertTrue((self.req_dir / "indice.yaml").is_file())

    def test_exportar_docx_de_requisitos(self):
        ok(self, self.api.coleccion("requisitos"))
        salida = Path(self._tmp.name) / "req.docx"
        d = ok(self, self.api.exportar_docx(1, str(salida)))
        self.assertIn("4 requisitos", d["resumen"])
        tabla = Document(str(salida)).tables[0]
        self.assertEqual(tabla.cell(0, 0).text, "R-1")
        self.assertEqual(tabla.cell(0, 2).text, "Cubierto")
        self.assertEqual(tabla.cell(4, 1).text, "Criterio validación: Test")


class Word(unittest.TestCase):
    def test_tabla_de_la_plantilla(self):
        with TemporaryDirectory() as tmp:
            salida = Path(tmp) / "r.docx"
            with redirect_stdout(io.StringIO()):
                yaml_a_docx.main([str(EJEMPLOS / "requisitos"), "-o", str(salida)])
            tabla = Document(str(salida)).tables[1]          # la primera es el historial
            textos = [[c.text for c in fila.cells] for fila in tabla.rows]
            self.assertEqual(len(textos), 6)
            self.assertEqual(textos[0][:2], ["REQ-GEN-001", "REQ-GEN-001"])   # ID en dos columnas
            self.assertEqual(textos[3][0], "Referencia: PT-DEMO-A1 v2, apartado 1.2")
            self.assertEqual(textos[5], ["Fase: ", "Fecha: 2026-01-05", "Versión: 1.0"])


if __name__ == "__main__":
    unittest.main()
