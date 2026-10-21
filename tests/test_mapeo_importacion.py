"""El mapeo de importación del proyecto (app/mapeo_importacion.py, importacion.yaml)."""
import os
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import Workbook

from app import ajustes
from app import esquema
from app import importacion as imp
from app import mapeo_importacion as mi
from app.importacion_api import AJUSTE_MAPEO_EXCEL, EditorImportacionApi


def _excel(ruta, cabecera, filas):
    libro = Workbook()
    libro.active.title = "Llamadas"
    libro.active.append(cabecera)
    for f in filas:
        libro.active.append(f)
    libro.save(ruta)


class MapeoDelProyecto(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)
        anterior = os.environ.get("AZTARNA_AJUSTES")
        os.environ["AZTARNA_AJUSTES"] = str(self.tmp / "ajustes.json")
        self.addCleanup(lambda: os.environ.__setitem__("AZTARNA_AJUSTES", anterior) if anterior else os.environ.pop("AZTARNA_AJUSTES", None))
        self.addCleanup(esquema.configurar, None)

    def ok(self, r):
        self.assertTrue(r["ok"], r.get("error"))
        return r["datos"]

    def test_lo_elegido_por_uno_lo_usa_otro_pc(self):
        _excel(self.tmp / "ana.xlsx", ["ID", "Título", "Pasos", "Comentario"], [["T-1", "Una", "Pulsar", "ver"]])
        api = EditorImportacionApi()
        mapeo = self.ok(api.importar_analizar(str(self.tmp / "ana.xlsx"), "prueba"))["mapeo"]
        self.assertEqual(mapeo["campos"]["prueba"]["comentario"], "adicional")
        mapeo["campos"]["prueba"]["comentario"] = "ignorar"                 # Ana decide no importarla
        r = self.ok(api.importar_ejecutar(str(self.tmp / "ana.xlsx"), mapeo, "nuevo", {}, "p", str(self.tmp / "p")))
        guardado = Path(r["mapeo_guardado"])
        self.assertEqual(guardado, self.tmp / "p" / "importacion.yaml")
        texto = guardado.read_text(encoding="utf-8")
        self.assertIn("Comentario: ignorar", texto)
        self.assertIn("Llamadas: prueba", texto)
        # Otro PC (sin nada recordado) abre el proyecto e importa otro Excel con la misma columna.
        ajustes.guardar({AJUSTE_MAPEO_EXCEL: None})
        _excel(self.tmp / "luis.xlsx", ["ID", "Título", "Comentario"], [["T-2", "Otra", "no"]])
        otro = EditorImportacionApi(self.tmp / "p")
        self.ok(otro.iniciar())
        a = self.ok(otro.importar_analizar(str(self.tmp / "luis.xlsx"), "prueba"))
        self.assertEqual(a["mapeo"]["campos"]["prueba"]["comentario"], "ignorar")
        # Mismo archivo aunque la ruta se escriba distinto (en Windows, RUNNER~1 frente a runneradmin).
        self.assertEqual(Path(a["mapeo_proyecto"]["archivo"]).resolve(), guardado.resolve())
        # Sin guardar: el archivo no cambia.
        a["mapeo"]["campos"]["prueba"]["comentario"] = "adicional"
        self.ok(otro.importar_ejecutar(str(self.tmp / "luis.xlsx"), a["mapeo"], "anadir", {}, "", None, {}, False))
        self.assertIn("Comentario: ignorar", guardado.read_text(encoding="utf-8"))

    def test_editado_a_mano_conserva_comentarios_y_avisa_de_lo_que_no_vale(self):
        (self.tmp / "pruebas").mkdir()
        ruta = self.tmp / "importacion.yaml"
        ruta.write_text("# Mapeo del equipo de llamadas\nexcel:\n  hojas:\n    Llamadas: prueba\n  columnas:\n    prueba:\n"
                        "      Pasos a seguir: pasos   # así lo llama el equipo\n      Color: rojo\n", encoding="utf-8")
        mapeo, leido, problemas = mi.leer(self.tmp / "pruebas", "excel")
        self.assertEqual(leido, ruta)
        self.assertEqual((mapeo["tipos"], mapeo["campos"]["prueba"]), ({"hoja:Llamadas": "prueba"}, {"pasosaseguir": "pasos"}))
        self.assertEqual(problemas, ["Color: «rojo» no es un campo de las pruebas"])
        # Lo del proyecto manda sobre lo del PC.
        combinado = mi.combinar({"campos": {"prueba": {"pasosaseguir": "adicional", "otra": "ignorar"}}}, mapeo)
        self.assertEqual(combinado["campos"]["prueba"], {"pasosaseguir": "pasos", "otra": "ignorar"})
        _excel(self.tmp / "a.xlsx", ["ID", "PASOS A SEGUIR", "Nueva"], [["T-1", "uno"]])
        analisis = imp.analizar(self.tmp / "a.xlsx", "prueba")
        mi.guardar(self.tmp / "pruebas", "excel", imp.mapeo_propuesto(analisis, mapeo), analisis)
        texto = ruta.read_text(encoding="utf-8")
        self.assertIn("# Mapeo del equipo de llamadas", texto)
        self.assertIn("Pasos a seguir: pasos   # así lo llama el equipo", texto)       # como estaba escrito
        self.assertIn("Nueva: adicional", texto)
        self.assertNotIn("PASOS A SEGUIR", texto)

    def test_word(self):
        doc = Document()
        doc.add_heading("Pruebas", 1)
        t = doc.add_table(rows=3, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "T-1", "Una"
        t.cell(1, 0).text = "ENSAYO(s) / Proceso(s): pulsar"
        t.cell(2, 0).text = "Criticidad: alta"
        doc.save(self.tmp / "a.docx")
        analisis = imp.analizar(self.tmp / "a.docx", "prueba")
        mapeo = imp.mapeo_propuesto(analisis)
        mi.guardar(self.tmp, "word", mapeo, analisis)
        leido, _ruta, problemas = mi.leer(self.tmp, "word")
        self.assertEqual(problemas, [])
        self.assertEqual(leido["campos"]["prueba"], {"ensayoproceso": "pasos", "criticidad": "ignorar"})   # «(s)» no cuenta
        self.assertEqual(mi.leer(self.tmp, "excel")[0]["campos"]["prueba"], {})              # cada origen, lo suyo


if __name__ == "__main__":
    unittest.main()
