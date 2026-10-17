"""Comparar validaciones y ejecuciones (app/comparar.py) y requisitos en la validación."""
import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document

from app import comparar
from app import ejecucion_core as ej
from app import validacion_core as vc
from app.informe_validacion import generar_informe
from app.validacion_api import EditorValidacionApi

RAIZ = Path(__file__).resolve().parents[1]


def ok(test, r):
    test.assertTrue(r["ok"], r.get("error"))
    return r.get("datos")


class Categorias(unittest.TestCase):
    def test_categoria(self):
        self.assertEqual(comparar.categoria("validada", "no_ok"), "nuevo_fallo")
        self.assertEqual(comparar.categoria("no_ok", "validada"), "arreglada")
        self.assertEqual(comparar.categoria("no_ok", "no_aplica"), "arreglada")
        self.assertEqual(comparar.categoria("no_ok", "no_ok"), "sigue_fallando")
        self.assertEqual(comparar.categoria("bloqueada", "validada"), "otros_cambios")
        self.assertEqual(comparar.categoria("no_ok", "pendiente"), "otros_cambios")
        self.assertEqual(comparar.categoria(None, "ok"), "nueva")
        self.assertEqual(comparar.categoria("ok", None), "ya_no_esta")
        self.assertEqual(comparar.categoria("ok", "ok"), "sin_cambios")

    def test_comparar(self):
        r = comparar.comparar({"A": {"estado": "ok", "titulo": "a"}, "B": {"estado": "no_ok", "titulo": "b"}},
                              {"B": {"estado": "ok", "titulo": "b"}, "C": {"estado": "no_ok", "titulo": "c"}})
        self.assertEqual([(f["id"], f["categoria"]) for f in r["filas"]], [("B", "arreglada"), ("C", "nueva"), ("A", "ya_no_esta")])
        self.assertEqual(r["cuentas"]["arreglada"], 1)


class ConLosEjemplos(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.base = Path(self._tmp.name) / "ejemplos"
        shutil.copytree(RAIZ / "ejemplos", self.base)
        self.v23 = self.base / "validaciones" / "validacion_v2.3"
        self.v22 = self.base / "validaciones" / "validacion_v2.2"

    def test_comparacion_de_la_validacion_de_ejemplo(self):
        c = vc.consolidar(vc.cargar(self.v23))["comparacion"]
        self.assertTrue(c["encontrada"])
        self.assertEqual(c["nombre"], "Validación v2.2 (ejemplo)")
        cambios = {f["id"]: f["categoria"] for f in c["filas"] if f["categoria"] != "sin_cambios"}
        self.assertEqual(cambios, {"DEMO-NIV-001": "nuevo_fallo", "DEMO-PUER-001": "arreglada", "DEMO-CAN-001": "otros_cambios",
                                   "DEMO-CAN-002": "otros_cambios", "DEMO-SEG-001": "otros_cambios"})
        # Quitarla y no poder compararla consigo misma.
        v = vc.cargar(self.v23)
        with self.assertRaises(ValueError):
            vc.comparar_con(v, self.v23)
        vc.comparar_con(v, None)
        self.assertIsNone(vc.consolidar(vc.cargar(self.v23))["comparacion"])

    def test_ejecuciones(self):
        r = comparar.comparar_carpetas(self.base / "ejecuciones" / "2026-10-16_Torre_v2.2", self.base / "ejecuciones" / "2026-11-05_Torre_Luis")
        self.assertEqual(r["anterior"]["tipo"], "ejecucion")
        cats = {f["id"]: f["categoria"] for f in r["filas"]}
        self.assertEqual(cats["DEMO-PUER-001"], "arreglada")
        self.assertEqual(cats["DEMO-NIV-001"], "sin_cambios")
        with self.assertRaises(ValueError):
            comparar.comparar_carpetas(self.v22, self.base / "ejecuciones" / "2026-11-05_Torre_Luis")

    def test_requisitos_en_la_validacion(self):
        r = vc.consolidar(vc.cargar(self.v23))["requisitos"]
        self.assertEqual(r["totales"], {"fallido": 1, "bloqueado": 1, "pendiente": 3, "verificado": 4, "sin_pruebas": 6, "no_aplica": 3})
        por_id = {x["id"]: x for x in r["lista"]}
        self.assertEqual(por_id["REQ-LLAM-001"]["estado"], "verificado")
        self.assertEqual(por_id["REQ-LLAM-003"]["estado"], "pendiente")
        self.assertEqual(por_id["REQ-GEN-001"]["estado"], "no_aplica")
        self.assertEqual(r["impacto"], ["DEMO-PUER-002"])             # REQ-PUER-002 ha cambiado
        self.assertEqual(vc._estado_requisito(True, ["validada", "no_ok", "bloqueada"]), "fallido")
        self.assertEqual(vc._estado_requisito(True, ["validada", "bloqueada"]), "bloqueado")
        self.assertEqual(vc._estado_requisito(True, ["no_aplica"]), "no_aplica")
        self.assertEqual(vc._estado_requisito(True, []), "sin_pruebas")

    def test_informe_con_comparacion_y_requisitos(self):
        salida = self.base / "informe.docx"
        generar_informe(vc.cargar(self.v23), salida)
        texto = "\n".join(p.text for p in Document(salida).paragraphs)
        celdas = " ".join(c.text for t in Document(salida).tables for f in t.rows for c in f.cells)
        self.assertIn("Comparación con la versión anterior", texto)
        self.assertIn("Requisitos", texto)
        self.assertIn("Nuevos fallos", celdas)
        self.assertIn("REQ-PUER-002 (cambiado)", celdas)

    def test_api_comparar_e_impacto(self):
        api = EditorValidacionApi(self.base / "yaml")
        ok(self, api.iniciar())
        v = ok(self, api.validacion_abrir(str(self.v23)))
        self.assertEqual(v["comparacion"]["cuentas"]["nuevo_fallo"], 1)
        self.assertIsNone(ok(self, api.validacion_quitar_comparacion())["comparacion"])
        self.assertTrue(ok(self, api.validacion_comparar_con(str(self.v22)))["comparacion"]["encontrada"])
        # Ejecución de impacto: solo las pruebas de los requisitos cambiados, dentro de la validación.
        e = ok(self, api.validacion_nueva_ejecucion("Impacto", {"solo_ids": ["DEMO-PUER-002"]}, {"entorno_pruebas": "Torre"}))
        self.assertEqual([p["id"] for p in ej.cargar(e["carpeta"])["datos"]["pruebas"]], ["DEMO-PUER-002"])
        self.assertTrue(e["carpeta"].startswith(str(self.v23)))
        r = ok(self, api.ejecucion_comparar(str(self.base / "ejecuciones" / "2026-10-16_Torre_v2.2")))
        self.assertEqual(r["cuentas"]["nueva"], 0)


if __name__ == "__main__":
    unittest.main()
