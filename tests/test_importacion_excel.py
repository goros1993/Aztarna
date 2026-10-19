"""Importar pruebas y requisitos de un Excel (app/importacion_excel.py)."""
import tempfile
import unittest
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app import esquema
from app import importacion as imp
from app import importacion_excel as ix
from app.editor_api import to_plain
from app.editor_core import items_in_chapter, read_chapters

CABECERA = ["ID", "Título", "Sección", "Descripción", "Pasos", "Requisitos", "Tags", "Entorno", "Automatizable", "Fecha", "Notas"]


def _libro(ruta, filas, otra_hoja=True):
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Llamadas"
    hoja.append(["Plan de pruebas de llamadas"])          # un título encima de la cabecera
    hoja.append(CABECERA)
    for fila in filas:
        hoja.append(fila)
    if otra_hoja:
        notas = libro.create_sheet("Notas")
        notas.append(["Esta hoja solo tiene notas."])
    libro.save(ruta)


FILAS = [
    ["T-001", "Llamada de cabina", "Cabina", "Una llamada desde la cabina.", "1. Pulsar el botón\n2. La cabina llega",
     "REQ-1, REQ-2", "smoke; regresion", "A/M", "Sí", date(2026, 10, 15), "Revisar con Ana"],
    [None, None, None, None, "3. Se abren las puertas"],      # continúa la de arriba
    ["T-002", "Llamada de piso", None, "Desde el piso.", "- Pulsar\n- Esperar", "REQ-3", None, "Torre", "no"],   # misma sección
    ["T-003", "Alarma", "Cabina > Alarmas", None, "Pulsar alarma", None, None, "HIL"],
]


class Excel(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.addCleanup(esquema.configurar, None)
        esquema.configurar(None)
        self.tmp = Path(self._tmp.name)
        self.xlsx = self.tmp / "pruebas.xlsx"
        _libro(self.xlsx, FILAS)

    def test_analizar_propone_hojas_y_columnas(self):
        a = imp.analizar(self.xlsx, "prueba")
        self.assertEqual([(f["hoja"], f["tipo"], f["tablas"]) for f in a["familias"]], [("Llamadas", "prueba", 4), ("Notas", "ignorar", 0)])
        propuestas = {e["etiqueta"]: e["propuesta"]["prueba"] for e in a["familias"][0]["etiquetas"]}
        self.assertEqual(propuestas, {"ID": "id", "Título": "titulo", "Sección": "seccion", "Descripción": "descripcion",
                                      "Pasos": "pasos", "Requisitos": "requisitos", "Tags": "etiquetas", "Entorno": "entorno",
                                      "Automatizable": "automatizable", "Fecha": "fecha", "Notas": "adicional"})
        self.assertEqual(a["familias"][0]["ejemplos"], ["T-001", "T-002", "T-003"])

    def test_convertir(self):
        mapeo = imp.mapeo_propuesto(imp.analizar(self.xlsx, "prueba"))
        conversion = imp.convertir_documento_ruta(self.xlsx, mapeo)
        [capitulo] = conversion["capitulos"]
        self.assertEqual(capitulo["capitulo"], "Llamadas")                    # sin columna Capítulo: la hoja
        resumen = [(e.get("seccion") or e["prueba"]["id"], e.get("nivel") or e["prueba"]["nivel"]) for e in capitulo["contenido"]]
        self.assertEqual(resumen, [("Cabina", 2), ("T-001", 3), ("T-002", 3), ("Alarmas", 3), ("T-003", 4)])
        t1, t2 = (e["prueba"] for e in capitulo["contenido"][1:3])
        self.assertEqual(t1["pasos"], ["Pulsar el botón", "La cabina llega", "Se abren las puertas"])
        self.assertEqual((t1["requisitos"], t1["etiquetas"]), (["REQ-1", "REQ-2"], ["smoke", "regresion"]))
        self.assertEqual((t1["entorno"], t1["automatizable"], t1["fecha"]), (["Automático", "Maqueta"], "si", "15/10/2026"))
        self.assertEqual(t1["campos_adicionales"], {"Notas": "Revisar con Ana"})
        self.assertEqual(t2["pasos"], ["Pulsar", "Esperar"])
        # «HIL» no es un entorno: se avisa y se conserva como campo adicional.
        t3 = capitulo["contenido"][4]["prueba"]
        self.assertEqual(t3["campos_adicionales"], {"Entorno de pruebas": "HIL"})
        avisos = [(a.tipo, a.capitulo) for a in conversion["avisos"]]
        self.assertIn(("valor_descartado", "hoja «Llamadas», fila 6"), avisos)

    def test_importar_y_despues_anadir_otra_tanda(self):
        mapeo = imp.mapeo_propuesto(imp.analizar(self.xlsx, "prueba"))
        r = imp.importar(self.xlsx, mapeo, self.tmp / "p")
        self.assertEqual(r["totales"], {"prueba": 3})
        # La semana siguiente: la misma hoja con una prueba más en «Cabina» y T-002 cambiada.
        segunda = self.tmp / "segunda.xlsx"
        filas = [list(f) for f in FILAS]
        filas[2][3] = "Desde el piso, con la puerta abierta."
        filas.append(["T-004", "Llamada doble", "Cabina", "Dos a la vez.", "Pulsar dos"])
        _libro(segunda, filas)
        p = imp.preparar(segunda, mapeo, self.tmp / "p", "anadir")
        rev = imp.revision(p["conversion"], p["existentes"], p["carpetas"], p["imagenes"])["tipos"]["prueba"]
        self.assertEqual([(c["id"], c["cambios"]) for c in rev["existentes"]], [("T-001", []), ("T-002", ["Descripción"]), ("T-003", [])])
        # Lo que propone el asistente: actualizar los que cambian y no tocar los demás.
        acciones = {str(c["orden"]): "actualizar" if c["cambios"] else "omitir" for c in rev["existentes"]}
        r = imp.importar(segunda, mapeo, self.tmp / "p", "anadir", acciones={"prueba": acciones})
        self.assertEqual((r["totales"], r["actualizados"], r["archivos"]["prueba"]), ({"prueba": 1}, {"prueba": 1}, []))
        [capitulo] = read_chapters(self.tmp / "p" / "pruebas")
        ids = [e["id"] for _i, e in items_in_chapter(capitulo, "prueba")]
        self.assertEqual(ids, ["T-001", "T-002", "T-004", "T-003"])            # T-004 al final de «Cabina», antes de «Alarmas»
        t2 = next(to_plain(e) for _i, e in items_in_chapter(capitulo, "prueba") if e["id"] == "T-002")
        self.assertEqual(t2["descripcion"], "Desde el piso, con la puerta abierta.")
        self.assertNotIn("_columnas", (self.tmp / "p" / "pruebas" / capitulo["path"].name).read_text(encoding="utf-8"))

    def test_actualizar_desde_excel_conserva_lo_que_no_tiene_columna(self):
        """El Excel trae menos columnas que el esquema: al actualizar, lo demás no se borra."""
        import shutil
        from app.editor_core import save_chapter
        mapeo = imp.mapeo_propuesto(imp.analizar(self.xlsx, "prueba"))
        imp.importar(self.xlsx, mapeo, self.tmp / "p")
        [capitulo] = read_chapters(self.tmp / "p" / "pruebas")
        t1 = next(e for _i, e in items_in_chapter(capitulo, "prueba") if e["id"] == "T-001")
        t1["situacion_inicial"] = ["En reposo"]                    # escrito después en el editor
        t1["campos_adicionales"]["Prioridad"] = "alta"
        save_chapter(capitulo)
        # Un Excel solo con ID, título y descripción.
        libro = Workbook()
        libro.active.title = "Llamadas"
        libro.active.append(["ID", "Título", "Descripción"])
        libro.active.append(["T-001", "Llamada de cabina", "Descripción nueva."])
        libro.save(self.tmp / "pocas.xlsx")
        mapeo2 = imp.mapeo_propuesto(imp.analizar(self.tmp / "pocas.xlsx", "prueba"))
        p = imp.preparar(self.tmp / "pocas.xlsx", mapeo2, self.tmp / "p", "anadir")
        rev = imp.revision(p["conversion"], p["existentes"], p["carpetas"], p["imagenes"])["tipos"]["prueba"]
        self.assertEqual([(c["id"], c["cambios"]) for c in rev["existentes"]], [("T-001", ["Descripción"])])
        imp.importar(self.tmp / "pocas.xlsx", mapeo2, self.tmp / "p", "anadir", acciones={"prueba": {"0": "actualizar"}})
        t1 = next(to_plain(e) for _i, e in items_in_chapter(read_chapters(self.tmp / "p" / "pruebas")[0], "prueba") if e["id"] == "T-001")
        self.assertEqual(t1["descripcion"], "Descripción nueva.")
        self.assertEqual((t1["situacion_inicial"], t1["pasos"][0], t1["requisitos"]), (["En reposo"], "Pulsar el botón", ["REQ-1", "REQ-2"]))
        self.assertEqual(t1["campos_adicionales"], {"Notas": "Revisar con Ana", "Prioridad": "alta"})
        shutil.rmtree(self.tmp / "p")

    def test_columna_capitulo_y_requisitos(self):
        libro = Workbook()
        hoja = libro.active
        hoja.append(["Código", "Requisito", "Módulo", "Tipo", "Método de verificación", "Texto"])
        hoja.append(["R-1", "Atiende llamadas", "Llamadas", "Obligatorio", "Test", "Debe atenderlas."])
        hoja.append(["R-2", "Avisa", None, "obligatorio", "Inspección / Test", "Debe avisar."])      # mismo módulo
        hoja.append(["R-3", "Puertas", "Puertas", "Informativo", None, "Se abren."])
        libro.save(self.tmp / "reqs.xlsx")
        a = imp.analizar(self.tmp / "reqs.xlsx", "requisito")
        self.assertEqual(a["familias"][0]["tipo"], "requisito")
        mapeo = imp.mapeo_propuesto(a)
        caps = imp.convertir_documento_ruta(self.tmp / "reqs.xlsx", mapeo)["capitulos"]
        self.assertEqual([(c["capitulo"], [e["requisito"]["id"] for e in c["contenido"]]) for c in caps],
                         [("Llamadas", ["R-1", "R-2"]), ("Puertas", ["R-3"])])
        r2 = caps[0]["contenido"][1]["requisito"]
        self.assertEqual((r2["tipo"], r2["criterio_validacion"], r2["descripcion"]), ("Obligatorio", ["Inspección", "Test"], "Debe avisar."))

    def test_plantilla_se_importa_tal_cual(self):
        ruta = ix.plantilla(self.tmp / "plantilla.xlsx", "prueba")
        hoja = load_workbook(ruta).active
        hoja.append(["Llamadas", "Cabina", "T-1", "Una prueba", "Descripción", "Ajuste", "En reposo", "Pulsar\nEsperar"])
        hoja.parent.save(ruta)
        a = imp.analizar(ruta)
        self.assertEqual([f["tipo"] for f in a["familias"]], ["prueba", "ignorar"])
        self.assertFalse([e for e in a["familias"][0]["etiquetas"] if e["propuesta"]["prueba"] in ("adicional", "ignorar")])
        [capitulo] = imp.convertir_documento_ruta(ruta, imp.mapeo_propuesto(a))["capitulos"]
        prueba = capitulo["contenido"][1]["prueba"]
        self.assertEqual((prueba["id"], prueba["situacion_inicial"], prueba["pasos"]), ("T-1", ["En reposo"], ["Pulsar", "Esperar"]))


class ExcelApi(unittest.TestCase):
    """El asistente del editor con un Excel: análisis, importación, mapeo recordado aparte y plantilla."""

    def setUp(self):
        import os
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)
        anterior = os.environ.get("AZTARNA_AJUSTES")
        os.environ["AZTARNA_AJUSTES"] = str(self.tmp / "ajustes.json")
        self.addCleanup(lambda: os.environ.__setitem__("AZTARNA_AJUSTES", anterior) if anterior else os.environ.pop("AZTARNA_AJUSTES", None))
        self.addCleanup(esquema.configurar, None)
        self.xlsx = self.tmp / "pruebas.xlsx"
        _libro(self.xlsx, FILAS)

    def ok(self, r):
        self.assertTrue(r["ok"], r.get("error"))
        return r["datos"]

    def test_flujo_completo(self):
        from app import ajustes
        from app.importacion_api import AJUSTE_MAPEO, AJUSTE_MAPEO_EXCEL, EditorImportacionApi
        api = EditorImportacionApi()
        a = self.ok(api.importar_analizar(str(self.xlsx), "prueba"))
        self.assertEqual(a["origen"], "excel")
        self.assertNotIn("tabla", [t["valor"] for t in a["tipos"]])           # una hoja no es una tabla informativa
        self.assertEqual([o["valor"] for o in a["opciones"]["prueba"][:4]], ["id", "titulo", "capitulo", "seccion"])
        mapeo = a["mapeo"]
        mapeo["campos"]["prueba"]["notas"] = "ignorar"
        rev = self.ok(api.importar_revisar(str(self.xlsx), mapeo, "nuevo"))
        self.assertEqual(rev["tipos"]["prueba"]["total"], 3)
        r = self.ok(api.importar_ejecutar(str(self.xlsx), mapeo, "nuevo", {}, "proyecto", str(self.tmp / "proyecto")))
        self.assertEqual(r["resultado"]["totales"], {"prueba": 3})
        self.assertEqual(ajustes.leer()[AJUSTE_MAPEO_EXCEL]["campos"]["prueba"]["notas"], "ignorar")
        self.assertNotIn(AJUSTE_MAPEO, ajustes.leer())                         # el del Word no se toca
        # Resultados: solo desde un Word.
        self.assertIn("desde un Word", api.importar_analizar(str(self.xlsx), "resultado")["error"])

    def test_plantilla_desde_la_api(self):
        from app.importacion_api import EditorImportacionApi
        r = self.ok(EditorImportacionApi().importar_plantilla_excel("requisito", str(self.tmp / "reqs")))
        self.assertTrue(r["archivo"].endswith("reqs.xlsx"))
        cabecera = [c.value for c in load_workbook(r["archivo"]).active[1]]
        self.assertEqual(cabecera[:4], ["Capítulo", "Sección", "ID", "Título"])
        self.assertIn("Criterio validación", cabecera)


if __name__ == "__main__":
    unittest.main()
