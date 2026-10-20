"""Pruebas de las ejecuciones (ejecucion_core.py)."""
import os
import shutil
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

import base64
import io

from PIL import Image

from app import ejecucion_core as ej
from app import esquema
from app.editor_core import read_chapters, save_chapter
from app.ejecucion_api import EditorEjecucionApi
from app import informe_docx
from docx import Document

EJEMPLOS = Path(__file__).resolve().parent.parent / "ejemplos" / "yaml" / "pruebas"


class EjecucionBase(unittest.TestCase):
    def setUp(self):
        self._tmp = TemporaryDirectory()
        self.tmp = Path(self._tmp.name).resolve()   # en Windows, Temp puede venir con nombre corto (RUNNER~1)
        self.pruebas = self.tmp / "pruebas"
        shutil.copytree(EJEMPLOS, self.pruebas)
        shutil.copy(EJEMPLOS.parent / esquema.ARCHIVO_CONFIGURACION, self.tmp)
        esquema.configurar(self.pruebas)
        self.carpeta = self.tmp / "ejecuciones" / "regresion"

    def tearDown(self):
        self._tmp.cleanup()

    def capitulos(self):
        return read_chapters(self.pruebas)

    def crear(self, filtro=None, **datos):
        return ej.crear(self.carpeta, self.pruebas, "Regresión v2.3",
                        filtro if filtro is not None else {"etiquetas": ["regresion"]}, datos)


class SeleccionTests(EjecucionBase):
    def ids(self, filtro):
        return [item["id"] for item in ej.seleccionar(self.capitulos(), filtro)]

    def test_etiquetas_basta_con_una(self):
        ids = self.ids({"etiquetas": ["smoke", "manual"]})
        self.assertEqual(ids, ["DEMO-LLAM-001", "DEMO-NIV-001", "DEMO-PUER-001", "DEMO-SEG-001", "DEMO-VAR-002"])

    def test_se_combinan_criterios(self):
        ids = self.ids({"etiquetas": ["regresion"], "entorno": ["Torre"], "capitulos": ["01_ejemplo_basico.yaml"]})
        self.assertEqual(ids, ["DEMO-NIV-001"])

    def test_entorno_y_maniobra_basta_con_uno(self):
        ids = self.ids({"maniobra": ["Dúplex"], "automatizable": ["si", "hecho"]})
        self.assertEqual(ids, ["DEMO-LLAM-001", "DEMO-LLAM-003", "DEMO-NIV-001"])

    def test_incluir_y_excluir_ids_mandan(self):
        ids = self.ids({"etiquetas": ["smoke"], "incluir_ids": ["DEMO-DIAG-001"], "excluir_ids": ["DEMO-NIV-001"]})
        self.assertIn("DEMO-DIAG-001", ids)
        self.assertNotIn("DEMO-NIV-001", ids)

    def test_sin_criterio_entran_todas(self):
        self.assertEqual(len(self.ids({})), 14)

    def test_opciones_para_la_interfaz(self):
        opciones = ej.opciones_filtro(self.capitulos())
        self.assertIn("regresion", opciones["etiquetas"])
        self.assertEqual(opciones["entorno"], ["Automático", "Maqueta", "Torre"])
        self.assertEqual(opciones["maniobra"], ["Simplex", "Dúplex"])
        self.assertEqual(len(opciones["capitulos"]), 4)


class CrearTests(EjecucionBase):
    def test_crea_la_ejecucion_con_la_lista_congelada(self):
        ejecucion = self.crear(objeto="Firmware 2.3.0-rc2", responsable="Unai")
        cargada = ej.cargar(self.carpeta)
        self.assertEqual(cargada["datos"]["nombre"], "Regresión v2.3")
        self.assertEqual(cargada["datos"]["objeto"], "Firmware 2.3.0-rc2")
        self.assertEqual(cargada["datos"]["estado"], "abierta")
        self.assertEqual(len(ej.ids_de(cargada)), 10)
        self.assertEqual(ej.ids_de(cargada), ej.ids_de(ejecucion))
        self.assertTrue(all(len(p["huella"]) == 12 for p in cargada["datos"]["pruebas"]))
        self.assertTrue((self.carpeta / ej.CARPETA_EVIDENCIAS).is_dir())

    def test_no_toca_los_yaml_de_pruebas(self):
        antes = {p.name: p.read_bytes() for p in self.pruebas.glob("*.yaml")}
        ejecucion = self.crear()
        ej.registrar(ejecucion, "DEMO-LLAM-001", "ok")
        ej.registrar(ejecucion, "DEMO-LLAM-002", "no_ok", observaciones="No se ilumina")
        self.assertEqual(antes, {p.name: p.read_bytes() for p in self.pruebas.glob("*.yaml")})

    def test_rechaza_carpeta_con_contenido_y_filtro_sin_pruebas(self):
        self.carpeta.mkdir(parents=True)
        (self.carpeta / "otra.txt").write_text("x")
        with self.assertRaisesRegex(ValueError, "carpeta vacía"):
            self.crear()
        with self.assertRaisesRegex(ValueError, "Ninguna prueba"):
            ej.crear(self.tmp / "vacia", self.pruebas, "X", {"etiquetas": ["no-existe"]})

    def test_rechaza_ids_repetidos(self):
        capitulo = self.capitulos()[0]
        capitulo["data"]["contenido"][3]["prueba"]["id"] = "DEMO-LLAM-002"
        save_chapter(capitulo)
        with self.assertRaisesRegex(ValueError, "repetido"):
            self.crear()

    def test_la_huella_ignora_el_nivel_pero_no_el_contenido(self):
        prueba = {"id": "A", "titulo": "T", "nivel": 3}
        self.assertEqual(ej.huella(prueba), ej.huella({**prueba, "nivel": 5}))
        self.assertNotEqual(ej.huella(prueba), ej.huella({**prueba, "titulo": "Otro"}))


class DatosTests(EjecucionBase):
    def test_se_guarda_como_ejecucion_yaml(self):
        self.crear(entorno_pruebas="Torre", entorno="Torre de Zaragoza")
        self.assertTrue((self.carpeta / ej.ARCHIVO_EJECUCION).is_file())
        self.assertFalse((self.carpeta / ej.ARCHIVO_CAMPAÑA).exists())
        datos = ej.cargar(self.carpeta)["datos"]
        self.assertEqual((datos["entorno_pruebas"], datos["entorno"]), ("Torre", "Torre de Zaragoza"))

    def test_las_carpetas_antiguas_se_siguen_leyendo_y_guardando(self):
        self.crear()
        (self.carpeta / ej.ARCHIVO_EJECUCION).rename(self.carpeta / ej.ARCHIVO_CAMPAÑA)
        ejecucion = ej.cargar(self.carpeta)
        ej.editar_datos(ejecucion, {"entorno_pruebas": "Maqueta"})
        self.assertFalse((self.carpeta / ej.ARCHIVO_EJECUCION).exists())
        self.assertEqual(ej.cargar(self.carpeta)["datos"]["entorno_pruebas"], "Maqueta")

    def test_entorno_de_pruebas_de_la_lista(self):
        with self.assertRaises(ValueError):
            self.crear(entorno_pruebas="Laboratorio")
        ejecucion = self.crear()
        with self.assertRaises(ValueError):
            ej.editar_datos(ejecucion, {"entorno_pruebas": "Laboratorio"})
        ej.editar_datos(ejecucion, {"responsable": "Ana", "nombre": "no se cambia"})
        self.assertEqual(ej.cargar(self.carpeta)["datos"]["responsable"], "Ana")
        self.assertEqual(ej.cargar(self.carpeta)["datos"]["nombre"], "Regresión v2.3")


class ResultadosTests(EjecucionBase):
    def setUp(self):
        super().setUp()
        self.ejecucion = self.crear()

    def test_registrar_y_cargar(self):
        ej.registrar(self.ejecucion, "DEMO-LLAM-001", "ok", ejecutor="Unai", observaciones="Sin incidencias")
        cargada = ej.cargar(self.carpeta)
        self.assertEqual(ej.estado_de(cargada, "DEMO-LLAM-001"), "ok")
        self.assertEqual(cargada["resultados"]["DEMO-LLAM-001"]["ejecutor"], "Unai")
        self.assertEqual(ej.estado_de(cargada, "DEMO-LLAM-002"), "pendiente")

    def test_un_nuevo_resultado_guarda_el_anterior_en_el_historial(self):
        ej.registrar(self.ejecucion, "DEMO-LLAM-001", "no_ok", observaciones="No nivela", incidencia="BUG-1")
        ej.registrar(self.ejecucion, "DEMO-LLAM-001", "ok")
        resultado = ej.cargar(self.carpeta)["resultados"]["DEMO-LLAM-001"]
        self.assertEqual(resultado["estado"], "ok")
        self.assertEqual(len(resultado["intentos"]), 1)
        self.assertEqual(resultado["intentos"][0]["incidencia"], "BUG-1")

    def test_volver_a_pendiente_recupera_el_resultado_anterior(self):
        ej.registrar(self.ejecucion, "DEMO-LLAM-001", "no_ok", observaciones="Falla")
        ej.registrar(self.ejecucion, "DEMO-LLAM-001", "ok")
        ej.registrar(self.ejecucion, "DEMO-LLAM-001", "pendiente")
        self.assertEqual(ej.estado_de(self.ejecucion, "DEMO-LLAM-001"), "no_ok")
        ej.registrar(self.ejecucion, "DEMO-LLAM-001", "pendiente")
        self.assertEqual(ej.estado_de(self.ejecucion, "DEMO-LLAM-001"), "pendiente")
        self.assertNotIn("DEMO-LLAM-001", ej.cargar(self.carpeta)["resultados"])

    def test_editar_notas_no_crea_otro_intento(self):
        ej.registrar(self.ejecucion, "DEMO-LLAM-001", "no_ok")
        ej.editar_notas(self.ejecucion, "DEMO-LLAM-001", observaciones="Detalle", incidencia="BUG-7")
        resultado = ej.cargar(self.carpeta)["resultados"]["DEMO-LLAM-001"]
        self.assertEqual((resultado["observaciones"], resultado["incidencia"]), ("Detalle", "BUG-7"))
        self.assertNotIn("intentos", resultado)

    def test_rechaza_estado_o_prueba_desconocidos(self):
        with self.assertRaisesRegex(ValueError, "Estado no válido"):
            ej.registrar(self.ejecucion, "DEMO-LLAM-001", "regular")
        with self.assertRaisesRegex(ValueError, "no forma parte"):
            ej.registrar(self.ejecucion, "DEMO-DIAG-001", "ok")


class ResumenTests(EjecucionBase):
    def setUp(self):
        super().setUp()
        self.ejecucion = self.crear()

    def test_totales_y_porcentajes(self):
        for test_id, estado in (("DEMO-LLAM-001", "ok"), ("DEMO-LLAM-002", "ok"), ("DEMO-LLAM-003", "no_ok"),
                                ("DEMO-NIV-001", "bloqueada"), ("DEMO-CAN-001", "no_aplica")):
            ej.registrar(self.ejecucion, test_id, estado, observaciones="x")
        r = ej.resumen(self.ejecucion, self.capitulos())
        self.assertEqual(r["total"], 10)
        self.assertEqual(r["totales"], {"pendiente": 5, "ok": 2, "no_ok": 1, "bloqueada": 1, "no_aplica": 1})
        self.assertEqual(r["ejecutadas"], 5)
        self.assertEqual(r["porcentaje_ok"], 22.2)        # 2 OK de 9 que aplican
        self.assertEqual(r["por_capitulo"]["Ejemplo básico"]["ok"], 2)

    def test_avisa_de_pruebas_cambiadas_borradas_y_fallos_sin_observaciones(self):
        ej.registrar(self.ejecucion, "DEMO-LLAM-001", "ok", huella_actual=self.ejecucion["datos"]["pruebas"][0]["huella"])
        ej.registrar(self.ejecucion, "DEMO-LLAM-002", "no_ok")
        capitulos = self.capitulos()
        capitulos[0]["data"]["contenido"][3]["prueba"]["titulo"] = "Título cambiado"
        del capitulos[0]["data"]["contenido"][4]          # DEMO-LLAM-002 desaparece
        resumen = ej.resumen(self.ejecucion, capitulos)
        mensajes = {(p["id"], p["mensaje"]) for p in resumen["problemas"]}
        self.assertIn(("DEMO-LLAM-001", "La prueba ha cambiado después de ejecutarse."), mensajes)
        self.assertEqual([p["id"] for p in resumen["pruebas"] if p["cambiada"]], ["DEMO-LLAM-001"])
        self.assertIn(("DEMO-LLAM-002", "La prueba ya no existe en los YAML de pruebas."), mensajes)
        self.assertIn(("DEMO-LLAM-002", "Resultado NO OK sin observaciones."), mensajes)

    def test_dar_por_revisada_una_prueba_cambiada(self):
        ej.registrar(self.ejecucion, "DEMO-LLAM-001", "ok", huella_actual="antigua")
        capitulos = self.capitulos()
        actual = ej.huella(ej.indice_por_id(capitulos)["DEMO-LLAM-001"]["prueba"])
        self.assertTrue(next(p for p in ej.resumen(self.ejecucion, capitulos)["pruebas"] if p["id"] == "DEMO-LLAM-001")["cambiada"])
        ej.dar_por_revisada(self.ejecucion, "DEMO-LLAM-001", actual)
        recargada = ej.cargar(self.ejecucion["carpeta"])
        self.assertEqual(recargada["resultados"]["DEMO-LLAM-001"]["huella"], actual)
        self.assertIn("revisada", recargada["resultados"]["DEMO-LLAM-001"])
        self.assertFalse(any(p["cambiada"] for p in ej.resumen(recargada, capitulos)["pruebas"]))
        with self.assertRaises(ValueError):          # sin resultado no hay nada que revisar
            ej.dar_por_revisada(self.ejecucion, "DEMO-LLAM-002", actual)


class CicloTests(EjecucionBase):
    def test_resincronizar_añade_las_pruebas_nuevas(self):
        ejecucion = self.crear()
        capitulo = self.capitulos()[1]
        capitulo["data"]["contenido"][4]["prueba"]["etiquetas"] = ["regresion"]   # DEMO-DIAG-001
        save_chapter(capitulo)
        cambios = ej.resincronizar(ejecucion, self.capitulos())
        self.assertEqual(cambios["añadidas"], ["DEMO-DIAG-001"])
        self.assertIn("DEMO-DIAG-001", ej.ids_de(ej.cargar(self.carpeta)))

    def test_cerrar_exige_que_no_quede_nada_pendiente_y_bloquea_cambios(self):
        ejecucion = self.crear({"etiquetas": ["manual"]})
        with self.assertRaisesRegex(ValueError, "pendientes"):
            ej.cerrar(ejecucion, self.capitulos())
        ej.registrar(ejecucion, "DEMO-VAR-002", "ok")
        ej.cerrar(ejecucion, self.capitulos())
        cargada = ej.cargar(self.carpeta)
        self.assertEqual(cargada["datos"]["estado"], "cerrada")
        with self.assertRaisesRegex(ValueError, "cerrada"):
            ej.registrar(cargada, "DEMO-VAR-002", "no_ok")

    def test_ruta_relativa_a_las_pruebas(self):
        ejecucion = self.crear()
        self.assertEqual(ejecucion["datos"]["pruebas_yaml"], "../../pruebas")
        self.assertEqual(ejecucion["datos"]["pruebas_yaml_absoluta"], str(self.pruebas.resolve()))
        self.assertEqual(ej.carpeta_pruebas(ejecucion), self.pruebas.resolve())


class UbicacionTests(EjecucionBase):
    def test_moverlo_todo_junto_usa_la_ruta_relativa(self):
        self.crear()
        otro = self.tmp / "otro_pc"
        shutil.copytree(self.tmp / "ejecuciones", otro / "ejecuciones")
        shutil.copytree(self.pruebas, otro / "pruebas")
        shutil.rmtree(self.pruebas)
        ejecucion = ej.cargar(otro / "ejecuciones" / "regresion")
        self.assertEqual(ej.carpeta_pruebas(ejecucion), (otro / "pruebas").resolve())

    def test_mover_solo_la_ejecucion_usa_la_ruta_absoluta(self):
        self.crear()
        movida = self.tmp / "en" / "otro" / "sitio"
        shutil.move(str(self.carpeta), movida)
        self.assertEqual(ej.carpeta_pruebas(ej.cargar(movida)), self.pruebas.resolve())

    def test_sin_pruebas_se_puede_reubicar(self):
        self.crear()
        nueva = self.tmp / "pruebas_nuevas"
        shutil.move(str(self.pruebas), nueva)
        ejecucion = ej.cargar(self.carpeta)
        self.assertIsNone(ej.carpeta_pruebas(ejecucion))
        vacia = self.tmp / "vacia"
        vacia.mkdir()
        with self.assertRaisesRegex(ValueError, "no contiene YAML"):
            ej.reubicar_pruebas(ejecucion, vacia)
        antes = (self.carpeta / ej.ARCHIVO_EJECUCION).read_bytes()
        ej.reubicar_pruebas(ejecucion, nueva)
        self.assertEqual(ej.carpeta_pruebas(ejecucion), nueva.resolve())
        ej.guardar(ejecucion)
        # La carpeta de este PC no se guarda en la ejecución: la de otra persona no se pisa.
        self.assertEqual((self.carpeta / ej.ARCHIVO_EJECUCION).read_bytes(), antes)
        self.assertIsNone(ej.carpeta_pruebas(ej.cargar(self.carpeta)))

    def test_reconoce_sus_pruebas_en_otra_carpeta(self):
        ejecucion = self.crear()
        copia = self.tmp / "clon" / "pruebas"
        shutil.copytree(self.pruebas, copia)
        self.assertTrue(ej.tiene_sus_pruebas(ejecucion, copia))
        capitulo = read_chapters(copia)[0]
        capitulo["data"]["contenido"] = []           # le faltan pruebas: no es la suya
        save_chapter(capitulo)
        self.assertFalse(ej.tiene_sus_pruebas(ejecucion, copia))


class CongelarTests(EjecucionBase):
    def cerrar_todo(self, filtro):
        ejecucion = self.crear(filtro)
        for test_id in ej.ids_de(ejecucion):
            ej.registrar(ejecucion, test_id, "ok")
        ej.cerrar(ejecucion, self.capitulos())
        return ej.cargar(self.carpeta)

    def test_al_cerrar_se_congelan_las_pruebas_con_secciones_e_imagenes(self):
        ejecucion = self.cerrar_todo({"etiquetas": ["puertas"]})
        congelada = self.carpeta / ej.CARPETA_CONGELADA
        capitulos = read_chapters(congelada)
        self.assertEqual([c["path"].name for c in capitulos], ["03_ejemplo_secciones_imagenes.yaml"])
        contenido = capitulos[0]["data"]["contenido"]
        self.assertEqual([e.get("seccion") or e["prueba"]["id"] for e in contenido],
                         ["Puertas", "Apertura", "DEMO-PUER-001", "Operador de puertas A",
                          "Reapertura por obstáculo", "DEMO-PUER-002"])
        self.assertTrue((congelada / "imagenes" / "cronograma_puertas.png").is_file())
        self.assertTrue((congelada / "imagenes" / "foto_banco_hil.jpg").is_file())
        self.assertFalse((congelada / "imagenes" / "trama_can.png").exists())
        self.assertEqual(ej.capitulos_de(ejecucion)[0]["path"].parent, congelada)

    def test_cambiar_las_pruebas_despues_no_afecta_a_una_ejecucion_cerrada(self):
        ejecucion = self.cerrar_todo({"etiquetas": ["smoke"]})
        capitulo = self.capitulos()[0]
        capitulo["data"]["contenido"][3]["prueba"]["titulo"] = "Otro título"
        save_chapter(capitulo)
        self.assertEqual(ej.resumen(ejecucion, ej.capitulos_de(ejecucion))["problemas"], [])


class InformeTests(EjecucionBase):
    def textos(self, ruta):
        doc = Document(ruta)
        celdas = [c.text for t in doc.tables for r in t.rows for c in r.cells]
        return [p.text for p in doc.paragraphs], celdas

    def test_informe_de_una_ejecucion_abierta_es_un_borrador(self):
        ejecucion = self.crear({"etiquetas": ["smoke"]})
        ej.registrar(ejecucion, "DEMO-LLAM-001", "no_ok", observaciones="No para en P2", incidencia="BUG-9")
        ej.registrar(ejecucion, "DEMO-LLAM-001", "ok", observaciones="Repetida")
        datos = informe_docx.generar_informe(ejecucion, self.tmp / "informe")
        self.assertTrue(datos["archivo"].endswith(".docx"))
        self.assertEqual(datos["pruebas"], 4)
        parrafos, celdas = self.textos(datos["archivo"])
        self.assertIn("Informe de ejecución", parrafos)
        self.assertTrue(any(p.startswith("BORRADOR") for p in parrafos))
        self.assertTrue(any(c.startswith("Resultado de la ejecución: OK") for c in celdas))
        self.assertTrue(any("Intentos anteriores (1)" in c and "BUG-9" in c for c in celdas))
        self.assertTrue(any(c.startswith("Resultado de la ejecución: Pendiente") for c in celdas))

    def test_informe_de_una_ejecucion_cerrada_usa_las_pruebas_congeladas(self):
        ejecucion = self.crear({"etiquetas": ["manual"]})
        ej.registrar(ejecucion, "DEMO-VAR-002", "ok")
        ej.cerrar(ejecucion, self.capitulos())
        capitulo = self.capitulos()[3]
        for e in capitulo["data"]["contenido"]:
            if isinstance(e, dict) and e.get("prueba", {}).get("id") == "DEMO-VAR-002":
                e["prueba"]["titulo"] = "Cambiado después de cerrar"
        save_chapter(capitulo)
        datos = informe_docx.generar_informe(ej.cargar(self.carpeta), self.tmp / "informe.docx")
        parrafos, celdas = self.textos(datos["archivo"])
        self.assertFalse(any(p.startswith("BORRADOR") for p in parrafos))
        self.assertFalse(any("Cambiado después de cerrar" in c for c in celdas))
        self.assertTrue(any("procedimiento de rescate" in c for c in celdas))

def ok(testcase, respuesta):
    testcase.assertTrue(respuesta["ok"], respuesta.get("error"))
    return respuesta["datos"]


class ApiTests(EjecucionBase):
    def setUp(self):
        super().setUp()
        anterior = os.environ.get("AZTARNA_AJUSTES")          # los ajustes del PC, en la carpeta temporal
        os.environ["AZTARNA_AJUSTES"] = str(self.tmp / "ajustes.json")
        self.addCleanup(lambda: os.environ.__setitem__("AZTARNA_AJUSTES", anterior) if anterior else os.environ.pop("AZTARNA_AJUSTES", None))
        self.api = EditorEjecucionApi(self.pruebas)
        ok(self, self.api.iniciar())

    def crear_por_api(self):
        return ok(self, self.api.ejecucion_crear("Regresión", {"etiquetas": ["smoke"]},
                                                  {"objeto": "FW 2.3"}, str(self.carpeta)))

    def test_previsualizar_y_crear(self):
        vista = ok(self, self.api.ejecucion_previsualizar({"etiquetas": ["smoke"]}))
        self.assertEqual([p["id"] for p in vista], ["DEMO-LLAM-001", "DEMO-NIV-001", "DEMO-PUER-001", "DEMO-SEG-001"])
        estado = self.crear_por_api()
        self.assertEqual(estado["datos"]["objeto"], "FW 2.3")
        self.assertEqual(estado["resumen"]["totales"]["pendiente"], 4)
        self.assertIn("smoke", ok(self, self.api.ejecucion_opciones())["etiquetas"])

    def test_registrar_y_corregir_notas_sin_otro_intento(self):
        self.crear_por_api()
        ok(self, self.api.ejecucion_registrar("DEMO-LLAM-001", "no_ok", "Falla", "BUG-1", [], "Unai"))
        estado = ok(self, self.api.ejecucion_registrar("DEMO-LLAM-001", "no_ok", "Falla en P2", "BUG-1", [], "Unai"))
        self.assertEqual(estado["resumen"]["totales"]["no_ok"], 1)
        datos = ok(self, self.api.ejecucion_resultado("DEMO-LLAM-001"))
        self.assertEqual(datos["resultado"]["observaciones"], "Falla en P2")
        self.assertNotIn("intentos", datos["resultado"])
        self.assertTrue(datos["en_ejecucion"])
        self.assertFalse(datos["prueba_cambiada"])

    def test_detecta_prueba_cambiada_despues_de_ejecutar(self):
        self.crear_por_api()
        ok(self, self.api.ejecucion_registrar("DEMO-LLAM-001", "ok"))
        ok(self, self.api.guardar_prueba(0, 3, {"titulo": "Título nuevo"}))
        self.assertTrue(ok(self, self.api.ejecucion_resultado("DEMO-LLAM-001"))["prueba_cambiada"])

    def test_evidencias_se_guardan_en_la_ejecucion(self):
        self.crear_por_api()
        buffer = io.BytesIO()
        Image.new("RGB", (40, 20), "red").save(buffer, format="PNG")
        datos = "data:image/png;base64," + base64.b64encode(buffer.getvalue()).decode()
        ruta = ok(self, self.api.ejecucion_evidencia_base64("DEMO-LLAM-001", "captura.png", datos))["ruta"]
        self.assertEqual(ruta, "evidencias/DEMO-LLAM-001_01.png")
        log = base64.b64encode(b"linea 1").decode()
        self.assertEqual(ok(self, self.api.ejecucion_evidencia_base64("DEMO-LLAM-001", "bus.log", log))["ruta"],
                         "evidencias/DEMO-LLAM-001_02.log")
        self.assertEqual(ok(self, self.api.ejecucion_imagen(ruta, 100))["ancho"], 40)
        self.assertFalse(self.api.ejecucion_imagen("../campaña.yaml")["ok"])
        ok(self, self.api.ejecucion_registrar("DEMO-LLAM-001", "ok", "", "", [ruta]))
        self.assertEqual(ej.cargar(self.carpeta)["resultados"]["DEMO-LLAM-001"]["evidencias"], [ruta])

    def test_abrir_ejecucion_de_otra_carpeta_carga_sus_pruebas(self):
        self.crear_por_api()
        otra = EditorEjecucionApi()
        estado = ok(self, otra.ejecucion_abrir(str(self.carpeta)))
        self.assertEqual(estado["carpeta_pruebas"]["carpeta"], str(self.pruebas.resolve()))
        self.assertEqual(estado["resumen"]["total"], 4)
        self.assertIsNone(ok(self, otra.ejecucion_abrir(str(self.carpeta)))["carpeta_pruebas"])

    def test_la_ejecucion_de_ejemplo_abre_sin_avisos(self):
        ejemplo = Path(__file__).resolve().parent.parent / "ejemplos" / "ejecuciones" / "regresion_v2.3"
        estado = ok(self, EditorEjecucionApi().ejecucion_abrir(str(ejemplo)))
        self.assertEqual(estado["resumen"]["total"], 10)
        self.assertEqual(estado["resumen"]["problemas"], [],
                         "Si cambias las pruebas de ejemplo, vuelve a generar la ejecución de ejemplo")

    def test_cerrar_por_api_deja_las_pruebas_en_solo_lectura(self):
        estado = ok(self, self.api.ejecucion_crear("Manual", {"etiquetas": ["manual"]}, {}, str(self.carpeta)))
        self.assertEqual(estado["resumen"]["total"], 1)
        self.assertFalse(self.api.ejecucion_cerrar()["ok"])            # queda una pendiente
        ok(self, self.api.ejecucion_registrar("DEMO-VAR-002", "ok"))
        cerrada = ok(self, self.api.ejecucion_cerrar())
        self.assertTrue(cerrada["cerrada"])
        self.assertTrue(cerrada["carpeta_pruebas"]["solo_lectura"])
        fallo = self.api.guardar_prueba(0, 1, {"titulo": "x"})
        self.assertFalse(fallo["ok"])
        self.assertIn("copia congelada", fallo["error"])
        informe = ok(self, self.api.ejecucion_informe(str(self.tmp / "inf.docx")))
        self.assertTrue(Path(informe["archivo"]).is_file())
        # al abrir otra carpeta se puede volver a editar
        ok(self, self.api.abrir_carpeta(str(self.pruebas)))
        ok(self, self.api.guardar_prueba(0, 3, {"titulo": "Editable de nuevo"}))

    def test_abrir_sin_encontrar_las_pruebas_y_reubicarlas(self):
        self.crear_por_api()
        nueva = self.tmp / "movidas"
        shutil.move(str(self.pruebas), nueva)
        otra = EditorEjecucionApi()
        datos = ok(self, otra.ejecucion_abrir(str(self.carpeta)))
        self.assertTrue(datos["faltan_pruebas"])
        estado = ok(self, otra.ejecucion_reubicar(str(self.carpeta), str(nueva)))
        self.assertEqual(estado["carpeta_pruebas"]["carpeta"], str(nueva.resolve()))
        # Este PC lo recuerda: la próxima vez se abre sin preguntar.
        datos = ok(self, EditorEjecucionApi().ejecucion_abrir(str(self.carpeta)))
        self.assertNotIn("faltan_pruebas", datos)
        self.assertEqual(datos["carpeta_pruebas"]["carpeta"], str(nueva.resolve()))

    def test_abrir_la_de_otro_pc_con_las_pruebas_abiertas(self):
        """Cada uno con su clon de las pruebas: la ejecución de otro se abre con las tuyas."""
        self.crear_por_api()
        mias = self.tmp / "mi_clon" / "pruebas"
        shutil.copytree(self.pruebas, mias)
        shutil.rmtree(self.pruebas)                 # la ruta guardada (la del otro PC) no existe aquí
        api = EditorEjecucionApi(mias)
        ok(self, api.iniciar())
        datos = ok(self, api.ejecucion_abrir(str(self.carpeta)))
        self.assertNotIn("faltan_pruebas", datos)
        ok(self, api.ejecucion_registrar("DEMO-LLAM-001", "ok"))
        self.assertEqual(ej.estado_de(ej.cargar(self.carpeta), "DEMO-LLAM-001"), "ok")
        self.assertNotIn("mi_clon", (self.carpeta / ej.ARCHIVO_EJECUCION).read_text(encoding="utf-8"))

    def test_sin_ejecucion_abierta(self):
        self.assertIsNone(ok(self, self.api.ejecucion_estado()))
        self.assertFalse(self.api.ejecucion_registrar("DEMO-LLAM-001", "ok")["ok"])


if __name__ == "__main__":
    unittest.main()
