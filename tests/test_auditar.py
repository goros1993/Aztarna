"""Pruebas de auditar.py: detectar lo que se pierde al importar el Word."""
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from docx import Document

from migracion import auditar
from migracion import convertir
from tests.test_conversion import documento_mixto


def documento_con_problemas():
    """Word con un poco de todo: formato antiguo, campos que se eliminan, valores
    que no encajan, texto que el importador no sabe colocar y una tabla anidada."""
    doc = Document()
    doc.add_heading("Capítulo", level=1)
    doc.add_paragraph("Introducción del capítulo.")
    doc.add_heading("Prueba antigua", level=2)
    antigua = doc.add_table(rows=5, cols=2)
    antigua.cell(0, 0).text, antigua.cell(0, 1).text = "OLD_001", "Formato antiguo"
    antigua.cell(1, 0).text = "Descripción: Comprobar la alarma."
    antigua.cell(1, 1).text = "Componente: CPU_A"
    antigua.cell(2, 0).text = "Situación inicial:\nReposo.\nPasos:\nProvocar fallo.\nSituación final:\nSe genera alarma."
    antigua.cell(3, 0).text = "Entorno de pruebas: HIL"
    antigua.cell(3, 1).text = "Resultado esperado: Alarma visible."
    antigua.cell(4, 0).text = "Texto suelto que nadie recoge"
    doc.add_heading("Prueba con anidada", level=2)
    anidada = doc.add_table(rows=2, cols=1)
    anidada.cell(0, 0).text = "NEST_001"
    anidada.cell(1, 0).text = "Pasos:\nUno."
    anidada.cell(1, 0).add_table(rows=1, cols=1).cell(0, 0).text = "Dentro"
    doc.add_heading("Prueba buena", level=2)
    buena = doc.add_table(rows=3, cols=2)
    buena.cell(0, 0).text, buena.cell(0, 1).text = "GOOD_001", "Formato nuevo"
    buena.cell(1, 0).text = "Ensayo:\nHacerlo."
    buena.cell(2, 0).text = "Entorno de pruebas: A/M"
    buena.cell(2, 1).text = "Automatizable: Sí"
    return doc


class AuditarTests(unittest.TestCase):
    def convertir_y_auditar(self, doc, tmp):
        ruta = Path(tmp) / "doc.docx"
        doc.save(ruta)
        capitulos = convertir.extraer_capitulos(Document(ruta), [])
        salida = Path(tmp) / "yaml"
        salida.mkdir()
        from ruamel.yaml import YAML
        yaml = YAML()
        indice = []
        for i, cap in enumerate(capitulos, 1):
            with open(salida / f"{i:02d}.yaml", "w", encoding="utf-8") as f:
                yaml.dump(cap, f)
            indice.append({"capitulo": cap["capitulo"], "archivo": f"{i:02d}.yaml"})
        with open(salida / "indice.yaml", "w", encoding="utf-8") as f:
            yaml.dump(indice, f)
        return auditar.auditar(ruta, salida)

    def test_documento_correcto(self):
        with TemporaryDirectory() as tmp:
            datos = self.convertir_y_auditar(documento_mixto(), tmp)
        r = datos["recuentos"]
        self.assertEqual((r["pruebas_word"], r["pruebas_yaml"], r["tablas_word"], r["tablas_yaml"]), (3, 3, 4, 4))
        self.assertEqual(r["ids_repetidos"], ["LLAM_001"])
        self.assertFalse([t for t in datos["tablas"] if t["estado"] == "perdida"])
        self.assertEqual(datos["parrafos"], [])

    def test_clasifica_lo_que_falta(self):
        with TemporaryDirectory() as tmp:
            datos = self.convertir_y_auditar(documento_con_problemas(), tmp)
        tablas = {t["id"]: t for t in datos["tablas"]}
        antigua = tablas["OLD_001"]
        self.assertEqual(antigua["estado"], "perdida")
        self.assertEqual(antigua["perdidas"], ["nadie", "que", "recoge", "suelto", "texto"])
        self.assertEqual(antigua["lugares"], [[5, 1, "sin etiqueta"]])
        self.assertIn("cpu_a", antigua["eliminadas"])            # componente: eliminado por regla
        self.assertIn("visible", antigua["eliminadas"])           # resultado esperado
        self.assertEqual(antigua["descartadas"], ["hil"])         # entorno que no es una opción
        self.assertNotIn("genera", antigua["eliminadas"])         # situación final -> ensayo: no se pierde
        anidada = tablas["NEST_001"]
        self.assertEqual(anidada["estado"], "perdida")
        self.assertTrue(any("tablas dentro de celdas" in a for a in anidada["avisos"]))
        # A/M -> Automático, Maqueta y Sí -> si no cuentan como eliminados.
        self.assertEqual(tablas["GOOD_001"]["estado"], "ok")

    def test_informe_html_y_resumen_sin_contenido(self):
        with TemporaryDirectory() as tmp:
            datos = self.convertir_y_auditar(documento_con_problemas(), tmp)
        html = auditar.informe_html(datos)
        self.assertIn("Texto suelto que nadie recoge", html)      # el informe sí lleva el texto…
        self.assertNotIn("{{datos}}", html)
        resumen = auditar.resumen(datos)
        self.assertNotIn("nadie recoge", resumen)                 # …el resumen no
        self.assertIn("fila  5 col 1  sin etiqueta", resumen)

    def test_parrafos_que_no_estan_en_los_yaml(self):
        with TemporaryDirectory() as tmp:
            doc = documento_con_problemas()
            ruta = Path(tmp) / "doc.docx"
            datos = self.convertir_y_auditar(doc, tmp)
            self.assertEqual(datos["parrafos"], [])
            # Si se borra texto del YAML después de importar, se detecta.
            yaml = (Path(tmp) / "yaml" / "01.yaml")
            yaml.write_text(yaml.read_text(encoding="utf-8").replace("Introducción del capítulo.", "Otra cosa."), encoding="utf-8")
            datos = auditar.auditar(ruta, Path(tmp) / "yaml")
        self.assertEqual(datos["parrafos"][0]["perdidas"], ["capitulo", "del", "introduccion"])


class TextoDelWordTests(unittest.TestCase):
    """Casos que daban falsas pérdidas o perdían texto de verdad."""

    def tabla(self):
        from docx.oxml import OxmlElement
        from docx.shared import RGBColor
        doc = Document()
        t = doc.add_table(rows=4, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "V_001", "Casos raros"
        celda = t.cell(1, 0).merge(t.cell(2, 0))           # unida en vertical
        celda.paragraphs[0].add_run("Pasos:")
        p = celda.add_paragraph()
        p.add_run("Ascen")
        p.add_run("sor").bold = True                       # palabra con parte en negrita
        p.add_run(" en ")
        p.add_run("P").font.color.rgb = RGBColor(0xFF, 0, 0)
        p.add_run("2.")
        p = celda.add_paragraph()
        p.add_run("Ver ")
        for contenedor, texto in (("w:hyperlink", "documento"), ("w:ins", " insertado")):
            envoltura, run, t_ = OxmlElement(contenedor), OxmlElement("w:r"), OxmlElement("w:t")
            t_.text = texto
            t_.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            run.append(t_)
            envoltura.append(run)
            p._p.append(envoltura)
        t.cell(1, 1).text = "Automatizable: Sí"
        t.cell(3, 0).text = "Fecha: 1/1/2026"
        return t

    def test_el_importador_lee_todo_el_texto_una_vez(self):
        prueba = convertir.parsear(self.tabla(), [])
        self.assertEqual(prueba["pasos"], ["Ascen**sor** en [[rojo|P]]2.", "Ver documento insertado"])

    def test_sin_falsas_perdidas(self):
        t = self.tabla()
        r = auditar.auditar_prueba(t, convertir.parsear(t, []))
        self.assertEqual((r["perdidas"], r["eliminadas"]), ([], []))
        filas = auditar.celdas(t)
        self.assertTrue(filas[2][0]["unida"])
        self.assertEqual(filas[2][0]["texto"], "")
