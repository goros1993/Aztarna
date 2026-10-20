"""Pruebas de la validación de una versión (validacion_core.py)."""
import shutil
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from app import ejecucion_core as ej
from app import validacion_core as vc
from app.validacion_api import EditorValidacionApi


def ok(testcase, resultado):
    testcase.assertTrue(resultado["ok"], resultado.get("error"))
    return resultado["datos"]

EJEMPLOS = Path(__file__).resolve().parent.parent / "ejemplos" / "yaml"


class ValidacionBase(unittest.TestCase):
    def setUp(self):
        self._tmp = TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)
        shutil.copytree(EJEMPLOS, self.tmp / "yaml")
        self.pruebas = self.tmp / "yaml" / "pruebas"
        self.val = vc.crear(self.tmp / "validacion", self.tmp / "yaml", "Validación v2.3", "2.3.0",
                            {"incluir_ids": ["DEMO-LLAM-001", "DEMO-NIV-001", "DEMO-PUER-002", "DEMO-MANT-002"],
                             "excluir_ids": [], "capitulos": ["__ninguno__"]})

    def ejecucion(self, nombre, entorno, ids, anadir=True):
        c = ej.crear(self.tmp / "ejec" / nombre, self.pruebas, nombre, {"incluir_ids": ids, "capitulos": ["__ninguno__"]},
                     {"entorno_pruebas": entorno} if entorno else {})
        if anadir:
            vc.anadir_ejecucion(self.val, c["carpeta"])
        return c

    def prueba(self, datos, test_id):
        return next(p for p in datos["pruebas"] if p["id"] == test_id)


class ArchivoTests(ValidacionBase):
    def test_crear_cargar_y_rutas_relativas(self):
        datos = vc.cargar(self.tmp / "validacion")["datos"]
        self.assertEqual((datos["nombre"], datos["version"], datos["pruebas_yaml"]), ("Validación v2.3", "2.3.0", "../yaml/pruebas"))
        self.assertEqual(vc.carpeta_pruebas(self.val), self.pruebas.resolve())

    def test_anadir_quitar_y_no_repetir(self):
        c = self.ejecucion("torre", "Torre", ["DEMO-NIV-001"])
        with self.assertRaises(ValueError):
            vc.anadir_ejecucion(self.val, c["carpeta"])
        with self.assertRaises(ValueError):
            vc.anadir_ejecucion(self.val, self.tmp)
        self.assertEqual(vc.cargar(self.tmp / "validacion")["datos"]["ejecuciones"][0]["ruta"], "../ejec/torre")
        vc.quitar_ejecucion(self.val, 0)
        self.assertEqual(vc.cargar(self.tmp / "validacion")["datos"]["ejecuciones"], [])

    def test_anadir_todas_las_de_una_carpeta(self):
        for nombre, entorno in (("ana", "Maqueta"), ("luis", "Torre")):
            ej.crear(self.tmp / "compartida" / nombre, self.pruebas, nombre, {"incluir_ids": ["DEMO-NIV-001"], "capitulos": ["x"]}, {"entorno_pruebas": entorno})
        ej.crear(self.tmp / "compartida" / "pc_de_eva" / "torre", self.pruebas, "eva", {"incluir_ids": ["DEMO-NIV-001"], "capitulos": ["x"]}, {"entorno_pruebas": "Torre"})
        (self.tmp / "compartida" / "otra_cosa").mkdir()
        r = vc.anadir_ejecuciones(self.val, self.tmp / "compartida")
        self.assertEqual(r, {"anadidas": ["ana", "luis", "torre"], "ya_estaban": []})
        r = vc.anadir_ejecuciones(self.val, self.tmp / "compartida")
        self.assertEqual(r["anadidas"], [])
        self.assertEqual(len(r["ya_estaban"]), 3)
        with self.assertRaises(ValueError):
            vc.anadir_ejecuciones(self.val, self.tmp / "compartida" / "otra_cosa")

    def test_editar(self):
        vc.editar(self.val, {"nombre": "Otra", "version": ""})
        datos = vc.cargar(self.tmp / "validacion")["datos"]
        self.assertEqual(datos["nombre"], "Otra")
        self.assertNotIn("version", datos)
        with self.assertRaises(ValueError):
            vc.editar(self.val, {"nombre": " "})


class ConsolidarTests(ValidacionBase):
    def test_sin_ejecuciones_todo_sin_planificar(self):
        d = vc.consolidar(self.val)
        self.assertEqual(d["total"], 4)
        self.assertEqual(d["totales"]["pendiente"], 4)
        self.assertTrue(d["borrador"])
        # DEMO-NIV-001 pide Automático y Torre
        niv = self.prueba(d, "DEMO-NIV-001")
        self.assertEqual(niv["requeridos"], ["Automático", "Torre"])
        self.assertEqual({e: c["estado"] for e, c in niv["entornos"].items()}, {"Automático": "sin_planificar", "Torre": "sin_planificar"})

    def test_validada_cuando_esta_ok_en_todos_sus_entornos(self):
        auto = self.ejecucion("auto", "Automático", ["DEMO-NIV-001"])
        torre = self.ejecucion("torre", "Torre", ["DEMO-NIV-001"])
        ej.registrar(auto, "DEMO-NIV-001", "ok", fecha="2026-11-01T10:00")
        d = vc.consolidar(self.val)
        self.assertEqual(self.prueba(d, "DEMO-NIV-001")["estado"], "pendiente")
        self.assertEqual(d["por_entorno"]["Torre"]["pendiente"], 1)          # planificada, sin ejecutar
        ej.registrar(torre, "DEMO-NIV-001", "ok", fecha="2026-11-02T10:00")
        self.assertEqual(self.prueba(vc.consolidar(self.val), "DEMO-NIV-001")["estado"], "validada")

    def test_en_el_mismo_entorno_vale_el_mas_reciente(self):
        lunes = self.ejecucion("lunes", "Torre", ["DEMO-PUER-002"])
        jueves = self.ejecucion("jueves", "Torre", ["DEMO-PUER-002"])
        ej.registrar(jueves, "DEMO-PUER-002", "ok", fecha="2026-11-06T10:00")
        ej.registrar(lunes, "DEMO-PUER-002", "no_ok", observaciones="falla", fecha="2026-11-03T10:00")
        p = self.prueba(vc.consolidar(self.val), "DEMO-PUER-002")
        self.assertEqual(p["estado"], "validada")
        self.assertEqual([r["estado"] for r in p["entornos"]["Torre"]["historial"]], ["ok", "no_ok"])

    def test_otro_entorno_se_muestra_pero_no_cuenta(self):
        maqueta = self.ejecucion("maqueta", "Maqueta", ["DEMO-PUER-002"])
        ej.registrar(maqueta, "DEMO-PUER-002", "no_ok", observaciones="x", fecha="2026-11-03T10:00")
        d = vc.consolidar(self.val)
        p = self.prueba(d, "DEMO-PUER-002")
        self.assertEqual(p["estado"], "pendiente")                          # sigue faltando Torre
        self.assertFalse(p["entornos"]["Maqueta"]["cuenta"])
        self.assertTrue(any("un entorno que la prueba no pide" in a["mensaje"] for a in d["avisos"]))

    def test_peor_estado_entre_entornos(self):
        auto = self.ejecucion("auto", "Automático", ["DEMO-NIV-001"])
        torre = self.ejecucion("torre", "Torre", ["DEMO-NIV-001"])
        ej.registrar(auto, "DEMO-NIV-001", "ok", fecha="2026-11-01T10:00")
        ej.registrar(torre, "DEMO-NIV-001", "bloqueada", observaciones="banco", fecha="2026-11-01T11:00")
        self.assertEqual(self.prueba(vc.consolidar(self.val), "DEMO-NIV-001")["estado"], "bloqueada")
        ej.registrar(torre, "DEMO-NIV-001", "no_aplica", fecha="2026-11-02T11:00")
        self.assertEqual(self.prueba(vc.consolidar(self.val), "DEMO-NIV-001")["estado"], "validada")

    def test_prueba_sin_entorno_vale_cualquiera_con_aviso(self):
        texto = (self.pruebas / "04_ejemplo_edicion_variantes.yaml").read_text(encoding="utf-8")
        self.assertIn("DEMO-MANT-002", texto)
        maqueta = self.ejecucion("maqueta", "Maqueta", ["DEMO-MANT-002"])
        ej.registrar(maqueta, "DEMO-MANT-002", "ok", fecha="2026-11-01T10:00")
        d = vc.consolidar(self.val)
        p = self.prueba(d, "DEMO-MANT-002")
        if p["requeridos"]:
            self.skipTest("el ejemplo ya indica su entorno")
        self.assertEqual(p["estado"], "validada")
        self.assertIn(vc.CUALQUIERA, d["entornos"])
        self.assertTrue(any(a["id"] == "DEMO-MANT-002" and "no indica su entorno" in a["mensaje"] for a in d["avisos"]))

    def test_ejecucion_sin_entorno_no_cuenta_y_avisa(self):
        c = self.ejecucion("vieja", "", ["DEMO-PUER-002"])
        ej.registrar(c, "DEMO-PUER-002", "ok", fecha="2026-11-01T10:00")
        d = vc.consolidar(self.val)
        self.assertEqual(self.prueba(d, "DEMO-PUER-002")["estado"], "pendiente")
        self.assertTrue(any("no indica su entorno de pruebas" in a["mensaje"] for a in d["avisos"]))

    def test_prueba_cambiada_despues_de_ejecutar(self):
        torre = self.ejecucion("torre", "Torre", ["DEMO-PUER-002"])
        ej.registrar(torre, "DEMO-PUER-002", "ok", fecha="2026-11-01T10:00", huella_actual="otra")
        d = vc.consolidar(self.val)
        self.assertTrue(any(a["id"] == "DEMO-PUER-002" and "ha cambiado" in a["mensaje"] for a in d["avisos"]))
        self.assertEqual(self.prueba(d, "DEMO-PUER-002")["cambiada_en"], ["Torre"])
        self.assertEqual(self.prueba(d, "DEMO-NIV-001")["cambiada_en"], [])

    def test_dar_por_revisada_quita_el_aviso_hasta_que_vuelva_a_cambiar(self):
        torre = self.ejecucion("torre", "Torre", ["DEMO-PUER-002"])
        ej.registrar(torre, "DEMO-PUER-002", "ok", fecha="2026-11-01T10:00", huella_actual="otra")
        actual = ej.huella(ej.indice_por_id(ej.read_chapters(self.pruebas))["DEMO-PUER-002"]["prueba"])
        vc.dar_por_revisada(self.val, "DEMO-PUER-002", actual)
        self.assertEqual(vc.cargar(self.tmp / "validacion")["datos"]["revisadas"], {"DEMO-PUER-002": actual})
        d = vc.consolidar(vc.cargar(self.tmp / "validacion"))
        self.assertEqual(self.prueba(d, "DEMO-PUER-002")["cambiada_en"], [])
        self.assertFalse(any(a["id"] == "DEMO-PUER-002" and "ha cambiado" in a["mensaje"] for a in d["avisos"]))
        self.assertEqual(self.prueba(d, "DEMO-PUER-002")["estado"], "validada")
        # Si vuelve a cambiar, vuelve a avisar.
        vc.dar_por_revisada(self.val, "DEMO-PUER-002", "anterior")
        self.assertEqual(self.prueba(vc.consolidar(self.val), "DEMO-PUER-002")["cambiada_en"], ["Torre"])

    def test_borrador_hasta_cerrar_todas(self):
        torre = self.ejecucion("torre", "Torre", ["DEMO-PUER-002"])
        ej.registrar(torre, "DEMO-PUER-002", "ok", fecha="2026-11-01T10:00")
        self.assertTrue(vc.consolidar(self.val)["borrador"])
        ej.cerrar(torre, ej.read_chapters(self.pruebas))
        self.assertFalse(vc.consolidar(self.val)["borrador"])

    def test_ejecucion_que_no_se_encuentra(self):
        c = self.ejecucion("torre", "Torre", ["DEMO-PUER-002"])
        shutil.rmtree(c["carpeta"])
        d = vc.consolidar(self.val)
        self.assertFalse(d["ejecuciones"][0]["encontrada"])
        self.assertEqual(d["avisos"][0]["tipo"], "error")


class EjemploTests(unittest.TestCase):
    def test_la_validacion_de_ejemplo(self):
        """ejemplos/validaciones/validacion_v2.3: Robot (Automático), torre cerrada, repetición y la de maqueta."""
        carpeta = EJEMPLOS.parent / "validaciones" / "validacion_v2.3"
        d = vc.consolidar(vc.cargar(carpeta))
        self.assertTrue(all(e["encontrada"] for e in d["ejecuciones"]))
        self.assertEqual([e["entorno_pruebas"] for e in d["ejecuciones"]], ["Maqueta", "Automático", "Torre", "Torre"])
        self.assertEqual(d["totales"], {"no_ok": 1, "bloqueada": 1, "pendiente": 3, "validada": 4, "no_aplica": 1})
        puer = next(p for p in d["pruebas"] if p["id"] == "DEMO-PUER-002")
        self.assertEqual([r["estado"] for r in puer["entornos"]["Torre"]["historial"]], ["ok", "no_ok"])   # vale la repetición
        self.assertEqual([a["id"] for a in d["avisos"]], ["DEMO-NIV-001"])


class InformeTests(ValidacionBase):
    def test_informe_con_matriz_y_no_validadas(self):
        from docx import Document
        from app import informe_validacion
        torre = self.ejecucion("torre", "Torre", ["DEMO-PUER-002", "DEMO-NIV-001"])
        ej.registrar(torre, "DEMO-PUER-002", "no_ok", observaciones="No reabre", fecha="2026-11-01T10:00")
        salida = self.tmp / "informe.docx"
        datos = informe_validacion.generar_informe(self.val, salida)
        self.assertEqual(datos["pruebas"], 4)
        doc = Document(str(salida))
        textos = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("BORRADOR", textos)
        matrices = [t for t in doc.tables if t.rows[0].cells[-1].text == "Versión"]
        self.assertTrue(matrices)
        fila = next(r for t in matrices for r in t.rows if r.cells[0].text == "DEMO-PUER-002")
        self.assertEqual(fila.cells[-1].text, "NO OK")
        detalle = "\n".join(c.text for tabla in doc.tables for fila in tabla.rows for c in fila.cells)
        self.assertIn("No reabre", detalle)
        self.assertIn("Detalle por capítulo", textos)

    def test_resumen_por_capitulo_y_solo_resumen(self):
        from docx import Document
        from app import informe_validacion
        torre = self.ejecucion("torre", "Torre", ["DEMO-PUER-002", "DEMO-NIV-001"])
        ej.registrar(torre, "DEMO-PUER-002", "no_ok", observaciones="No reabre", fecha="2026-11-01T10:00")
        capitulos = {p["capitulo"] for p in vc.consolidar(self.val)["pruebas"]}
        salida = self.tmp / "resumen.docx"
        informe_validacion.generar_informe(self.val, salida, detalle=False)
        doc = Document(str(salida))
        textos = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("Detalle por capítulo", textos)
        self.assertFalse([t for t in doc.tables if t.rows[0].cells[-1].text == "Versión"])
        if len(capitulos) > 1:
            tabla = next(t for t in doc.tables if t.rows[0].cells[0].text == "Capítulo")
            self.assertEqual({r.cells[0].text for r in tabla.rows[1:]}, capitulos)
            self.assertEqual(sum(int(r.cells[1].text) for r in tabla.rows[1:]), 4)


class ApiTests(unittest.TestCase):
    def setUp(self):
        self._tmp = TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)
        shutil.copytree(EJEMPLOS, self.tmp / "yaml")
        self.api = EditorValidacionApi(self.tmp / "yaml")
        ok(self, self.api.iniciar())

    def test_crear_nueva_ejecucion_dentro_y_consolidar(self):
        v = ok(self, self.api.validacion_crear("Validación v2.3", "2.3.0", {"etiquetas": ["regresion"]}, str(self.tmp / "val")))
        self.assertEqual((v["total"], v["ejecuciones"]), (10, []))
        e = ok(self, self.api.validacion_nueva_ejecucion("En Torre", {"etiquetas": ["regresion"], "entorno": ["Torre"]},
                                                          {"entorno_pruebas": "Torre", "entorno": "Torre 2"}))
        self.assertTrue(e["carpeta"].startswith(str(self.tmp / "val")))
        ok(self, self.api.ejecucion_registrar("DEMO-PUER-001", "ok"))
        v = ok(self, self.api.validacion_estado())
        puer = next(p for p in v["pruebas"] if p["id"] == "DEMO-PUER-001")
        self.assertEqual(puer["estado"], "validada")
        self.assertEqual((v["ejecuciones"][0]["entorno_pruebas"], v["ejecuciones"][0]["banco"]), ("Torre", "Torre 2"))

    def test_abrir_quitar_y_volver_desde_una_ejecucion_cerrada(self):
        ok(self, self.api.validacion_crear("V", "", {"incluir_ids": ["DEMO-PUER-001"], "capitulos": ["__x__"]}, str(self.tmp / "val")))
        ok(self, self.api.validacion_nueva_ejecucion("T", {"incluir_ids": ["DEMO-PUER-001"], "capitulos": ["__x__"]}, {"entorno_pruebas": "Torre"}))
        ok(self, self.api.ejecucion_registrar("DEMO-PUER-001", "ok"))
        ok(self, self.api.ejecucion_cerrar())
        self.assertTrue(ok(self, self.api.iniciar())["solo_lectura"])       # pruebas congeladas
        r = ok(self, self.api.ejecucion_soltar())
        self.assertFalse(r["carpeta_pruebas"]["solo_lectura"])              # vuelve a las de la validación
        self.assertFalse(ok(self, self.api.validacion_estado())["borrador"])
        otra = EditorValidacionApi(self.tmp / "yaml")
        v = ok(self, otra.validacion_abrir(str(self.tmp / "val")))
        self.assertEqual(len(v["ejecuciones"]), 1)
        v = ok(self, otra.validacion_quitar_ejecucion(0))
        self.assertEqual(v["ejecuciones"], [])

    def test_ejecucion_abierta_por_su_cuenta_se_marca_y_se_puede_anadir(self):
        """Con una validación abierta, una ejecución creada o abierta aparte no cuenta hasta añadirla."""
        ok(self, self.api.validacion_crear("V", "", {}, str(self.tmp / "val")))
        e = ok(self, self.api.ejecucion_crear("Suelta", {"etiquetas": ["regresion"]}, {"entorno_pruebas": "Torre"}, str(self.tmp / "suelta")))
        self.assertIs(e["en_validacion"], False)
        r = ok(self, self.api.validacion_anadir_ejecucion(str(self.tmp / "suelta")))
        self.assertEqual(r["anadidas"], ["suelta"])
        self.assertIs(r["ejecucion"]["en_validacion"], True)
        self.assertIs(ok(self, self.api.ejecucion_estado())["en_validacion"], True)
        ok(self, self.api.validacion_soltar())
        self.assertIsNone(ok(self, self.api.ejecucion_estado())["en_validacion"])

    def _guardar_prueba(self, test_id, cambios):
        p = next(p for p in ok(self, self.api.validacion_estado())["pruebas"] if p["id"] == test_id)
        capitulo = next(i for i, c in enumerate(self.api._chapters) if c["path"].name == p["archivo"])
        ok(self, self.api.guardar_prueba(capitulo, p["indice"], cambios))

    def test_cambiar_el_entorno_de_la_prueba_recoge_lo_ejecutado_y_se_puede_dar_por_revisada(self):
        """Preparada para Maqueta y pasada en Torre: al cambiarla a Torre cuenta lo de Torre."""
        ok(self, self.api.validacion_crear("V", "", {"incluir_ids": ["DEMO-LLAM-003"], "capitulos": ["__x__"]}, str(self.tmp / "val")))
        ok(self, self.api.validacion_nueva_ejecucion("T", {"incluir_ids": ["DEMO-LLAM-003"], "capitulos": ["__x__"]}, {"entorno_pruebas": "Torre"}))
        ok(self, self.api.ejecucion_registrar("DEMO-LLAM-003", "ok"))
        p = ok(self, self.api.validacion_estado())["pruebas"][0]
        self.assertEqual((p["requeridos"], p["estado"]), (["Maqueta"], "pendiente"))
        self._guardar_prueba("DEMO-LLAM-003", {"entorno": ["Torre"]})
        p = ok(self, self.api.validacion_estado())["pruebas"][0]
        self.assertEqual((p["requeridos"], p["estado"], p["cambiada_en"]), (["Torre"], "validada", ["Torre"]))
        # Desde la ejecución abierta: el resultado pasa a valer para la prueba actual.
        e = ok(self, self.api.ejecucion_dar_por_revisada("DEMO-LLAM-003"))
        self.assertEqual([x["id"] for x in e["resumen"]["pruebas"] if x["cambiada"]], [])
        self.assertFalse(ok(self, self.api.ejecucion_resultado("DEMO-LLAM-003"))["prueba_cambiada"])
        self.assertEqual(ok(self, self.api.validacion_estado())["pruebas"][0]["cambiada_en"], [])
        # Desde la validación (por ejemplo, con la ejecución ya cerrada).
        self._guardar_prueba("DEMO-LLAM-003", {"titulo": "Otro título"})
        ok(self, self.api.ejecucion_cerrar())
        ok(self, self.api.ejecucion_soltar())
        self.assertEqual(ok(self, self.api.validacion_estado())["pruebas"][0]["cambiada_en"], ["Torre"])
        v = ok(self, self.api.validacion_dar_por_revisada("DEMO-LLAM-003"))
        self.assertEqual((v["pruebas"][0]["cambiada_en"], v["pruebas"][0]["estado"]), ([], "validada"))


if __name__ == "__main__":
    unittest.main()
