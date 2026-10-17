"""Importar los resultados de una ejecución desde un Word (app/importacion_resultados.py)."""
import os
import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document

from app import ejecucion_core as ej
from app import importacion_resultados as ir
from app import validacion_core as vc
from app import yaml_a_docx
from app.editor_core import read_chapters

RAIZ = Path(__file__).resolve().parents[1]
EJEMPLOS = RAIZ / "ejemplos" / "yaml"


def documento_con_resultados(ruta: Path, resultados: dict[str, list[str]], extra: list[tuple[str, list[str]]] = ()) -> None:
    """Word de las pruebas de ejemplo con filas de resultado al final de las tablas indicadas."""
    yaml_a_docx.generar_docx(EJEMPLOS / "pruebas", ruta)
    doc = Document(ruta)
    for tabla in doc.tables:
        filas = resultados.get(tabla.cell(0, 0).text.strip())
        for texto in filas or []:
            fila = tabla.add_row()
            fila.cells[0].merge(fila.cells[-1]).text = texto
    for ident, filas in extra:                     # tablas de pruebas que no están en el proyecto
        tabla = doc.add_table(rows=1 + len(filas), cols=2)
        tabla.cell(0, 0).text, tabla.cell(0, 1).text = ident, "Prueba de otro proyecto"
        for i, texto in enumerate(filas, start=1):
            tabla.cell(i, 0).merge(tabla.cell(i, 1)).text = texto
    doc.save(ruta)


def filas(resultado, observaciones="", ejecutor="Luis", fecha="15/09/2026 10:30"):
    return [f"Resultado obtenido: {resultado}", f"Observaciones: {observaciones}", f"Ejecutor prueba: {ejecutor}", f"Fecha ejecución: {fecha}"]


class Valores(unittest.TestCase):
    def test_fechas(self):
        self.assertEqual(ir.leer_fecha("15/09/2026"), "2026-09-15T00:00")
        self.assertEqual(ir.leer_fecha("15-9-26 9:05"), "2026-09-15T09:05")
        self.assertEqual(ir.leer_fecha("2026-09-15 10:30"), "2026-09-15T10:30")
        self.assertIsNone(ir.leer_fecha("septiembre"))
        self.assertIsNone(ir.leer_fecha("31/02/2026"))

    def test_propuestas(self):
        self.assertEqual(ir.proponer_valor("Correcto"), "ok")
        self.assertEqual(ir.proponer_valor("No OK"), "no_ok")
        self.assertEqual(ir.proponer_valor("N/A"), "no_aplica")
        self.assertEqual(ir.proponer_valor("OK, sin incidencias"), "ok")
        self.assertEqual(ir.proponer_valor(""), ir.SIN_RESULTADO)
        self.assertIsNone(ir.proponer_valor("Parcial"))
        self.assertEqual(ir.proponer_campo("Resultado obtenido"), "estado")
        self.assertEqual(ir.proponer_campo("Resultado esperado"), "ignorar")       # es de la prueba
        self.assertEqual(ir.proponer_campo("Fecha"), "ignorar")
        self.assertEqual(ir.proponer_campo("Fecha ejecución"), "fecha")


class Importar(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)
        shutil.copytree(EJEMPLOS, self.tmp / "yaml")
        self.docx = self.tmp / "resultados.docx"
        documento_con_resultados(self.docx, {
            "DEMO-LLAM-001": filas("Correcto"),
            "DEMO-LLAM-002": filas("NOK", "La puerta tarda 4 s", fecha="septiembre"),
            "DEMO-LLAM-003": filas("N/A"),
            "DEMO-NIV-001": filas("Parcial"),
            "DEMO-PUER-001": filas(""),
        }, extra=[("OTRO-001", filas("OK"))])
        self.analisis = ir.analizar(self.docx)
        self.mapeo = ir.mapeo_propuesto(self.analisis)

    def test_analisis_y_revision(self):
        familia = next(f for f in self.analisis["familias"] if f["tipo"] == "resultado")
        valores = {v["valor"]: v["propuesta"] for e in familia["etiquetas"] if e["clave"] == "resultadoobtenido" for v in e["valores"]}
        self.assertEqual(valores, {"Correcto": "ok", "NOK": "no_ok", "N/A": "no_aplica", "Parcial": None, "": ir.SIN_RESULTADO, "OK": "ok"})
        self.assertEqual(self.mapeo["campos"]["resultado"]["descripcion"], "ignorar")
        rev = ir.revision(ir.preparar(self.docx, self.mapeo, read_chapters(self.tmp / "yaml" / "pruebas")))
        # 14 pruebas del proyecto con tabla en el Word; OTRO-001 no está en el proyecto.
        self.assertEqual(rev["pruebas"], 14)
        self.assertEqual({k: v for k, v in rev["por_estado"].items() if v},
                         {"ok": 1, "no_ok": 1, "no_aplica": 1, "pendiente": 11})
        self.assertEqual({g["tipo"] for g in rev["avisos"]}, {"no_en_proyecto", "valor_sin_mapear", "fecha_no_entendida"})

    def test_crear_la_ejecucion(self):
        self.mapeo["valores"][ir.clave_valor("Parcial")] = "bloqueada"           # el usuario lo elige
        r = ir.importar(self.docx, self.mapeo, self.tmp / "ejecucion", self.tmp / "yaml",
                        "Regresión importada", {"entorno_pruebas": "Torre", "responsable": "Ana"})
        self.assertEqual((r["pruebas"], r["con_resultado"]), (14, 4))
        campaña = ej.cargar(self.tmp / "ejecucion")
        self.assertEqual(campaña["datos"]["estado"], "abierta")
        self.assertEqual(campaña["datos"]["entorno_pruebas"], "Torre")
        self.assertEqual(campaña["datos"]["seleccion"]["importado_de"], "resultados.docx")
        res = campaña["resultados"]
        self.assertEqual({k: v["estado"] for k, v in res.items()},
                         {"DEMO-LLAM-001": "ok", "DEMO-LLAM-002": "no_ok", "DEMO-LLAM-003": "no_aplica", "DEMO-NIV-001": "bloqueada"})
        self.assertEqual((res["DEMO-LLAM-001"]["fecha"], res["DEMO-LLAM-001"]["ejecutor"], res["DEMO-LLAM-001"]["origen"]),
                         ("2026-09-15T10:30", "Luis", "word"))
        # La fecha que no se entiende queda en observaciones (con la fecha de la importación).
        self.assertIn("La puerta tarda 4 s", res["DEMO-LLAM-002"]["observaciones"])
        self.assertIn("Fecha de ejecución: septiembre", res["DEMO-LLAM-002"]["observaciones"])
        # Las pruebas sin resultado siguen pendientes en la ejecución.
        resumen = ej.resumen(campaña, read_chapters(self.tmp / "yaml" / "pruebas"))
        self.assertEqual((resumen["total"], resumen["totales"]["pendiente"]), (14, 10))

    def test_api_dentro_de_la_validacion(self):
        from app.importacion_api import EditorImportacionApi
        anterior = os.environ.get("AZTARNA_AJUSTES")
        os.environ["AZTARNA_AJUSTES"] = str(self.tmp / "ajustes.json")
        self.addCleanup(lambda: os.environ.__setitem__("AZTARNA_AJUSTES", anterior) if anterior else os.environ.pop("AZTARNA_AJUSTES", None))
        api = EditorImportacionApi(self.tmp / "yaml")
        self.assertTrue(api.iniciar()["ok"])
        self.assertTrue(api.validacion_crear("V", "", {}, str(self.tmp / "val"))["ok"])
        a = api.importar_analizar(str(self.docx), "resultado")["datos"]
        self.assertEqual([e["valor"] for e in a["estados"]][-1], ir.SIN_RESULTADO)
        rev = api.importar_resultados_revisar(str(self.docx), a["mapeo"])["datos"]
        self.assertEqual(rev["pruebas"], 14)
        r = api.importar_resultados_ejecutar(str(self.docx), a["mapeo"], "Torre importada", {"entorno_pruebas": "Torre"}, True)
        self.assertTrue(r["ok"], r)
        self.assertIs(r["datos"]["campana"]["en_validacion"], True)
        self.assertTrue(r["datos"]["campana"]["carpeta"].startswith(str(self.tmp / "val")))
        v = vc.consolidar(vc.cargar(self.tmp / "val"))
        llam = next(p for p in v["pruebas"] if p["id"] == "DEMO-LLAM-001")
        self.assertEqual(llam["entornos"]["Torre"]["estado"], "ok")
        # Sin entorno de pruebas no se crea.
        r = api.importar_resultados_ejecutar(str(self.docx), a["mapeo"], "Sin entorno", {}, True)
        self.assertFalse(r["ok"])


if __name__ == "__main__":
    unittest.main()
