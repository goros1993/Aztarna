import base64
from io import BytesIO
from pathlib import Path
import struct
from tempfile import TemporaryDirectory
import unittest
from zipfile import ZipFile

from docx import Document

import convertir
import yaml_a_docx


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/"
    "p3sAAAAASUVORK5CYII="
)

EMF_MINIMAL = (
    struct.pack(
        "<II4i4iIIII",
        1, 108,
        0, 0, 1, 1,
        0, 0, 1000, 500,
        0x464D4520, 0x10000, 128, 2,
    )
    + bytes(108 - 56)
    + struct.pack("<IIIII", 14, 20, 0, 0, 20)
)


class ParsearTablaVarianteTests(unittest.TestCase):
    def test_exporta_emf_nativo_en_el_docx(self):
        with TemporaryDirectory() as temp:
            base_dir = Path(temp)
            emf = base_dir / "vector.emf"
            emf.write_bytes(EMF_MINIMAL)
            doc = Document()
            parrafo = doc.add_paragraph()
            yaml_a_docx.agregar_imagen(parrafo.add_run(), emf, yaml_a_docx.Cm(10))
            destino = base_dir / "salida.docx"
            doc.save(destino)

            with ZipFile(destino) as paquete:
                emf_en_docx = next(
                    nombre for nombre in paquete.namelist()
                    if nombre.endswith(".emf")
                )
                self.assertEqual(paquete.read(emf_en_docx), EMF_MINIMAL)

            reabierto = Document(destino)
            self.assertEqual(len(reabierto.inline_shapes), 1)

    def test_extrae_imagenes_del_contexto_y_de_la_tabla(self):
        doc = Document()
        doc.add_heading("Capítulo", level=1)
        doc.add_heading("Prueba", level=2)
        parrafo_contexto = doc.add_paragraph("Contexto con imagen ")
        parrafo_contexto.add_run().add_picture(BytesIO(PNG_1X1))

        tabla = doc.add_table(rows=2, cols=2)
        tabla.cell(0, 0).text = "TEST-IMG"
        tabla.cell(0, 1).text = "Prueba con imagen"
        tabla.cell(1, 0).text = "Situación inicial:"
        parrafo_seccion = tabla.cell(1, 0).add_paragraph("Ver el estado ")
        parrafo_seccion.add_run().add_picture(BytesIO(PNG_1X1))
        tabla.cell(1, 1).text = "Imagen general"
        tabla.cell(1, 1).paragraphs[0].add_run().add_picture(BytesIO(PNG_1X1))

        with TemporaryDirectory() as temp:
            salida = Path(temp)
            guardar = convertir.crear_guardador_imagenes(salida / "imagenes")
            capitulos = convertir.extraer_capitulos(doc, [], guardar)

            prueba = next(
                e["prueba"] for e in capitulos[0]["contenido"] if "prueba" in e
            )
            contexto = prueba["contexto"][0]

            self.assertEqual(contexto["texto"], "Contexto con imagen")
            self.assertEqual(contexto["imagenes"], ["imagenes/imagen_001.png"])
            self.assertEqual(
                prueba["situacion_inicial"][0]["imagenes"],
                ["imagenes/imagen_001.png"],
            )
            self.assertEqual(prueba["imagenes"], ["imagenes/imagen_001.png"])
            self.assertEqual(
                (salida / "imagenes" / "imagen_001.png").read_bytes(),
                PNG_1X1,
            )

    def test_exporta_imagenes_desde_contexto_lista_y_tabla(self):
        with TemporaryDirectory() as temp:
            base_dir = Path(temp)
            imagen = base_dir / "tiny.png"
            imagen.write_bytes(PNG_1X1)
            doc = Document()
            yaml_a_docx.escribir_contexto(
                doc,
                [{"texto": "Contexto", "imagenes": ["tiny.png"]}],
                base_dir,
            )
            yaml_a_docx.tabla_prueba(
                doc,
                {
                    "id": "TEST-IMG",
                    "situacion_inicial": [
                        {"texto": "Inicio", "imagenes": ["tiny.png"]}
                    ],
                    "imagenes": ["tiny.png"],
                },
                base_dir,
            )

            self.assertEqual(len(doc.inline_shapes), 3)

    def test_asocia_texto_suelto_a_pruebas_o_grupos_segun_titulos(self):
        doc = Document()
        doc.add_heading("Capítulo", level=1)
        doc.add_heading("Grupo", level=2)
        doc.add_paragraph("Contexto aplicable al grupo.")
        doc.add_heading("Subgrupo", level=3)
        doc.add_heading("Prueba uno", level=4)
        doc.add_paragraph("Contexto aplicable solo a la primera tabla.")
        primera = doc.add_table(rows=2, cols=2)
        primera.cell(0, 0).text = "TEST-001"
        primera.cell(0, 1).text = "Primera prueba"
        primera.cell(1, 0).text = "Pasos:\nHacer la prueba."
        doc.add_heading("Prueba dos", level=4)
        segunda = doc.add_table(rows=2, cols=2)
        segunda.cell(0, 0).text = "TEST-002"
        segunda.cell(0, 1).text = "Segunda prueba"
        segunda.cell(1, 0).text = "Pasos:\nHacer la prueba."

        avisos = []
        capitulos = convertir.extraer_capitulos(doc, avisos)
        elementos = capitulos[0]["contenido"]
        grupo = next(e for e in elementos if e.get("seccion") == "Grupo")
        primera_prueba = next(
            e["prueba"] for e in elementos if e.get("prueba", {}).get("id") == "TEST-001"
        )
        segunda_prueba = next(
            e["prueba"] for e in elementos if e.get("prueba", {}).get("id") == "TEST-002"
        )

        self.assertEqual(grupo["contexto"], ["Contexto aplicable al grupo."])
        self.assertEqual(
            primera_prueba["contexto"],
            ["Contexto aplicable solo a la primera tabla."],
        )
        self.assertNotIn("contexto", segunda_prueba)

    def test_mantiene_el_parseo_del_formato_estandar(self):
        doc = Document()
        tabla = doc.add_table(rows=4, cols=4)
        tabla.cell(0, 0).text = "TEST-001"
        tabla.cell(0, 1).text = "Prueba estándar"
        tabla.cell(1, 0).text = "Descripción: Verificar la llamada."
        tabla.cell(1, 1).text = "Componente: CPU_A"
        tabla.cell(1, 2).text = "Plataforma: HIL"
        tabla.cell(2, 0).text = (
            "Situación inicial:\nAscensor en reposo.\n"
            "Pasos:\nRealizar una llamada.\n"
            "Situación final:\nLa llamada se atiende."
        )
        tabla.cell(3, 0).text = "Automatizable: Sí"
        tabla.cell(3, 1).text = "Criticidad: Alta"
        tabla.cell(3, 2).text = "Entorno de pruebas: HIL"
        tabla.cell(3, 3).text = "Referencia: REF-1 Fecha: 2026 Versión: 1"

        avisos = []
        prueba = convertir.parsear(tabla, avisos)

        self.assertEqual(prueba["descripcion"], "Verificar la llamada.")
        self.assertEqual(prueba["situacion_inicial"], ["Ascensor en reposo."])
        # La situación final pasa al final del Ensayo; componente y plataforma ya no existen.
        self.assertEqual(prueba["pasos"], ["Realizar una llamada.", "La llamada se atiende."])
        self.assertNotIn("situacion_final", prueba)
        self.assertNotIn("componente", prueba)
        self.assertEqual(prueba["automatizable"], "si")
        self.assertEqual(prueba["entorno"], [])          # «HIL» no es una opción del esquema
        self.assertEqual(prueba["parametrizacion"], [])  # todas las pruebas con todos los campos
        self.assertEqual(list(prueba)[:4], ["id", "titulo", "descripcion", "parametrizacion"])
        self.assertEqual({a.tipo for a in avisos}, {"campos_eliminados", "valor_descartado"})

    def test_parsea_campos_de_formato_alternativo_y_preserva_campos_extra(self):
        doc = Document()
        tabla = doc.add_table(rows=6, cols=2)
        tabla.cell(0, 0).text = "DEMO-FUNC-007"
        tabla.cell(0, 1).text = "Accionamiento vertical regulado"
        tabla.cell(1, 0).text = (
            "Condiciones iniciales:\n"
            "• CPU_A en PLACA_B\n"
            "• CPU_A parametrizada"
        )
        tabla.cell(1, 1).text = (
            "ENSAYO(s) / Proceso(s): ENSAYO 1 (Tipos de movimiento)\n"
            "• Ascensor en una planta intermedia."
        )
        tabla.cell(2, 0).text = (
            "Resultado esperado: Indicado en ensayo. "
            "Criterio(s) de aceptación: Indicado en ensayo. "
            "Componente: CPU_A"
        )
        tabla.cell(2, 1).text = (
            "Maniobra(s): Triplex "
            "Criterio validación: Test "
            "Entorno de prueba(s): BANCO-02"
        )
        tabla.cell(3, 0).text = "Automatizable: Largo"
        tabla.cell(3, 1).text = "Entorno de pruebas: "
        tabla.cell(4, 0).text = "Ejecutor prueba:  Fecha ejecución: "
        tabla.cell(4, 1).text = "Resultado Obtenido: Decisión: Resultado:"
        tabla.cell(5, 0).text = "Criticidad: "
        tabla.cell(5, 1).text = "Referencia: <N.A> Fecha: 15/11/2024 Versión: 1.0"

        avisos = []
        prueba = convertir.parsear(tabla, avisos)

        self.assertEqual(prueba["id"], "DEMO-FUNC-007")
        self.assertNotIn("componente", prueba)
        self.assertEqual(prueba["automatizable"], "largo")
        self.assertEqual(prueba["situacion_inicial"], ["CPU_A en PLACA_B", "CPU_A parametrizada"])
        self.assertEqual(
            prueba["pasos"],
            ["ENSAYO 1 (Tipos de movimiento)", "Ascensor en una planta intermedia."],
        )
        # Valores que no encajan en las opciones se descartan con aviso.
        self.assertEqual((prueba["entorno"], prueba["maniobra"]), ([], []))
        # Resultado esperado, criterios y los campos de resultado se eliminan.
        self.assertNotIn("campos_adicionales", prueba)
        self.assertEqual((prueba["fecha"], prueba["version"]), ("15/11/2024", "1.0"))
        mensajes = [a.mensaje for a in avisos]
        self.assertIn("Maniobra: se descarta «Triplex»", mensajes)
        self.assertIn("Entorno de pruebas: se descarta «BANCO-02»", mensajes)
        self.assertIn("se elimina «Resultado esperado» (con datos)", mensajes)
        self.assertNotIn("se elimina «Decisión» (con datos)", mensajes)       # estaba vacío

    def test_lee_la_tabla_de_la_plantilla(self):
        """Formato nuevo (plantillas/tabla_prueba.docx): listas, varias opciones, requisitos y tags."""
        doc = Document()
        tabla = doc.add_table(rows=7, cols=3)
        tabla.cell(0, 0).text, tabla.cell(0, 1).text = "REQ_A_011", "Prueba nueva"
        tabla.cell(1, 0).text = "Descripción: Algo."
        tabla.cell(2, 0).text = "Parametrización:\nVelocidad 1 m/s\nCarga 0 %"
        tabla.cell(3, 0).text = "Situación inicial:\nEn reposo.\nEnsayo:\nPaso uno.\nPaso dos."
        tabla.cell(4, 0).text = "Maniobra: S/D"
        tabla.cell(4, 1).text = "Automatizable: Medio"
        tabla.cell(4, 2).text = "Entorno de pruebas: A / T"
        tabla.cell(5, 0).text = "Requisitos:\nREQ_1\nREQ_2"
        tabla.cell(5, 1).text = "Tags: smoke, puertas"
        tabla.cell(6, 0).text = "Fecha: 15/10/2026"
        tabla.cell(6, 1).text = "Versión: 1.2"
        avisos = []
        prueba = convertir.parsear(tabla, avisos)
        self.assertEqual(prueba["parametrizacion"], ["Velocidad 1 m/s", "Carga 0 %"])
        self.assertEqual(prueba["pasos"], ["Paso uno.", "Paso dos."])
        self.assertEqual(prueba["maniobra"], ["Simplex", "Dúplex"])
        self.assertEqual(prueba["entorno"], ["Automático", "Torre"])
        self.assertEqual(prueba["requisitos"], ["REQ_1", "REQ_2"])
        self.assertEqual(prueba["etiquetas"], ["smoke", "puertas"])
        self.assertEqual(prueba["automatizable"], "medio")
        self.assertFalse(avisos)

    def test_exporta_campos_adicionales_en_el_docx(self):
        doc = Document()
        yaml_a_docx.escribir_contexto(doc, ["Contexto asociado a la prueba."], docx_base_dir())
        yaml_a_docx.tabla_prueba(
            doc,
            {
                "id": "TEST-001",
                "campos_adicionales": {
                    "Resultado esperado": "Indicado en ensayo.",
                    "Requisito(s)": "REQ-001",
                },
            },
            docx_base_dir(),
        )

        self.assertEqual(doc.paragraphs[0].text, "Contexto asociado a la prueba.")
        textos = [celda.text for fila in doc.tables[0].rows for celda in fila.cells]
        self.assertTrue(any("Resultado esperado: Indicado en ensayo." in texto for texto in textos))
        self.assertTrue(any("Requisito(s): REQ-001" in texto for texto in textos))


def docx_base_dir():
    from pathlib import Path

    return Path(".")


class TamañoImagenesTests(unittest.TestCase):
    def test_las_imagenes_altas_no_superan_el_alto_maximo(self):
        from docx.shared import Cm
        from PIL import Image

        with TemporaryDirectory() as tmp:
            alta, ancha = Path(tmp) / "alta.png", Path(tmp) / "ancha.png"
            Image.new("RGB", (300, 900), "white").save(alta)
            Image.new("RGB", (900, 300), "white").save(ancha)
            doc = Document()
            yaml_a_docx.agregar_imagen(doc.add_paragraph().add_run(), alta, Cm(10))
            yaml_a_docx.agregar_imagen(doc.add_paragraph().add_run(), ancha, Cm(10))
            primera, segunda = doc.inline_shapes
            self.assertLessEqual(primera.height, Cm(yaml_a_docx.ALTO_MAX_IMAGEN_CM) + 1)
            self.assertLess(primera.width, Cm(10))
            self.assertEqual(segunda.width, Cm(10))


if __name__ == "__main__":
    unittest.main()


def documento_mixto():
    """Word con pruebas y tablas que no son pruebas, como las del documento real."""
    doc = Document()
    doc.add_heading("Capítulo", level=1)
    historial = doc.add_table(rows=2, cols=3)
    for i, texto in enumerate(("Modificado por:", "Cambio", "Edición")):
        historial.cell(0, i).text = texto
    for i, texto in enumerate(("Ana", "Primera versión", "1")):
        historial.cell(1, i).text = texto
    doc.add_paragraph("Texto después del historial.")
    doc.add_heading("Parámetros de rescate", level=2)
    parametros = doc.add_table(rows=3, cols=2)
    parametros.cell(0, 0).merge(parametros.cell(0, 1)).text = "Parametros de Rescates CPU"
    parametros.cell(1, 0).text = "Velocidad"
    parametros.cell(1, 1).text = "0,3 m/s"
    parametros.cell(2, 0).text = "Tiempo máximo"
    parametros.cell(2, 1).text = "30 s"
    doc.add_paragraph("Notas sobre los parámetros.")
    doc.add_heading("Prueba de llamada", level=3)
    prueba = doc.add_table(rows=3, cols=2)
    prueba.cell(0, 0).text = "LLAM_001"
    prueba.cell(0, 1).text = "Llamada de cabina"
    prueba.cell(1, 0).text = "Pasos:\nPulsar el botón."
    prueba.cell(2, 0).text = "Automatizable: Sí"
    doc.add_heading("Prueba repetida", level=3)
    repetida = doc.add_table(rows=2, cols=2)
    repetida.cell(0, 0).text = "LLAM_001"
    repetida.cell(0, 1).text = "Otra con el mismo ID"
    repetida.cell(1, 0).text = "Pasos:\nAlgo."
    doc.add_heading("Prueba con espacio", level=3)
    espacio = doc.add_table(rows=4, cols=1)
    espacio.cell(0, 0).text = "BOR_LLAM _002"
    espacio.cell(1, 0).text = "Componente: CPU"
    espacio.cell(2, 0).text = "Criticidad: Alta"
    espacio.cell(3, 0).text = "Plataforma: HIL"
    doc.add_table(rows=2, cols=2)                                   # vacía
    dudosa = doc.add_table(rows=2, cols=2)
    dudosa.cell(0, 0).text = "Revisión"
    dudosa.cell(1, 0).text = "Fecha: 01/02/2024"
    dudosa.cell(1, 1).text = "Versión: 3"
    return doc


class ClasificarTablasTests(unittest.TestCase):
    def test_clasifica_pruebas_y_tablas_que_no_lo_son(self):
        tablas = documento_mixto().tables
        clases = [convertir.clasificar_tabla(t)[0] for t in tablas]
        self.assertEqual(clases, ["informativa", "informativa", "prueba", "prueba", "prueba", "informativa", "dudosa"])
        self.assertEqual(convertir.clasificar_tabla(tablas[5])[1], "tabla vacía")

    def test_conserva_las_tablas_que_no_son_pruebas_en_su_contexto(self):
        avisos = []
        capitulo = convertir.extraer_capitulos(documento_mixto(), avisos)[0]
        contenido = capitulo["contenido"]
        claves = [next(iter(e)) for e in contenido]
        # Antes de cualquier título, la tabla y el texto son del capítulo, en orden.
        self.assertEqual(claves[:3], ["tabla", "texto", "seccion"])
        self.assertEqual(contenido[0]["tabla"]["filas"][1], ["Ana", "Primera versión", "1"])
        self.assertEqual(contenido[1]["texto"], "Texto después del historial.")
        # Tras un título, la tabla queda en el contexto de la sección, con el texto que la sigue.
        seccion = contenido[2]
        self.assertEqual(seccion["seccion"], "Parámetros de rescate")
        self.assertEqual(seccion["contexto"][0]["tabla"]["filas"][0], [{"texto": "Parametros de Rescates CPU", "columnas": 2}])
        self.assertEqual(seccion["contexto"][1], "Notas sobre los parámetros.")
        pruebas = [e["prueba"] for e in contenido if "prueba" in e]
        self.assertEqual([p["id"] for p in pruebas], ["LLAM_001", "LLAM_001", "BOR_LLAM _002"])
        self.assertEqual(pruebas[0]["nombre"], "Prueba de llamada")
        # Las tablas que no son pruebas no generan avisos de pruebas.
        self.assertFalse(any("Modificado" in a or "Parametros" in a for a in avisos if a.tipo != "tabla_dudosa"))
        tipos = {a.tipo for a in avisos}
        self.assertTrue({"tabla_informativa", "tabla_dudosa", "id_duplicado", "id_no_valido"} <= tipos)
        duplicado = next(a for a in avisos if a.tipo == "id_duplicado")
        self.assertEqual((duplicado.id, duplicado.tabla, duplicado.capitulo), ("LLAM_001", 4, "Capítulo"))
        self.assertIn("tabla 3", duplicado.mensaje)

    def test_avisos_agrupados_y_docx_con_las_tablas(self):
        avisos = []
        capitulos = convertir.extraer_capitulos(documento_mixto(), avisos)
        with TemporaryDirectory() as tmp:
            convertir.escribir_avisos(avisos, Path(tmp))
            texto = (Path(tmp) / "avisos.txt").read_text(encoding="utf-8")
            self.assertTrue(texto.startswith(f"{len(avisos)} avisos"))
            self.assertIn("== ID repetido (1) ==", texto)
            self.assertIn("Capítulo · tabla 4 · LLAM_001: ID repetido", texto)
            csv = (Path(tmp) / "avisos.csv").read_text(encoding="utf-8-sig").splitlines()
            self.assertEqual(csv[0], "gravedad;tipo;capitulo;tabla;id;mensaje")
            self.assertEqual(len(csv), len(avisos) + 1)

            doc = yaml_a_docx.preparar_documento(1)
            yaml_a_docx.escribir_capitulos(doc, [(c["capitulo"], c["contenido"], Path(tmp)) for c in capitulos])
            textos = [celda.text for tabla in doc.tables for fila in tabla.rows for celda in fila.cells]
            self.assertIn("Primera versión", textos)
            self.assertIn("Parametros de Rescates CPU", textos)
            self.assertEqual(len(doc.tables), 7)
            parrafos = [p.text for p in doc.paragraphs]
            self.assertIn("Texto después del historial.", parrafos)
            self.assertLess(parrafos.index("Parámetros de rescate"), parrafos.index("Notas sobre los parámetros."))

    def test_tabla_en_el_contexto_de_la_prueba_siguiente(self):
        doc = Document()
        doc.add_heading("Capítulo", level=1)
        doc.add_heading("Prueba con tabla previa", level=2)
        doc.add_paragraph("Valores de partida:")
        valores = doc.add_table(rows=2, cols=2)
        valores.cell(0, 0).text, valores.cell(0, 1).text = "Parámetro", "Valor"
        valores.cell(1, 0).text, valores.cell(1, 1).text = "Velocidad", "1 m/s"
        prueba = doc.add_table(rows=2, cols=2)
        prueba.cell(0, 0).text, prueba.cell(0, 1).text = "T_001", "Prueba"
        prueba.cell(1, 0).text = "Pasos:\nHacerlo."
        contenido = convertir.extraer_capitulos(doc, [])[0]["contenido"]
        self.assertEqual(len(contenido), 1)
        contexto = contenido[0]["prueba"]["contexto"]
        self.assertEqual(contexto[0], "Valores de partida:")
        self.assertEqual(contexto[1]["tabla"]["filas"][1], ["Velocidad", "1 m/s"])
